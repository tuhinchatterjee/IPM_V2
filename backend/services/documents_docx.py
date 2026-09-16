"""A working paper as a Word file.

Not the same job as the generated reports. Those are assembled from results
and their writer worries about evidence registers and content hashes. This
one takes a paper a person WROTE — Markdown in a column — and lays it out so
it can be attached to an email and read by somebody who will never open
CreditProbe.

So the decisions here are about a reader, not about reproducibility:

* The cover states status, owner, period and revision, because the first
  question about a paper is always whether it is the current one.
* The data and model versions are on the cover too. A paper about a
  59,449-facility book that reads like a paper about a 19,745-facility book
  is the defect §24 exists to prevent, and the only defence a reader holding
  the file has is the version printed on it.
* The comments are an appendix rather than margin notes. A committee reads
  the paper; a reviewer reads the paper and the thread. Margin notes in a
  .docx that came from a different comment model would be Word's tracked
  comments, which look like somebody edited the file.
* Markdown is rendered rather than printed. A heading that arrives as
  "## Findings" in a Word document tells the reader the tool did not read
  what it was given.
"""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document as WordDocument
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from backend.models.platform import DOCUMENT_STATUS_LABEL

INK = RGBColor(0x16, 0x23, 0x2F)
MUTED = RGBColor(0x6C, 0x7A, 0x8C)

DOCX_VERSION = "retail-document-docx-1.0.0"

BULLET = re.compile(r"^\s*[-*]\s+(.*)$")
NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
HEADING = re.compile(r"^(#{1,4})\s+(.*)$")
TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
RULE = re.compile(r"^\s*([-*_])\1{2,}\s*$")


def _field(paragraph: Any, instruction: str) -> None:
    """A Word field — PAGE, NUMPAGES — rather than a number we guessed."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instruction} "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(text)
    run._r.append(end)


def _cells(line: str) -> list[str]:
    inner = TABLE_ROW.match(line)
    return [one.strip() for one in inner.group(1).split("|")] if inner else []


def _is_divider(cells: list[str]) -> bool:
    return bool(cells) and all(
        set(one) <= set("-: ") and "-" in one for one in cells)


def _table(document: Any, block: list[list[str]]) -> None:
    if not block:
        return
    width = max(len(row) for row in block)
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    for at, row in enumerate(block):
        cells = table.add_row().cells
        for column in range(width):
            cell = cells[column]
            cell.text = row[column] if column < len(row) else ""
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
                    if at == 0:
                        run.bold = True
    document.add_paragraph()


def _body(document: Any, markdown: str) -> None:
    """Markdown as Word. Headings, lists, tables, rules and paragraphs."""
    pending: list[list[str]] = []
    for line in (markdown or "").splitlines():
        cells = _cells(line)
        if cells:
            if not _is_divider(cells):
                pending.append(cells)
            continue
        if pending:
            _table(document, pending)
            pending = []

        if not line.strip():
            continue
        if RULE.match(line):
            document.add_paragraph()
            continue
        head = HEADING.match(line)
        if head:
            level = min(len(head.group(1)), 4)
            written = document.add_heading(head.group(2).strip(), level=level)
            for run in written.runs:
                run.font.color.rgb = INK
            continue
        bullet = BULLET.match(line)
        if bullet:
            document.add_paragraph(bullet.group(1).strip(),
                                   style="List Bullet")
            continue
        numbered = NUMBERED.match(line)
        if numbered:
            document.add_paragraph(numbered.group(1).strip(),
                                   style="List Number")
            continue
        document.add_paragraph(line.strip())
    if pending:
        _table(document, pending)


def write(session: Any, row: Any) -> bytes:
    """The document, as .docx bytes."""
    from backend.services import documents as docs
    from backend.services import workflow as wf

    document = WordDocument()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK

    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = (f"{row.title} · "
                   f"{DOCUMENT_STATUS_LABEL.get(row.status, row.status)}"
                   + (f" · {row.as_of}" if row.as_of else ""))
    if header.runs:
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"Revision {row.revision} · Internal · Page ")
    run.font.size = Pt(7)
    run.font.color.rgb = MUTED
    _field(footer, "PAGE")

    title = document.add_heading(row.title, level=0)
    for one in title.runs:
        one.font.color.rgb = INK
    if row.summary:
        document.add_paragraph(row.summary)

    facts = [
        ("Type", row.kind or "not recorded"),
        ("Product", row.product or "Retail"),
        ("Status", DOCUMENT_STATUS_LABEL.get(row.status, row.status)),
        ("As of", row.as_of or "not recorded"),
        ("Owner", row.owner_name or "not recorded"),
        ("Revision", f"{row.revision}"
         + ("" if row.is_current else " — superseded, kept for the record")),
    ]
    if row.approved_at:
        facts.append(("Approved",
                      f"{row.approved_by or 'unattributed'} · "
                      f"{row.approved_at:%Y-%m-%d}"))
    # The versions the figures rest on, on the face of the file. A reader
    # holding only the .docx has no other way to tell which book it is about.
    for key, value in sorted((row.data_versions or {}).items()):
        facts.append((key.replace("_", " ").title(), str(value)))
    _table(document, [["Item", "Value"], *[[a, b] for a, b in facts]])

    document.add_section(WD_SECTION.NEW_PAGE)
    _body(document, row.body)

    evidence = list(row.evidence or [])
    attachments = sorted(row.attachments, key=lambda one: one.id)
    if evidence or attachments:
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_heading("Supporting evidence", level=1)
        rows = [["What", "Kind", "Where"]]
        for one in evidence:
            rows.append([str(one.get("label") or one.get("id") or ""),
                         str(one.get("kind") or "reference"),
                         str(one.get("href") or "in CreditProbe")])
        for one in attachments:
            rows.append([one.label or one.filename, one.role or "file",
                         f"{one.filename} ({one.size_bytes:,} bytes)"])
        _table(document, rows)

    comments = wf.comments("document", str(row.id))
    if comments:
        document.add_heading("Review comments", level=1)
        note = document.add_paragraph()
        said = note.add_run(
            "Recorded in CreditProbe against this document. Reproduced here "
            "so the paper and the review travel together.")
        said.italic = True
        said.font.size = Pt(8)
        said.font.color.rgb = MUTED
        rows = [["Author", "When", "Status", "Comment"]]
        for one in comments:
            rows.append([
                str(one.get("author") or one.get("author_id") or "—"),
                str(one.get("created_at") or "")[:19].replace("T", " "),
                "resolved" if one.get("resolved") else "open",
                str(one.get("body") or ""),
            ])
        _table(document, rows)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


__all__ = ["DOCX_VERSION", "write"]
