"""
A Document's tables as a workbook. Playbook §11, PB-021.

A report's numbers, in the form somebody can actually work with. One sheet per
table, plus a contents sheet naming where each came from, so a workbook opened
six months later still says which report and which section produced it.

Values are written as they were computed. A figure that reached the document as
"19.20" is written as the number 19.2 where it parses cleanly as one, and as
text where it does not — because writing "n/a" into a numeric cell as 0 is the
spreadsheet version of inventing a result.
"""

from __future__ import annotations

import io
import re

from backend.playbook import document as D

_NUMERIC = re.compile(r"^-?[\d,]+(\.\d+)?%?$")
#: Excel refuses these in a sheet name, and truncates past 31 characters.
_BAD_SHEET = re.compile(r"[\\/*?:\[\]]")


def _sheet_name(raw: str, used: set[str]) -> str:
    name = _BAD_SHEET.sub("-", (raw or "Table").strip())[:31] or "Table"
    candidate, n = name, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = name[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _value(text: str):
    raw = (text or "").strip()
    if not raw or not _NUMERIC.match(raw):
        return raw
    if raw.endswith("%"):
        return raw  # a percentage keeps its sign rather than becoming a ratio
    try:
        return float(raw.replace(",", ""))
    except ValueError:
        return raw


def write(doc: D.Document) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FF0B2436")
    head_fill = PatternFill("solid", fgColor="FFEEF2F7")

    wb = Workbook()
    contents = wb.active
    contents.title = "CONTENTS"
    contents.append(["Report", doc.title or ""])
    for key, value in (doc.meta or {}).items():
        contents.append([key.replace("_", " ").title(), str(value)])
    contents.append([])
    contents.append(["Sheet", "Section"])
    for cell in contents[contents.max_row]:
        cell.font = head_font
        cell.fill = head_fill

    used: set[str] = {"contents"}
    written = 0
    for section in doc.sections:
        for block in section.blocks:
            if block.kind != D.TABLE:
                continue
            columns = block.data.get("columns") or []
            if not columns:
                continue
            name = _sheet_name(section.heading or "Table", used)
            ws = wb.create_sheet(name)
            ws.append([str(c) for c in columns])
            for cell in ws[1]:
                cell.font = head_font
                cell.fill = head_fill
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            for row in block.data.get("rows") or []:
                ws.append([_value(v) for v in list(row)[: len(columns)]])
            for i, column in enumerate(columns, start=1):
                width = max(
                    [len(str(column))]
                    + [len(str((r or [])[i - 1])) if len(r) >= i else 0
                       for r in (block.data.get("rows") or [])]
                )
                ws.column_dimensions[get_column_letter(i)].width = min(
                    max(width + 2, 10), 50)
            ws.freeze_panes = "A2"
            contents.append([name, section.heading or ""])
            written += 1

    if not written:
        note = wb.create_sheet("NOTE")
        note["A1"] = ("This report contains no tabular results, so there is "
                      "nothing to put in a workbook.")

    contents.column_dimensions["A"].width = 34
    contents.column_dimensions["B"].width = 60

    if doc.sources:
        src = wb.create_sheet("SOURCES")
        src.append(["Locator"])
        for cell in src[1]:
            cell.font = head_font
            cell.fill = head_fill
        for locator in doc.sources:
            src.append([locator])
        src.column_dimensions["A"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
