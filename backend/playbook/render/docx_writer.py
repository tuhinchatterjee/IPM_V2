"""
A Document as a Word file. Playbook §11.

"Do not deliver raw Markdown inside a DOCX container" is the requirement, and
it is a real risk: it is entirely possible to write a .docx whose every
paragraph begins with a literal `##`. So this writer uses Word's own machinery —
built-in heading styles so the navigation pane and any later table of contents
work, a real table style with a repeating header row, page numbering in the
footer as a field rather than as typed text, and margins a committee paper can
be printed at.

Sources are collected into a references appendix rather than left inline,
because `[[xlsx://ECL!B12]]` in the middle of a sentence is a debugging artefact,
not a citation somebody can read.
"""

from __future__ import annotations

import io

from backend.playbook import document as D
from backend.playbook.render import inline
from backend.playbook.render import shell as shell_of


def _add_field(paragraph, instruction: str) -> None:
    """Insert a Word field (used for PAGE / NUMPAGES).

    Typed page numbers are wrong the moment the document reflows, so the footer
    carries the field Word updates itself.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _runs(paragraph, text: str, *, size=None, colour=None, italic=False):
    """Write `text` into `paragraph` with its inline marks applied.

    Every paragraph used to be one plain run, so `**Model ID:**` reached the
    reader with its asterisks. The canonical Block has no inline model and
    deliberately still does not — see `render/inline.py` for why this belongs
    in the writer.
    """
    from docx.shared import Pt

    for run in inline.runs(text):
        r = paragraph.add_run(run.text)
        r.bold = run.bold or None
        r.italic = run.italic or italic or None
        if run.code:
            r.font.name = "Consolas"
        if size is not None:
            r.font.size = Pt(size)
        if colour is not None:
            r.font.color.rgb = colour
    return paragraph


def _toc_field(paragraph) -> None:
    """A real Word TOC field, not a typed list.

    Word builds and renumbers it, so the entries stay right when the document
    reflows — which a typed list cannot do, and which is the whole reason
    chapter 09 asks for "usable navigation" rather than "a list of headings".
    `\\o "1-3"` takes heading levels 1 to 3, `\\h` makes them links, `\\z`
    hides page numbers in web view and `\\u` uses the outline levels.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and choose Update Field to build the contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(placeholder)
    run._r.append(end)


def _repeat_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def write(doc: D.Document) -> bytes:
    from docx import Document as Docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    out = Docx()
    front = shell_of.shell_of(doc)
    out.core_properties.title = front.title
    out.core_properties.author = "CreditProbe"

    for section in out.sections:
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)

        header = section.header.paragraphs[0]
        # `front.title` is one line with no marks. The raw title used to go in
        # here, newlines and all, and a newline in a running header is what
        # rendered as black boxes in the PDF.
        header.text = front.title
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x6C, 0x7A, 0x8C)

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Page ")
        run.font.size = Pt(8)
        _add_field(footer, "PAGE")
        run = footer.add_run(" of ")
        run.font.size = Pt(8)
        _add_field(footer, "NUMPAGES")
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x6C, 0x7A, 0x8C)

    # ---- cover -------------------------------------------------------
    # The title block sits a third of the way down, the same as the PDF's, so
    # the two formats deliver the same document rather than two arrangements
    # of it. Empty paragraphs rather than a frame: python-docx has no vertical
    # positioning, and a spacer a reader can delete is better than a text box
    # they cannot edit.
    if front.has_contents:
        for _ in range(6):
            out.add_paragraph()
    title = out.add_heading(front.title, level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x0B, 0x24, 0x36)
    if front.subtitle:
        _runs(out.add_paragraph(), front.subtitle, size=11,
              colour=RGBColor(0x6C, 0x7A, 0x8C))

    if front.facts:
        # A two-column record rather than the middot-joined line this used to
        # be. "Model ID · Owner · Development date · Status" run together
        # reads as a caption; a committee paper's front matter is a table.
        facts = out.add_table(rows=0, cols=2)
        facts.style = "Light List Accent 1"
        for name, value in front.facts:
            cells = facts.add_row().cells
            cells[0].width = Inches(1.6)
            for cell, text, bold in ((cells[0], name, True),
                                     (cells[1], value, False)):
                cell.text = ""
                run = cell.paragraphs[0].add_run(text)
                run.bold = bold or None
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x16, 0x23, 0x2F)
        out.add_paragraph()

    # ---- contents ----------------------------------------------------
    if front.has_contents:
        out.add_page_break()
        out.add_heading("Contents", level=1)
        _toc_field(out.add_paragraph())
        out.add_page_break()

    for section in doc.sections:
        if section.heading:
            # A heading holds no runs, so its marks are stripped rather than
            # applied — `## **Scope**` must not print its asterisks either.
            out.add_heading(inline.plain(section.heading),
                            level=min(max(section.level, 1), 4))
        for block in section.blocks:
            if block.kind == D.TABLE:
                _write_table(out, block)
            elif block.kind == D.BULLETS:
                for item in block.data.get("items", []):
                    _runs(out.add_paragraph(style="List Bullet"), item)
            elif block.kind == D.NUMBERS:
                for item in block.data.get("items", []):
                    _runs(out.add_paragraph(style="List Number"), item)
            elif block.kind == D.CALLOUT:
                _runs(out.add_paragraph(), block.text, italic=True,
                      colour=RGBColor(0x0B, 0x24, 0x36))
            elif block.text:
                _runs(out.add_paragraph(), block.text)

    sources = doc.sources
    if sources:
        out.add_page_break()
        out.add_heading("Sources", level=1)
        p = out.add_paragraph(
            "Every figure in this report is traceable to one of the following."
        )
        for run in p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x6C, 0x7A, 0x8C)
        for locator in sources:
            out.add_paragraph(locator, style="List Bullet")

    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def _write_table(out, block: D.Block) -> None:
    from docx.shared import Pt, RGBColor

    columns = block.data.get("columns") or []
    rows = block.data.get("rows") or []
    if not columns:
        return
    table = out.add_table(rows=1, cols=len(columns))
    table.style = "Light Grid Accent 1"
    header = table.rows[0]
    for i, name in enumerate(columns):
        cell = header.cells[i]
        cell.text = inline.plain(str(name))
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x0B, 0x24, 0x36)
    _repeat_header(header)

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row[: len(columns)]):
            cells[i].text = ("" if value is None
                             else inline.plain(str(value)))
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    out.add_paragraph()
