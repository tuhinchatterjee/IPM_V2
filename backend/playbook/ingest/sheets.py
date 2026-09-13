"""
Reading a workbook or a CSV. Playbook §7.

The rule that shapes this reader
--------------------------------
**A formula is not a result.** openpyxl will hand you either the formula string
or the value Excel last cached, depending on how the file is opened. A workbook
saved by a tool that never calculated has formulas and no cached values, and a
reader that takes `"=B4*C4"` as a figure will produce a report full of confident
nonsense. So the workbook is opened TWICE — once for formulas, once for cached
values — and a cell with a formula but no cached value is recorded as
*uncomputed* rather than as a number.

**A hidden sheet is a decision somebody made.** It is not read as if it were
visible, and it is not silently dropped either: it is listed in the manifest so
an answer can say what it did not use.

Merged headings are resolved to the value of their top-left anchor, because a
merged header cell reads as `None` everywhere except its anchor, and a column
whose header is None is a column nobody can cite.
"""

from __future__ import annotations

import csv
import io
import re
from decimal import Decimal

from backend.playbook.ingest.types import (
    SHEET_RANGE,
    TABLE,
    Chunk,
    Manifest,
    ReadResult,
    UnreadableSource,
)

#: How many rows of a sheet are carried as evidence. A sheet larger than this is
#: summarised and its full extent recorded, rather than truncated in silence.
MAX_ROWS = 500


def _shown(value: object, number_format: str) -> str:
    """The value as the workbook shows it, or exactly as stored.

    A spreadsheet holds 0.5593220338983 and displays 0.559, because its author
    set the format that says so. Both are true and they are the same fact; only
    one of them belongs in a committee report. Reading the value and discarding
    the format left the evidence ledger quoting seventeen significant digits,
    which is what a grounded report then had to print to survive the check.

    The format is the workbook's own statement about presentation, so it is
    honoured and nothing is invented: a cell with no stated format (`General`)
    comes back exactly as stored.
    """
    from backend.playbook import calc

    if value is None:
        return ""
    if not isinstance(value, (int, float, Decimal)) or isinstance(value, bool):
        return str(value)
    dp = _decimals_in(number_format)
    if dp is None:
        return str(value)
    try:
        return calc.present(str(value), dp)
    except calc.CalculationError:
        return str(value)


def _decimals_in(number_format: str) -> int | None:
    """How many decimals a number format shows, or None when it says nothing.

    Excel formats are a small language; this reads the one thing needed from
    it — the digits after the decimal point in the positive section — and
    declines to guess at anything else. `General` states no precision, so it
    gets none.
    """
    fmt = (number_format or "").split(";")[0].strip()
    if not fmt or fmt.lower() == "general":
        return None
    fmt = re.sub(r'"[^"]*"', "", fmt)          # literal text
    fmt = re.sub(r"\\.", "", fmt)              # escaped characters
    if "%" in fmt or "E+" in fmt.upper():
        # A percentage or scientific format scales the value as well as
        # shaping it, which is a conversion rather than a presentation.
        return None
    match = re.search(r"\.([0#?]+)", fmt)
    return len(match.group(1)) if match else 0


def _merged_lookup(ws) -> dict[tuple[int, int], object]:
    """Map every covered cell of a merged range to its anchor's value."""
    out: dict[tuple[int, int], object] = {}
    for rng in ws.merged_cells.ranges:
        anchor = ws.cell(row=rng.min_row, column=rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                out[(r, c)] = anchor
    return out


def _read_workbook(content: bytes, filename: str) -> ReadResult:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - declared dependency
        raise UnreadableSource("openpyxl is not installed") from exc

    manifest = Manifest(format="xlsx")
    try:
        # data_only=True gives Excel's cached results; the default gives the
        # formulas. Both are needed to tell a computed 19.2 from an uncomputed
        # "=B4*C4" that merely looks like one.
        values_wb = load_workbook(io.BytesIO(content), data_only=True,
                                  read_only=False)
        formula_wb = load_workbook(io.BytesIO(content), data_only=False,
                                   read_only=False)
    except Exception as exc:
        raise UnreadableSource(
            f"{filename or 'This file'} could not be opened as a workbook."
        ) from exc

    chunks: list[Chunk] = []
    uncomputed = 0

    for name in values_wb.sheetnames:
        vws = values_wb[name]
        fws = formula_wb[name]
        if vws.sheet_state != "visible":
            manifest.skip(f"sheet {name!r}",
                          f"the sheet is {vws.sheet_state} in the workbook")
            continue
        if vws.max_row is None or vws.max_row < 1:
            manifest.skip(f"sheet {name!r}", "the sheet is empty")
            continue

        merged = _merged_lookup(vws)
        rows: list[list[str]] = []
        raw_rows: list[list[str]] = []
        refs: list[list[str]] = []
        limit = min(vws.max_row, MAX_ROWS)
        for r in range(1, limit + 1):
            row: list[str] = []
            raw_row: list[str] = []
            ref_row: list[str] = []
            for c in range(1, (vws.max_column or 1) + 1):
                cell = vws.cell(row=r, column=c)
                cached = cell.value
                if cached is None and (r, c) in merged:
                    cached = merged[(r, c)]
                raw = fws.cell(row=r, column=c).value
                ref_row.append(f"{get_column_letter(c)}{r}")
                if cached is None and isinstance(raw, str) and raw.startswith("="):
                    uncomputed += 1
                    row.append("<uncomputed formula>")
                    raw_row.append("<uncomputed formula>")
                else:
                    exact = "" if cached is None else str(cached)
                    raw_row.append(exact)
                    row.append(_shown(cached, cell.number_format))
            if any(cell for cell in row):
                rows.append(row)
                raw_rows.append(raw_row)
                refs.append(ref_row)

        if not rows:
            manifest.skip(f"sheet {name!r}", "the sheet holds no values")
            continue

        header, *body = rows
        anchor = f"{get_column_letter(1)}1"
        chunks.append(Chunk(
            SHEET_RANGE,
            f"xlsx://{name}!{anchor}",
            f"{name}: " + " | ".join(header),
            [name],
            # `rows` is what the workbook SHOWS; `raw_rows` is what it
            # STORES, cell for cell, and `cells` is where each one lives.
            # Provenance keeps all three, so a report may quote the governed
            # presentation of a fact without losing the exact source value or
            # the address that proves where it came from.
            {"sheet": name, "columns": header, "rows": body,
             "raw_columns": raw_rows[0] if raw_rows else header,
             "raw_rows": raw_rows[1:],
             "cells": refs[1:], "header_cells": refs[0] if refs else [],
             "total_rows": vws.max_row, "rows_read": len(rows)},
            ordinal=len(chunks),
        ))
        manifest.read.append(f"sheet {name!r} ({len(rows)} of {vws.max_row} rows)")
        if vws.max_row > MAX_ROWS:
            manifest.warnings.append(
                f"sheet {name!r} has {vws.max_row} rows; the first {MAX_ROWS} "
                "were read and the remainder were not."
            )

    if not chunks:
        raise UnreadableSource(
            f"{filename or 'This workbook'} has no readable visible sheets."
        )
    if uncomputed:
        manifest.warnings.append(
            f"{uncomputed} cell(s) hold a formula with no cached result. Those "
            "are reported as uncomputed and are not used as figures."
        )

    links = getattr(values_wb, "_external_links", None)
    if links:
        manifest.warnings.append(
            f"{len(links)} external workbook link(s) were not followed."
        )
    return ReadResult(chunks, manifest)


def _read_csv(content: bytes, filename: str) -> ReadResult:
    manifest = Manifest(format="csv")
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - check() has already refused undecodable text
        raise UnreadableSource(f"{filename or 'This file'} is not readable text.")

    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(r)]
    if not rows:
        raise UnreadableSource(f"{filename or 'This file'} holds no rows.")

    header, *body = rows
    read_rows = body[:MAX_ROWS]
    manifest.read.append(f"{len(read_rows)} of {len(body)} data rows")
    if len(body) > MAX_ROWS:
        manifest.warnings.append(
            f"{len(body)} rows present; the first {MAX_ROWS} were read."
        )
    chunk = Chunk(TABLE, "csv://rows", " | ".join(header), [],
                  {"columns": header, "rows": read_rows,
                   "total_rows": len(body), "rows_read": len(read_rows)})
    return ReadResult([chunk], manifest)


def read(content: bytes, *, filename: str = "", kind: str = "xlsx") -> ReadResult:
    if kind == "csv":
        return _read_csv(content, filename)
    return _read_workbook(content, filename)
