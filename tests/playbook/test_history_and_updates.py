"""The three contracts the dashboard needed. §20, §23, §24.

**History** is assembled from the trails the rows already keep, never stored
beside them. A second event log could disagree with what it describes, and the
one that disagrees is always the one somebody reads.

**Check for updates** is a read that produces a proposal. New data is a reason
to tell somebody, not a licence to edit a governed pack — so the test that
matters most here is that nothing changed.

**Uploaded metric updates** are suggestions and say so. A column header
resembling a tracked metric is not evidence that it is that metric.
"""

from __future__ import annotations

import io

import pytest

from backend.exports import playbook_contract as contract
from backend.models.playbook import PlaybookMetricBinding
from backend.playbook import document as D
from backend.playbook import repository as repo
from backend.playbook import service
from backend.playbook.intelligence import adopt, history, refresh
from backend.playbook.intelligence import binding as bind
from backend.playbook.intelligence import governance as gov
from backend.playbook.intelligence import sections as sect

ACTOR = "user:7"
BODY = "considered judgement " * 12


def _doc(*headings: str) -> D.Document:
    return D.parse("".join(f"## {h}\n\n{BODY}\n\n" for h in headings),
                   title="Report")


def _versioned(db, workspace, doc, *, artifact=None, version=1):
    if artifact is None:
        artifact = repo.create_artifact(db, workspace.id, kind="report",
                                        title="Report")
        db.flush()
    row = repo.new_version(db, artifact, content=doc.as_dict(),
                           source_manifest={},
                           content_hash=doc.content_hash(), validation={})
    db.flush()
    adopt.adopt(db, workspace.id, artifact.id, version_id=row.id,
                version=version, doc=doc)
    return artifact, row


def _governed(db, workspace, *, metric_id="ifrs9.coverage_ratio",
              label="Coverage ratio", value="5.86", display="5.86%",
              section_key="", unit="percent"):
    [row] = bind.apply(db, workspace.id, bind.from_export(
        [contract.Metric(metric_id=metric_id, label=label, value=value,
                         display_value=display, unit=unit)]))
    row.section_key = section_key
    row.source_locator = "xlsx://Coverage!B12"
    db.flush()
    return row


# ======================================================== history


@pytest.mark.usefixtures("db")
class TestHistoryIsAssembledNotStored:

    def test_a_version_appears_with_what_it_changed(self, db, workspace):
        artifact, _ = _versioned(db, workspace, _doc("1. Summary"))
        db.flush()
        [event] = [e for e in history.events(db, workspace.id)
                   if e.kind == history.VERSION]
        assert event.title == "Version 1 created"
        assert event.ids["version"] == 1

    def test_a_seeded_version_is_not_attributed_to_anybody(self, db,
                                                           workspace):
        """§13: fixture text is never presented as something a model wrote."""
        artifact = repo.create_artifact(db, workspace.id, kind="report",
                                        title="R")
        db.flush()
        doc = _doc("1. Summary")
        repo.new_version(db, artifact, content=doc.as_dict(),
                         source_manifest={}, content_hash=doc.content_hash(),
                         validation={}, origin="seed_fixture")
        db.flush()
        [event] = [e for e in history.events(db, workspace.id)
                   if e.kind == history.VERSION]
        assert event.actor == ""
        assert event.ids["origin"] == "seed_fixture"

    def test_a_governed_act_carries_the_person_who_performed_it(
            self, db, workspace):
        finding = gov.raise_finding(db, workspace.id, origin=gov.FROM_RULE,
                                    title="Coverage fell")
        gov.move_finding(db, finding, to=gov.ACCEPTED, actor=ACTOR,
                         reason="agreed")
        db.flush()
        events = [e for e in history.events(db, workspace.id)
                  if e.kind == history.FINDING]
        accepted = [e for e in events if "accepted" in e.title]
        assert accepted and accepted[0].actor == ACTOR
        assert "agreed" in accepted[0].detail

    def test_a_system_act_shows_no_name_rather_than_a_fake_one(
            self, db, workspace):
        gov.raise_finding(db, workspace.id, origin=gov.FROM_RULE,
                          title="Coverage fell")
        db.flush()
        [raised] = [e for e in history.events(db, workspace.id)
                    if e.kind == history.FINDING]
        assert raised.actor == ""

    def test_a_confirmed_metric_link_is_an_event(self, db, workspace):
        row = _governed(db, workspace)
        bind.confirm(db, row, actor=ACTOR)
        db.flush()
        [event] = [e for e in history.events(db, workspace.id)
                   if e.kind == history.BINDING]
        assert event.actor == ACTOR
        assert "ifrs9.coverage_ratio" in event.detail

    def test_an_unconfirmed_link_is_not_an_event(self, db, workspace):
        """Nothing happened yet. A suggestion sitting in a queue is state, not
        history."""
        bind.apply(db, workspace.id, bind.from_labels(
            [("Gini", "0.415", "0.4152", "B2")], locator="xlsx://P!A1"))
        db.flush()
        assert not [e for e in history.events(db, workspace.id)
                    if e.kind == history.BINDING]

    def test_uploads_and_re_reads_both_appear(self, db, scope, workspace,
                                              results_workbook_xlsx):
        from backend.playbook import reparse

        source = service.add_source(db, scope, workspace.id,
                                    filename="results.xlsx",
                                    content=results_workbook_xlsx,
                                    source_role="results")
        reparse.latest(db, source.id).parser_version = "0"
        db.flush()
        reparse.reread(db, scope, source.id)
        db.flush()

        events = history.events(db, workspace.id)
        assert [e for e in events if e.kind == history.SOURCE]
        parses = [e for e in events if e.kind == history.PARSE]
        assert len(parses) == 2
        assert any("again (revision 2)" in e.title for e in parses)

    def test_a_deleted_finding_takes_its_entries_with_it(self, db, workspace):
        """The property that makes assembling better than storing."""
        from backend.models.playbook import PlaybookFinding

        finding = gov.raise_finding(db, workspace.id, origin=gov.FROM_RULE,
                                    title="Coverage fell")
        db.flush()
        assert [e for e in history.events(db, workspace.id)
                if e.kind == history.FINDING]

        db.delete(db.get(PlaybookFinding, finding.id))
        db.flush()
        assert not [e for e in history.events(db, workspace.id)
                    if e.kind == history.FINDING]

    def test_newest_first_and_undated_entries_are_kept_last(self, db,
                                                            workspace):
        finding = gov.raise_finding(db, workspace.id, origin=gov.FROM_RULE,
                                    title="Coverage fell")
        finding.history = list(finding.history) + [
            {"at": "", "act": "noted", "from": "", "to": "", "actor": "",
             "reason": "no timestamp"}]
        db.flush()
        stamps = [e.at for e in history.events(db, workspace.id)]
        assert stamps == sorted(stamps, key=lambda s: (s != "", s),
                                reverse=True)
        assert stamps[-1] == ""

    def test_the_feed_counts_every_kind_it_knows(self, db, workspace):
        gov.raise_finding(db, workspace.id, origin=gov.FROM_RULE, title="X")
        db.flush()
        feed = history.feed(db, workspace.id)
        assert {k["kind"] for k in feed["kinds"]} == set(history.KINDS)
        assert feed["total"] >= 1

    def test_the_feed_filters_without_losing_the_totals(self, db, workspace):
        _versioned(db, workspace, _doc("1. Summary"))
        gov.raise_finding(db, workspace.id, origin=gov.FROM_RULE, title="X")
        db.flush()
        feed = history.feed(db, workspace.id, kinds=[history.FINDING])
        assert all(e["kind"] == history.FINDING for e in feed["events"])
        assert feed["total"] > feed["shown"]


# ======================================================== check for updates


@pytest.mark.usefixtures("db")
class TestCheckingForUpdatesChangesNothing:

    def test_a_quiet_document_says_so(self, db, workspace):
        _versioned(db, workspace, _doc("1. Summary"))
        proposal = refresh.check(db, workspace.id)
        assert proposal.anything is False
        assert proposal.as_dict()["message"].startswith("Nothing has changed")

    def test_a_newer_reading_is_proposed_with_its_section(self, db,
                                                          workspace):
        artifact, _ = _versioned(db, workspace, _doc("1. Summary"))
        key = sect.key_for("1. Summary", 1)
        row = _governed(db, workspace, section_key=key)
        row.freshness = bind.NEW_AVAILABLE
        db.flush()

        proposal = refresh.check(db, workspace.id)
        assert [m["label"] for m in proposal.metrics_with_newer_values] \
            == ["Coverage ratio"]
        assert [s["heading"] for s in proposal.sections_affected] \
            == ["1. Summary"]
        assert "1 metric has a newer value" in proposal.lines()

    def test_an_unconfirmed_suggestion_is_not_a_reason_to_refresh(
            self, db, workspace):
        """It is a different sentence with a different remedy."""
        _versioned(db, workspace, _doc("1. Summary"))
        row, = bind.apply(db, workspace.id, bind.from_labels(
            [("Gini", "0.415", "0.4152", "B2")], locator="xlsx://P!A1"))
        row.freshness = bind.NEW_AVAILABLE
        db.flush()

        proposal = refresh.check(db, workspace.id)
        assert proposal.metrics_with_newer_values == []
        assert proposal.sections_affected == []
        assert len(proposal.suggestions_awaiting_review) == 1

    def test_a_finding_raised_against_a_moved_figure_is_flagged(
            self, db, workspace):
        _versioned(db, workspace, _doc("1. Summary"))
        row = _governed(db, workspace)
        row.freshness = bind.NEW_AVAILABLE
        db.flush()
        gov.raise_finding(db, workspace.id, origin=gov.FROM_CHANGE,
                          title="Coverage fell", severity=gov.HIGH,
                          metric_id=row.metric_id, current_value="5.86%")
        db.flush()

        proposal = refresh.check(db, workspace.id)
        assert [f["title"] for f in proposal.findings_to_reconsider] \
            == ["Coverage fell"]
        assert proposal.findings_to_reconsider[0]["recorded_value"] == "5.86%"

    def test_a_decided_decision_is_never_reopened_by_new_data(self, db,
                                                              workspace):
        """History is not revised because the world moved on."""
        _versioned(db, workspace, _doc("1. Summary"))
        row = _governed(db, workspace)
        row.freshness = bind.NEW_AVAILABLE
        db.flush()
        finding = gov.raise_finding(db, workspace.id, origin=gov.FROM_CHANGE,
                                    title="Coverage fell",
                                    metric_id=row.metric_id)
        decision = gov.propose_decision(db, workspace.id, question="Hold?",
                                        related_finding_ids=[finding.id])
        gov.move_decision(db, decision, to=gov.READY_FOR_DECISION,
                          actor=ACTOR)
        gov.record(db, decision, outcome=gov.APPROVE, actor=ACTOR)
        db.flush()

        proposal = refresh.check(db, workspace.id)
        assert proposal.decisions_to_revisit == []

    def test_an_outstanding_decision_on_moved_evidence_is_flagged(
            self, db, workspace):
        _versioned(db, workspace, _doc("1. Summary"))
        row = _governed(db, workspace)
        row.freshness = bind.NEW_AVAILABLE
        db.flush()
        finding = gov.raise_finding(db, workspace.id, origin=gov.FROM_CHANGE,
                                    title="Coverage fell",
                                    metric_id=row.metric_id)
        gov.propose_decision(db, workspace.id, question="Hold the overlay?",
                             related_finding_ids=[finding.id])
        db.flush()

        proposal = refresh.check(db, workspace.id)
        assert [d["question"] for d in proposal.decisions_to_revisit] \
            == ["Hold the overlay?"]

    def test_checking_writes_nothing(self, db, scope, workspace,
                                     results_workbook_xlsx):
        """The whole contract. A check that edited would be the silent rewrite
        §16 forbids."""
        from backend.playbook import reparse

        artifact, version = _versioned(db, workspace, _doc("1. Summary"))
        source = service.add_source(db, scope, workspace.id,
                                    filename="results.xlsx",
                                    content=results_workbook_xlsx,
                                    source_role="results")
        row = _governed(db, workspace,
                        section_key=sect.key_for("1. Summary", 1))
        row.freshness = bind.NEW_AVAILABLE
        reparse.latest(db, source.id).parser_version = "0"
        db.flush()

        before = (version.content_hash,
                  [r.status for r in sect.rows_for(db, artifact.id)],
                  row.display_value, row.freshness,
                  reparse.latest(db, source.id).revision)

        refresh.check(db, workspace.id)
        refresh.check(db, workspace.id)

        after = (version.content_hash,
                 [r.status for r in sect.rows_for(db, artifact.id)],
                 row.display_value, row.freshness,
                 reparse.latest(db, source.id).revision)
        assert before == after
        assert len(repo.versions(db, artifact.id)) == 1

    def test_a_stale_source_appears_in_the_proposal(self, db, scope,
                                                    workspace,
                                                    results_workbook_xlsx):
        from backend.playbook import reparse

        _versioned(db, workspace, _doc("1. Summary"))
        source = service.add_source(db, scope, workspace.id,
                                    filename="results.xlsx",
                                    content=results_workbook_xlsx,
                                    source_role="results")
        reparse.latest(db, source.id).parser_version = "0"
        db.flush()

        proposal = refresh.check(db, workspace.id)
        assert [s["filename"] for s in proposal.sources_needing_reread] \
            == ["results.xlsx"]
        assert "1 source needs re-reading" in proposal.lines()


# ======================================================== a new upload


@pytest.mark.usefixtures("db")
class TestAnUploadProposesRatherThanUpdates:

    @staticmethod
    def _monitoring_workbook() -> bytes:
        from openpyxl import Workbook

        wb = Workbook()
        sheet = wb.active
        sheet.title = "Performance"
        sheet.append(["Metric", "Value"])
        sheet.append(["Coverage ratio", 0.0647])
        sheet.append(["Gini coefficient", 0.4770])
        sheet.append(["Something nobody tracks", 12])
        for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
            for cell in row:
                cell.number_format = "0.000"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_it_names_the_tracked_metrics_the_file_would_update(
            self, db, scope, workspace):
        _governed(db, workspace, metric_id="ifrs9.coverage_ratio",
                  label="Coverage ratio", value="0.0586", display="5.86%")
        source = service.add_source(db, scope, workspace.id,
                                    filename="monitoring.xlsx",
                                    content=self._monitoring_workbook(),
                                    source_role="results")
        db.flush()

        result = refresh.proposals_from_source(db, workspace.id, source.id)
        assert result["detected"] == 1
        [row] = result["rows"]
        assert row["document_metric"] == "Coverage ratio"
        assert row["current_document_value"] == "5.86%"
        assert row["uploaded_value"] == "0.065"
        assert row["changes_value"] is True
        assert "appear" in result["message"]

    def test_every_row_is_a_suggestion_and_says_so(self, db, scope,
                                                   workspace):
        """A column header resembling a tracked metric is not evidence that it
        is that metric."""
        _governed(db, workspace, metric_id="ifrs9.coverage_ratio",
                  label="Coverage ratio", value="0.0586", display="5.86%")
        source = service.add_source(db, scope, workspace.id,
                                    filename="monitoring.xlsx",
                                    content=self._monitoring_workbook(),
                                    source_role="results")
        db.flush()

        result = refresh.proposals_from_source(db, workspace.id, source.id)
        assert all(r["match_status"] == bind.SUGGESTED for r in result["rows"])

    def test_it_proposes_nothing_for_a_metric_nobody_tracks(self, db, scope,
                                                            workspace):
        source = service.add_source(db, scope, workspace.id,
                                    filename="monitoring.xlsx",
                                    content=self._monitoring_workbook(),
                                    source_role="results")
        db.flush()
        result = refresh.proposals_from_source(db, workspace.id, source.id)
        assert result["detected"] == 0
        assert "Nothing in this file" in result["message"]

    def test_proposing_changes_no_governed_value(self, db, scope, workspace):
        binding = _governed(db, workspace, metric_id="ifrs9.coverage_ratio",
                            label="Coverage ratio", value="0.0586",
                            display="5.86%")
        source = service.add_source(db, scope, workspace.id,
                                    filename="monitoring.xlsx",
                                    content=self._monitoring_workbook(),
                                    source_role="results")
        db.flush()
        refresh.proposals_from_source(db, workspace.id, source.id)
        assert binding.display_value == "5.86%"
        assert binding.value_in_document == "0.0586"

    def test_a_source_from_another_workspace_is_refused(self, db, scope,
                                                        workspace):
        other = repo.create_workspace(db, scope, title="Another")
        db.flush()
        source = service.add_source(db, scope, other.id,
                                    filename="monitoring.xlsx",
                                    content=self._monitoring_workbook(),
                                    source_role="results")
        db.flush()
        with pytest.raises(repo.NotFound):
            refresh.proposals_from_source(db, workspace.id, source.id)

    def test_an_unlinked_suggestion_never_becomes_a_governed_value(
            self, db, scope, workspace):
        """Rule 3, at the end of the journey: the uploaded value is visible and
        is not current until a person confirms it."""
        binding = _governed(db, workspace, metric_id="ifrs9.coverage_ratio",
                            label="Coverage ratio", value="0.0586",
                            display="5.86%")
        source = service.add_source(db, scope, workspace.id,
                                    filename="monitoring.xlsx",
                                    content=self._monitoring_workbook(),
                                    source_role="results")
        db.flush()
        refresh.proposals_from_source(db, workspace.id, source.id)

        [live] = (db.query(PlaybookMetricBinding)
                  .filter(PlaybookMetricBinding.id == binding.id).all())
        assert live.display_value == "5.86%"


@pytest.mark.usefixtures("db")
class TestApplyingAnUploadedReading:
    """§24's last step. Only confirmed mappings become governed values."""

    @staticmethod
    def _workbook() -> bytes:
        from openpyxl import Workbook

        wb = Workbook()
        sheet = wb.active
        sheet.title = "Performance"
        sheet.append(["Metric", "Value"])
        sheet.append(["Coverage ratio", 0.0647])
        for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
            for cell in row:
                cell.number_format = "0.000"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _ready(self, db, scope, workspace):
        binding = _governed(db, workspace, metric_id="ifrs9.coverage_ratio",
                            label="Coverage ratio", value="0.0586",
                            display="5.86%")
        source = service.add_source(db, scope, workspace.id,
                                    filename="monitoring.xlsx",
                                    content=self._workbook(),
                                    source_role="results")
        db.flush()
        return binding, source

    def test_a_confirmed_reading_becomes_the_current_value(self, db, scope,
                                                           workspace):
        binding, source = self._ready(db, scope, workspace)
        result = refresh.apply_uploaded(db, workspace.id, source.id,
                                        binding_ids=[binding.id],
                                        actor=ACTOR)
        assert result["updated"][0]["value"] == "0.065"
        assert binding.display_value == "0.065"
        assert binding.raw_value == "0.0647"
        assert binding.confirmed_by == ACTOR
        assert binding.source_locator.endswith("!B2")
        assert binding.freshness == bind.CURRENT

    def test_a_reading_nobody_confirmed_is_left_alone(self, db, scope,
                                                      workspace):
        binding, source = self._ready(db, scope, workspace)
        result = refresh.apply_uploaded(db, workspace.id, source.id,
                                        binding_ids=[], actor=ACTOR)
        assert result["updated"] == []
        assert binding.display_value == "5.86%"
        assert result["message"] == "Nothing was changed."

    def test_a_machine_may_not_apply_one(self, db, scope, workspace):
        binding, source = self._ready(db, scope, workspace)
        with pytest.raises(gov.NotPermitted):
            refresh.apply_uploaded(db, workspace.id, source.id,
                                   binding_ids=[binding.id], actor="claude")
        assert binding.display_value == "5.86%"

    def test_the_frozen_snapshot_is_never_rewritten(self, db, scope,
                                                    workspace):
        """THEN is what the document relied on. If applying a new reading
        moved it, the next comparison would show no change at all."""
        from backend.models.playbook import PlaybookMetricSnapshot

        binding, source = self._ready(db, scope, workspace)
        artifact, version = _versioned(db, workspace, _doc("1. Summary"))
        db.flush()
        [snapshot] = (db.query(PlaybookMetricSnapshot)
                      .filter(PlaybookMetricSnapshot.version_id == version.id)
                      .all())
        before = snapshot.display_value

        refresh.apply_uploaded(db, workspace.id, source.id,
                               binding_ids=[binding.id], actor=ACTOR)
        assert snapshot.display_value == before == "5.86%"

    def test_an_id_that_was_never_proposed_is_ignored(self, db, scope,
                                                      workspace):
        """A client cannot update an arbitrary binding by naming it here."""
        binding, source = self._ready(db, scope, workspace)
        other = _governed(db, workspace, metric_id="ifrs9.weighted_ecl",
                          label="Weighted ECL", value="22.77",
                          display="22.77")
        refresh.apply_uploaded(db, workspace.id, source.id,
                               binding_ids=[other.id], actor=ACTOR)
        assert other.display_value == "22.77"
