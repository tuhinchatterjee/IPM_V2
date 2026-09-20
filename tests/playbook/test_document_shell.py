"""
The shell a report is delivered in. Chapter 09.

A real Auto Loan report reached a user looking like this:

* the H1, the running header and the file name were the user's PROMPT —
  "Create a detailed Auto Loan Application Scorecard Model Development Report
  using the attached evidence. Treat AL-AS-v1." — because `author_document`
  was called with `title=workspace.title` and a workspace is named after its
  first message;
* that prompt contained a newline, which went into the PDF's running header
  where the font has no glyph for it and drew as two black boxes;
* `**Model ID:**`, `> STANDING CAVEAT` and a bare `---` printed verbatim,
  because a `Block` carries one flat string and every writer emitted it as a
  single plain run;
* and there was no cover, no contents and no way to navigate sixteen pages.

Chapter 09 requires "real headings, tables, sensible page layout, source notes
and usable navigation" in Word, and "correct pagination, no clipped tables and
no broken characters" in PDF. Both rows were failed, and nothing in the
repository looked: DC-15 was marked PASS on "38 KB Word, 2.8 KB PDF, both
downloaded", which proves a file exists and opens.
"""

from __future__ import annotations

import io
import re

import pytest

from backend.playbook import document as D
from backend.playbook import service
from backend.playbook.render import docx_writer, inline, pdf_writer, shell
from backend.playbook.repository import versions as repo_versions

REPORT = """# Auto Loan Application Scorecard — Model Development Report

## 1. Executive summary

**Model ID:** AL-AS-v1.0 **Owner:** Retail Credit Risk

> A standing caveat about *synthetic* evidence.

## 2. Scope and evidence

| Item | Definition |
| --- | --- |
| Portfolio | Retail auto loans |
| Geography | Saudi Arabia |

## 3. Target definition

The field `Default_12m` records 90+ DPD within twelve months.

## 4. Calibration

Observed default rates are not monotonic across grades.
"""

#: What the delivered draft actually looked like: the report names itself and
#: then says something before its first numbered section. That matters,
#: because `_title_is_not_a_section` recovers a mis-parsed H1 only when the
#: section it became is EMPTY — a title line followed immediately by the next
#: heading. Put one sentence under the title and that safety net does not
#: apply, which is precisely the shape that reached the user.
DRAFT = REPORT.replace(
    "## 1. Executive summary",
    "Prepared for the Model Risk Committee.\n\n## 1. Executive summary", 1)


#: The prompt that became a title, newline and all.
PROMPT = ("Create a detailed Auto Loan Application Scorecard Model "
          "Development Report using the attached evidence.\nTreat AL-AS-v1.")


@pytest.fixture
def report() -> D.Document:
    return D.parse(REPORT)


@pytest.fixture
def ledger():
    """Empty on purpose. These tests are about the shell a report is
    delivered in, not about what its figures are traceable to."""
    from backend.playbook import evidence as ev

    return ev.Ledger()


def _docx_text(raw: bytes) -> str:
    from docx import Document

    out = Document(io.BytesIO(raw))
    parts = [p.text for p in out.paragraphs]
    for table in out.tables:
        parts += [cell.text for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def _pdf_bold_runs(raw: bytes) -> int:
    """How many times the file switches into a bold font.

    Presence of a bold font proves nothing: `PbTitle`, `PbH1` and `PbH2`
    inherit reportlab's own heading styles, which are Helvetica-Bold, so every
    report that has a heading has the font whether or not a single word of its
    prose is emphasised. Each `/Fn size Tf` in a page's content stream is one
    switch, and `**bold**` inside a paragraph adds switches that deleting the
    markers does not — which is the whole difference between applying inline
    formatting and throwing it away.

    The reader is held in a local: pypdf reads stream bytes lazily, and a
    reader built inline is collected out from under `get_data`, which then
    returns nothing at all.
    """
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    switches = 0
    for page in reader.pages:
        fonts = {str(name): str(ref.get_object().get("/BaseFont") or "")
                 for name, ref in
                 ((page.get("/Resources") or {}).get("/Font") or {}).items()}
        data = page.get_contents().get_data()
        for name in re.findall(rb"(/\w+)\s+[\d.]+\s+Tf", data):
            if "Bold" in fonts.get(name.decode(), ""):
                switches += 1
    return switches


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    return "\n".join(p.extract_text() or ""
                     for p in PdfReader(io.BytesIO(raw)).pages)


class TestTheTitleIsTheReportsOwn:

    def test_the_model_names_its_own_report(self, report):
        assert report.title == (
            "Auto Loan Application Scorecard — Model Development Report")

    def test_the_prompt_is_only_a_fallback(self, report):
        assert shell.title_of(report, fallback=PROMPT) == report.title

    def test_a_fallback_is_still_one_line(self):
        """The newline is what drew as black boxes in the PDF header."""
        cleaned = shell.clean_title(PROMPT)
        assert "\n" not in cleaned
        assert len(cleaned) <= shell.MAX_TITLE

    def test_marks_never_reach_a_title(self):
        assert shell.clean_title("**Auto Loan** Scorecard") == \
            "Auto Loan Scorecard"

    def test_an_empty_title_is_named_rather_than_blank(self):
        assert shell.clean_title("") == "CreditProbe report"
        assert shell.clean_title("   \n  ") == "CreditProbe report"

    def test_a_good_title_is_left_exactly_alone(self):
        for title in ("IFRS 9 Committee Report — Q2 2026",
                      "Behavioural Scorecard Validation Report"):
            assert shell.clean_title(title) == title


class TestNoMarkdownReachesTheReader:

    @pytest.mark.parametrize("marker", ["**", "*", "`", "> "])
    def test_the_source_really_contains_it(self, marker):
        """So a passing test below means the writer removed it, not that the
        fixture never had it."""
        assert marker in REPORT

    def test_word_applies_bold_rather_than_printing_asterisks(self, report):
        from docx import Document

        out = Document(io.BytesIO(docx_writer.write(report)))
        text = _docx_text(docx_writer.write(report))
        assert "**" not in text and "__" not in text
        bold = [r.text for p in out.paragraphs for r in p.runs if r.bold]
        assert "Model ID:" in bold, "the marks were stripped, not applied"

    def test_the_pdf_applies_them_too(self, report):
        """Asserted by comparison, not by presence.

        `_escape` strips the marks before escaping, so "no asterisks in the
        extracted text" is equally true of a writer that simply deleted them,
        and a bold font is already in the file because the headings use one.
        What only an applying writer can produce is MORE bold than the same
        report with the marks taken out of the source — same headings, same
        cover, same table furniture, one difference.
        """
        raw = pdf_writer.write(report)
        text = _pdf_text(raw)
        assert "**" not in text
        assert "Model ID:" in text

        twin = D.parse(inline.plain(REPORT))
        assert _pdf_bold_runs(raw) > _pdf_bold_runs(pdf_writer.write(twin)), (
            "the marks were stripped, not applied: the emphasised report uses "
            "no more bold than the same report written without emphasis")

    def test_the_bold_is_the_paragraph_and_not_the_furniture(self):
        """The narrow version, on a document that has no furniture at all.

        No cover table, no front matter, too few sections for a contents page.
        The only bold either of these can carry beyond its two headings is the
        one inside the paragraph.
        """
        body = "**Model ID:** AL-AS-v1.0 and ordinary words.\n"
        head = "# A note\n\n## Only section\n\n"
        marked = _pdf_bold_runs(pdf_writer.write(D.parse(head + body)))
        flat = _pdf_bold_runs(
            pdf_writer.write(D.parse(head + inline.plain(body))))
        assert marked == flat + 1, (
            f"expected exactly one more bold run in the paragraph; "
            f"emphasised {marked}, plain {flat}")

    def test_a_callout_keeps_its_words_and_loses_its_marker(self, report):
        """`>` is the CALLOUT kind this model already had; `parse` simply
        never consumed it, so the marker printed."""
        callouts = [b for s in report.sections for b in s.blocks
                    if b.kind == D.CALLOUT]
        assert callouts, "the '>' was never consumed as a callout at all"
        block = callouts[0]
        assert not block.text.startswith(">")
        for text in (_docx_text(docx_writer.write(report)),
                     _pdf_text(pdf_writer.write(report))):
            assert "standing caveat" in text.lower()
            assert "> " not in text

    def test_a_horizontal_rule_is_a_separator_not_content(self):
        """`---` used to fall through to the paragraph buffer and print."""
        doc = D.parse("# R\n\n## S\n\nAbove.\n\n---\n\nBelow.\n")
        kinds = [b.text for s in doc.sections for b in s.blocks]
        assert kinds == ["Above.", "Below."]
        assert "---" not in _pdf_text(pdf_writer.write(doc))

    def test_a_table_cell_holds_no_marks(self):
        doc = D.parse("# R\n\n## S\n\n| A | B |\n| --- | --- |\n"
                      "| **bold** | `code` |\n")
        text = _docx_text(docx_writer.write(doc))
        assert "**" not in text and "`" not in text
        assert "bold" in text and "code" in text


class TestTheTitleSurvivesTheRealPath:
    """`title_of` alone cannot prove this.

    The defect was upstream of it: `author_document` parsed with
    `title=workspace.title`, so `parse` never took the model's own `# H1` as
    the title — it became the first SECTION, and the prompt became the title.
    A unit test that parses without a title cannot see that, which is why
    this one goes through the service.
    """

    def test_the_artifact_is_named_after_the_report(
            self, db, scope, workspace, ledger, scripted_author):
        from backend.playbook import repository as repo

        scripted_author(DRAFT)
        service.author_document(
            db, scope, workspace.id,
            instruction="Write it.", ledger=ledger,
            # What a real caller passes: the workspace's name, which is its
            # first message.
            title=PROMPT, formats=["docx"])

        artifact = repo.artifacts(db, workspace.id)[0]
        assert artifact.title == (
            "Auto Loan Application Scorecard — Model Development Report")
        assert "Create a detailed" not in artifact.title

    def test_the_heading_is_not_demoted_into_a_section(
            self, db, scope, workspace, ledger, scripted_author):
        """The other half of the same bug: with a title supplied, the model's
        H1 became section one and the report carried two titles — the prompt
        at the top and the real one below it."""
        scripted_author(DRAFT)
        outcome = service.author_document(
            db, scope, workspace.id, instruction="Write it.",
            ledger=ledger, title=PROMPT, formats=["docx"])

        version = repo_versions(db, outcome.artifact_id)[-1]
        stored = D.Document.from_dict(version.content or {})
        assert stored.title.startswith("Auto Loan Application Scorecard")
        assert "Create a detailed" not in stored.title
        assert all("Auto Loan Application Scorecard" not in s.heading
                   for s in stored.sections), (
            "the report's own title is also one of its sections")


class TestThereIsSomethingToNavigateBy:

    def test_word_carries_a_real_contents_field(self, report):
        """A field, not a typed list: Word renumbers it when the document
        reflows, which is what "usable navigation" means."""
        from docx import Document

        out = Document(io.BytesIO(docx_writer.write(report)))
        xml = out.element.xml
        assert 'TOC \\o "1-3"' in xml
        assert "Contents" in _docx_text(docx_writer.write(report))

    def test_the_cover_carries_no_running_header_or_page_number(self, report):
        """A cover is a cover.

        The running header repeated the title six centimetres above the
        title, and "Page 1" sat under a page that is not part of the reading.
        """
        from pypdf import PdfReader

        pages = PdfReader(io.BytesIO(pdf_writer.write(report))).pages
        first = pages[0].extract_text() or ""
        assert "Page 1" not in first
        # Once, as the title — not again as a header above it. The title
        # wraps in the extracted text, so it is counted by its opening.
        assert first.count("Auto Loan Application Scorecard") == 1

        # And the rest of the document is numbered as usual.
        assert "Page 2" in (pages[1].extract_text() or "")

    def test_a_short_document_still_gets_its_furniture(self):
        """No cover, no contents — and therefore the header and the page
        number from the first page, because there is no cover to keep them
        off."""
        from pypdf import PdfReader

        short = D.parse("# A note\n\n## Only section\n\nOne paragraph.\n")
        page = PdfReader(io.BytesIO(pdf_writer.write(short))).pages[0]
        assert "Page 1" in (page.extract_text() or "")

    def test_the_pdf_contents_lists_every_section_with_its_page(self, report):
        text = _pdf_text(pdf_writer.write(report))
        assert "Contents" in text
        for heading in ("1. Executive summary", "2. Scope and evidence",
                        "3. Target definition", "4. Calibration"):
            assert heading in text

    def test_the_contents_does_not_list_itself(self, report):
        text = _pdf_text(pdf_writer.write(report))
        assert text.count("Contents") == 1

    def test_a_short_note_gets_no_contents_page(self):
        """Below three sections a contents page is furniture, not navigation."""
        doc = D.parse("# A note\n\n## Only section\n\nBody.\n")
        assert shell.shell_of(doc).has_contents is False
        assert "Contents" not in _pdf_text(pdf_writer.write(doc))

    def test_the_front_matter_is_a_record_not_a_caption(self):
        doc = D.parse(REPORT)
        doc.meta = {"model_id": "AL-AS-v1.0", "owner": "Retail Credit Risk",
                    "status": "Draft"}
        front = shell.shell_of(doc)
        assert ("Model Id", "AL-AS-v1.0") in front.facts
        text = _docx_text(docx_writer.write(doc))
        assert "Retail Credit Risk" in text
        # The middot-joined line it used to be.
        assert "·" not in text


class TestTheContentsPageDoesNotBecomeAClaim:
    """Building navigation puts a column of page numbers into the PDF's text
    layer, and the evidence has no reason to support them."""

    def test_the_rendered_pdf_still_validates(self, report):
        from backend.playbook import validate

        result = validate.validate(pdf_writer.write(report), "pdf", report)
        assert result.ok, result.issues

    def test_a_bare_figure_that_is_not_a_page_reference_is_still_a_claim(self):
        from backend.playbook import validate

        doc = D.parse("# R\n\n## 1. Scope\n\nBody.\n\n## 2. Findings\n\n"
                      "Body.\n\n## 3. Conclusion\n\nBody.\n")
        cleaned, removed = validate.structural_numerals(
            "1. Scope\n3\n2. Findings\n4\n41.9\nPage 2", doc)
        assert removed == 5
        assert "41.9" in cleaned, (
            "a figure standing alone in a cell is not a page reference")


class TestInlineFormatting:
    """Kept apart from the writers, because it is the piece both depend on."""

    @pytest.mark.parametrize("source,plain", [
        ("**Model ID:** AL-AS-v1.0", "Model ID: AL-AS-v1.0"),
        ("An *italic* word", "An italic word"),
        ("A `code` span", "A code span"),
        ("**bold with *italic* inside**", "bold with italic inside"),
        ("__also bold__ and _also italic_", "also bold and also italic"),
    ])
    def test_marks_are_understood(self, source, plain):
        assert inline.plain(source) == plain

    @pytest.mark.parametrize("source", [
        "snake_case_name stays whole",
        "a * b * c is arithmetic",
        "An unmatched ** stays literal",
        "2026_Q1_results.xlsx",
    ])
    def test_what_is_not_emphasis_is_left_alone(self, source):
        assert inline.plain(source) == source

    def test_markup_cannot_be_injected_through_the_document(self):
        """A source document containing a tag must render as characters."""
        out = inline.markup("Compare <b>2 < 3</b> & 4")
        assert "&lt;b&gt;" in out and "&amp;" in out

    def test_a_tag_this_module_added_is_not_escaped_by_it(self):
        assert inline.markup("**bold**") == "<b>bold</b>"
