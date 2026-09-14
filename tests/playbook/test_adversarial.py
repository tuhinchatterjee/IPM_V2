"""Gate 11 — actively trying to break it. §21.

The soak harness proves the journeys land in the same place ten times running.
This is the other half: the cases chosen because they are the ones most likely
to produce a confident wrong answer.

The failure this file exists to prevent is not a crash. A crash is visible and
somebody fixes it. The failure that reaches a committee is a number that looks
right — two metrics with the same label bound because the words matched, a
percentage-point move printed as a percentage, SAR against SAR million, a
figure grounded because the same digits happened to appear in an unrelated
source. Every case below is either refused with a named reason, or handled
correctly; none is allowed to be silently approximated.
"""

from __future__ import annotations

import io

import pytest

from backend.exports import playbook_contract as contract
from backend.models.playbook import (
    PlaybookMetricBinding,
    PlaybookMetricSnapshot,
)
from backend.playbook import calc, grounding, merge, service, validate
from backend.playbook import document as D
from backend.playbook import evidence as ev
from backend.playbook import repository as repo
from backend.playbook.intelligence import adopt, compare
from backend.playbook.intelligence import binding as bind
from backend.playbook.intelligence import governance as gov
from backend.playbook.intelligence import readiness as score
from backend.playbook.intelligence import sections as sect

ACTOR = "user:7"
BODY = "considered judgement " * 12


def _artifact(db, workspace, doc, *, version=1, artifact=None):
    if artifact is None:
        artifact = repo.create_artifact(db, workspace.id, kind="report",
                                        title="Report")
        db.flush()
    row = repo.new_version(db, artifact, content=doc.as_dict(),
                           source_manifest={},
                           content_hash=doc.content_hash(), validation={})
    db.flush()
    adopt.adopt(db, workspace.id, artifact.id, version_id=row.id,
                version=version, doc=doc)
    return artifact, row


def _doc(*headings: str) -> D.Document:
    return D.parse("".join(f"## {h}\n\n{BODY}\n\n" for h in headings),
                   title="Report")


def _bound(db, workspace, **over):
    fields = dict(workspace_id=workspace.id, metric_id="risk.default_rate",
                  label="Default rate", value_in_document="6.47",
                  raw_value="6.47", display_value="6.47%", unit="percent",
                  binding_method=bind.FROM_EXPORT, confirmed_by_user=False)
    fields.update(over)
    row = PlaybookMetricBinding(**fields)
    db.add(row)
    db.flush()
    return row


def _snapshot(db, artifact, version_row, binding, **over):
    fields = dict(artifact_id=artifact.id, version_id=version_row.id,
                  version=1, metric_id=binding.metric_id,
                  label=binding.label, value=binding.value_in_document,
                  display_value=binding.display_value, unit=binding.unit,
                  section_key="", currency="", population="", segment="",
                  scenario="", reporting_period="")
    fields.update(over)
    row = PlaybookMetricSnapshot(**fields)
    db.add(row)
    db.flush()
    return row


# ================================================== identity and context


@pytest.mark.usefixtures("db")
class TestTwoThingsThatLookLikeOneThing:

    def test_the_same_label_in_two_populations_is_two_metrics(
            self, db, workspace):
        """The mistake that makes a governed pack wrong: binding by words."""
        retail = _bound(db, workspace, metric_id="risk.default_rate.retail",
                        population="Retail", display_value="6.47%")
        corporate = _bound(db, workspace,
                           metric_id="risk.default_rate.corporate",
                           population="Corporate", display_value="2.11%")
        assert retail.metric_id != corporate.metric_id
        assert bind.coverage([retail, corporate])["detected"] == 2

    def test_a_snapshot_will_not_compare_across_populations(
            self, db, workspace):
        artifact, version = _artifact(db, workspace, _doc("Summary"))
        binding = _bound(db, workspace, population="Corporate",
                         confirmed_by_user=True)
        snapshot = _snapshot(db, artifact, version, binding,
                             population="Retail")

        result = compare.compare_one(snapshot, binding)
        assert result.comparable is False
        assert "population differs" in result.reason
        assert "retail" in result.reason and "corporate" in result.reason

    def test_two_reporting_periods_are_not_one_series(self, db, workspace):
        artifact, version = _artifact(db, workspace, _doc("Summary"))
        binding = _bound(db, workspace, reporting_period="Q3 2026",
                         confirmed_by_user=True)
        snapshot = _snapshot(db, artifact, version, binding,
                             reporting_period="Q2 2026")
        # The period is identity, not decoration, but it is not one of the
        # dimensions that BLOCKS a comparison: a quarter-on-quarter move is
        # the comparison. What must not happen is it being lost.
        result = compare.compare_one(snapshot, binding)
        assert result.then_period == "Q2 2026"
        assert result.lineage["now"]["period"] == "Q3 2026"

    def test_a_currency_change_stops_the_comparison(self, db, workspace):
        artifact, version = _artifact(db, workspace, _doc("Summary"))
        binding = _bound(db, workspace, unit="currency", currency="USD",
                         value_in_document="22.77", display_value="22.77",
                         confirmed_by_user=True)
        snapshot = _snapshot(db, artifact, version, binding, currency="SAR",
                             unit="currency")
        result = compare.compare_one(snapshot, binding)
        assert result.comparable is False
        assert "currency differs" in result.reason

    def test_the_same_number_in_two_sources_is_not_two_facts(self):
        """Grounding is a set membership test by design. What must hold is
        that a number cannot be laundered by APPEARING somewhere — the class
        it appears as has to be evidence-bearing."""
        ledger = ev.Ledger()
        ledger.add(ev.Item("docx://para/2", "paragraph",
                           "Model LGD-2026-v3 was approved on 30 June 2026."))
        doc = D.parse("## Summary\n\nCoverage stood at 3 per cent.\n",
                      title="R")
        result = grounding.check(doc, ledger, remove=False)
        assert not result.ok
        assert "3" in {t for f in result.findings for t in f.figures}


# ================================================== units


@pytest.mark.usefixtures("db")
class TestUnitsAreNotInterchangeable:

    def test_a_difference_between_percentages_is_percentage_points(
            self, db, workspace):
        artifact, version = _artifact(db, workspace, _doc("Summary"))
        binding = _bound(db, workspace, value_in_document="6.47",
                         raw_value="6.47", display_value="6.47%",
                         confirmed_by_user=True)
        snapshot = _snapshot(db, artifact, version, binding, value="5.86",
                             display_value="5.86%")
        result = compare.compare_one(snapshot, binding)
        assert result.comparable is True
        assert result.change_unit == "pp"
        # The unit slip that survives review and reaches a committee: a
        # 0.61pp move printed as "0.61%".
        assert "%" not in result.change

    def test_a_unit_change_stops_the_comparison(self, db, workspace):
        artifact, version = _artifact(db, workspace, _doc("Summary"))
        binding = _bound(db, workspace, unit="basis_point",
                         value_in_document="647", display_value="647 bps",
                         confirmed_by_user=True)
        snapshot = _snapshot(db, artifact, version, binding, unit="percent",
                             value="6.47", display_value="6.47%")
        result = compare.compare_one(snapshot, binding)
        assert result.comparable is False
        assert "unit differs" in result.reason

    def test_an_unknown_unit_is_refused_rather_than_defaulted(self):
        """No precision is invented for a unit nobody declared. The refusal
        names the unit rather than falling back to two decimals."""
        with pytest.raises(calc.CalculationError) as raised:
            calc.display_dp("sar millions")
        assert "sar millions" in str(raised.value)

    def test_sar_and_sar_million_are_not_the_same_figure(self):
        """22.77 million and 22770000 are the same money and NOT the same
        token. Grounding is exact, so the document must state what the
        evidence states."""
        ledger = ev.Ledger()
        ledger.add(ev.Item("xlsx://ECL!B2", "sheet_range", "Weighted 22.77"))
        stated = D.parse("## Summary\n\nECL was SAR 22,770,000.\n", title="R")
        assert not grounding.check(stated, ledger, remove=False).ok

        correct = D.parse("## Summary\n\nECL was SAR 22.77 million.\n",
                          title="R")
        assert grounding.check(correct, ledger, remove=False).ok


# ================================================== governance edges


@pytest.mark.usefixtures("db")
class TestGovernanceUnderPressure:

    def test_a_committee_report_with_no_meeting_date_is_scored_not_crashed(
            self, db, workspace):
        from backend.playbook.intelligence import service as intel

        intel.ensure_profile(db, workspace.id, instruction="committee report")
        _artifact(db, workspace, _doc("Summary", "Recommendations"))
        result = score.compute(db, workspace.id)
        assert 0 <= result.completion_pct <= 100
        assert any("period" in c["explanation"] or "date" in c["explanation"]
                   for c in result.components)

    def test_a_finding_with_no_owner_still_blocks(self, db, workspace):
        finding = gov.raise_finding(db, workspace.id, origin=gov.FROM_RULE,
                                    title="Coverage fell", severity=gov.HIGH,
                                    blocking=True)
        assert finding.owner == ""
        assert gov.blocking_unresolved([finding]) == [finding]

    def test_a_deleted_source_does_not_take_the_workspace_with_it(
            self, db, scope, workspace, committee_report_docx):
        from backend.models.playbook import PlaybookSource

        source = service.add_source(db, scope, workspace.id,
                                    filename="report.docx",
                                    content=committee_report_docx,
                                    source_role="previous_report")
        _artifact(db, workspace, _doc("Summary"))
        db.delete(db.get(PlaybookSource, source.id))
        db.flush()

        from backend.playbook.intelligence import service as intel

        state = intel.dashboard(db, workspace.id).as_dict()
        assert state["available"] is True
        assert state["sources"]["sources"] == 0

    def test_restoring_backwards_moves_forwards(self, db, scope, workspace):
        first = _doc("Summary")
        artifact, _ = _artifact(db, workspace, first)
        second = _doc("Summary", "Appendix")
        _artifact(db, workspace, second, artifact=artifact, version=2)

        restored = service.restore_version(db, scope, artifact.id, 1)
        versions = repo.versions(db, artifact.id)
        assert [v.version for v in versions] == [1, 2, 3]
        assert restored["version"] == 3 and restored["restored_from"] == 1
        # v3 carries v1's content, and v1 is still exactly where it was.
        assert versions[-1].content_hash == first.content_hash()
        assert versions[0].content_hash == first.content_hash()

    def test_a_section_edited_after_a_refresh_reopens_its_sign_off(
            self, db, workspace):
        artifact, _ = _artifact(db, workspace, _doc("Summary"))
        [row] = sect.rows_for(db, artifact.id)
        sect.transition(db, row, to=sect.READY_FOR_REVIEW, actor=ACTOR)
        sect.transition(db, row, to=sect.APPROVED, actor=ACTOR)

        edited = D.parse(f"## Summary\n\nRewritten after refresh. {BODY}\n\n",
                         title="Report")
        _artifact(db, workspace, edited, artifact=artifact, version=2)
        assert sect.rows_for(db, artifact.id)[0].status == sect.NEEDS_REVIEW

    def test_the_wrong_metric_can_be_unbound_without_losing_the_figure(
            self, db, workspace):
        """A user maps the wrong metric. Ignoring it must keep the figure
        visible and stop it being a governed link — not delete the reading."""
        [row] = bind.apply(db, workspace.id, bind.from_labels(
            [("Gini", "0.415", "0.4152", "B2")], locator="xlsx://P!A1"))
        bind.ignore(db, row)
        assert bind.is_governed(row) is False
        assert row.display_value == "0.415"


# ================================================== scale


@pytest.mark.usefixtures("db")
class TestScale:

    def test_fifty_tracked_metrics_are_all_inventoried(self, db, workspace):
        bind.apply(db, workspace.id, bind.from_export([
            contract.Metric(metric_id=f"m.{n}", label=f"Metric {n}",
                            value=str(n), display_value=str(n))
            for n in range(50)]))
        state = bind.coverage(
            db.query(PlaybookMetricBinding)
            .filter(PlaybookMetricBinding.workspace_id == workspace.id).all())
        assert state["detected"] == 50 and state["confirmed"] == 50

    @pytest.mark.parametrize("sections", [20, 100])
    def test_a_long_document_keeps_every_section_distinct(self, db, workspace,
                                                          sections):
        doc = _doc(*[f"{n}. Section {n}" for n in range(1, sections + 1)])
        artifact, _ = _artifact(db, workspace, doc)
        rows = sect.rows_for(db, artifact.id)
        assert len(rows) == sections
        assert len({r.section_key for r in rows}) == sections

    def test_a_long_title_survives_the_round_trip(self, db):
        title = ("Auto Loan Application Scorecard (AL-AS-v1.0) — Model "
                 "Development, Validation and Independent Review Report for "
                 "the Credit Risk Committee, Q3 2026")
        doc = D.parse(f"# {title}\n\n## 1. Summary\n\n{BODY}\n", title=title)
        assert doc.title == title
        assert [s.heading for s in doc.sections] == ["1. Summary"]
        assert D.Document.from_dict(doc.as_dict()).title == title

    def test_a_long_heading_is_found_in_a_rendered_pdf(self):
        from backend.playbook import render

        heading = ("4. Independent validation of the calibration, "
                   "discrimination and stability of the behavioural "
                   "scorecard across the retail portfolio")
        doc = D.parse(f"## {heading}\n\n{BODY}\n", title="Report")
        result = validate.validate(render.render(doc, "pdf"), "pdf", doc)
        assert result.ok, result.issues


# ================================================== awkward workbooks


class TestAwkwardWorkbooks:

    @staticmethod
    def _workbook(build) -> bytes:
        from openpyxl import Workbook

        wb = Workbook()
        build(wb)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_a_formula_with_no_cached_value_is_not_reported_as_a_result(self):
        from backend.playbook import ingest

        def build(wb):
            sheet = wb.active
            sheet.title = "ECL"
            sheet.append(["Scenario", "ECL"])
            sheet.append(["Base", 19.20])
            sheet["B3"] = "=SUM(B2:B2)"

        _, result = ingest.read("book.xlsx", self._workbook(build))
        # A sheet chunk carries its header in `text` and its readings in
        # `data`; a formula string is never a verified result, so the cell
        # comes back empty rather than as "=SUM(...)".
        cells = [str(v) for c in result.chunks
                 for row in c.data.get("rows", []) for v in row]
        assert not any("=SUM" in v for v in cells)
        assert "19.2" in cells

    def test_a_hidden_sheet_is_reported_rather_than_silently_dropped(self):
        from backend.playbook import ingest

        def build(wb):
            wb.active.append(["a", 1])
            hidden = wb.create_sheet("Working")
            hidden.append(["scratch", 2])
            hidden.sheet_state = "hidden"

        _, result = ingest.read("book.xlsx", self._workbook(build))
        assert result.manifest.complete is False
        assert any("Working" in str(s) for s in result.manifest.skipped)

    def test_two_reporting_periods_in_one_workbook_both_survive(self):
        from backend.playbook import ingest

        def build(wb):
            sheet = wb.active
            sheet.title = "Exposure"
            sheet.append(["Period", "Exposure"])
            sheet.append(["Q2 2026", 1000])
            sheet.append(["Q3 2026", 1050])

        _, result = ingest.read("book.xlsx", self._workbook(build))
        cells = [str(v) for c in result.chunks
                 for row in c.data.get("rows", []) for v in row]
        assert "Q2 2026" in cells and "Q3 2026" in cells
        assert "1000" in cells and "1050" in cells


@pytest.mark.usefixtures("db")
class TestDuplicateUploads:

    def test_the_same_file_twice_is_two_sources_not_one_overwritten(
            self, db, scope, workspace, committee_report_docx):
        """Deduplicating by hash would be wrong: a user who uploads the same
        file under two roles means it twice. What must not happen is the
        second upload overwriting the first's parse."""
        first = service.add_source(db, scope, workspace.id,
                                   filename="report.docx",
                                   content=committee_report_docx,
                                   source_role="previous_report")
        second = service.add_source(db, scope, workspace.id,
                                    filename="report.docx",
                                    content=committee_report_docx,
                                    source_role="methodology")
        assert first.id != second.id
        assert first.sha256 == second.sha256
        assert len(repo.chunks(db, first.id)) == len(repo.chunks(db,
                                                                 second.id))
        assert first.source_role == "previous_report"
        assert second.source_role == "methodology"


# ================================================== precision and drift


@pytest.mark.usefixtures("db")
class TestPrecisionUnderChange:

    def test_a_re_read_that_changes_displayed_precision_changes_no_document(
            self, db, scope, workspace):
        """§18's hardest case: the reading improves, and the already-approved
        document is left exactly as it was."""
        from openpyxl import Workbook

        from backend.playbook import reparse

        wb = Workbook()
        sheet = wb.active
        sheet.title = "P"
        sheet.append(["AUC", 0.5593220338983])
        sheet["B1"].number_format = "0.000"
        buf = io.BytesIO()
        wb.save(buf)

        source = service.add_source(db, scope, workspace.id,
                                    filename="p.xlsx", content=buf.getvalue(),
                                    source_role="results")
        artifact, _ = _artifact(db, workspace, _doc("Summary"))
        before = repo.versions(db, artifact.id)[0].content_hash

        reparse.latest(db, source.id).parser_version = "0"
        db.flush()
        reparse.reread(db, scope, source.id)

        versions = repo.versions(db, artifact.id)
        assert len(versions) == 1
        assert versions[0].content_hash == before

    def test_a_scoped_edit_does_not_drift_across_ten_rounds(self, db):
        """Progressive formatting drift: each round re-parses the previous
        round's canonical form. A pipeline that normalised a little each time
        would show it by round ten."""
        doc = D.parse(
            "## 1. Summary\n\nECL was SAR 22.77 million.\n\n"
            "## 2. Detail\n\n| Scenario | ECL |\n| --- | --- |\n"
            "| Base | 19.20 |\n", title="Report")
        hashes = []
        for _ in range(10):
            doc = D.Document.from_dict(doc.as_dict())
            hashes.append(merge.section_hash(doc.sections[1]))
        assert len(set(hashes)) == 1

    def test_a_table_cell_survives_ten_canonical_round_trips_unchanged(self):
        doc = D.parse(
            "## Detail\n\n| Scenario | ECL |\n| --- | --- |\n"
            "| Base | 19.20 |\n| Downturn | 36.00 |\n", title="Report")
        original = doc.as_dict()
        for _ in range(10):
            doc = D.Document.from_dict(doc.as_dict())
        assert doc.as_dict() == original
