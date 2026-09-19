"""Fixtures for the Playbook suite.

Real files, built here rather than committed as binaries: a test that reads a
document nobody can regenerate is a test whose failure nobody can diagnose.
"""

from __future__ import annotations

import io
import time

import pytest


def _stub_matches(stand_in, real) -> None:
    """Refuse a stand-in the product would call differently.

    A stub whose signature has fallen behind the real function fails deep
    inside whichever test happens to call it first, with a TypeError that
    reads like a product defect. It has cost real time twice — once here and
    once in the soak harness — so the mismatch is caught at patch time,
    naming the parameter that is missing.
    """
    import inspect

    wanted = set(inspect.signature(real).parameters)
    have = set(inspect.signature(stand_in).parameters)
    missing = wanted - have - {"args", "kwargs"}
    assert not missing, (
        f"the stub for {real.__name__}() is missing {sorted(missing)}; add "
        f"it rather than letting every test that uses this fixture fail")


@pytest.fixture
def committee_report_docx() -> bytes:
    """A previous-period committee report, with the structure Playbook edits."""
    from docx import Document

    doc = Document()
    doc.core_properties.title = "IFRS 9 Committee Report — Q1 2026"
    doc.core_properties.author = "Credit Risk"

    doc.sections[0].header.paragraphs[0].text = "IFRS 9 Committee Report — Q1 2026"

    doc.add_heading("1. Executive summary", level=1)
    doc.add_paragraph(
        "Weighted ECL for the quarter was SAR 20.90 million against exposure of "
        "SAR 1,000 million, a coverage ratio of 2.09 per cent."
    )
    doc.add_heading("2. Scenario results", level=1)
    doc.add_heading("2.1 Weighted outcome", level=2)
    doc.add_paragraph("Scenario weights were unchanged at 60/15/25.")
    table = doc.add_table(rows=4, cols=2)
    rows = [("Scenario", "ECL, SAR million"), ("Base", "18.00"),
            ("Upturn", "14.00"), ("Downturn", "32.00")]
    for r, (a, b) in enumerate(rows):
        table.rows[r].cells[0].text = a
        table.rows[r].cells[1].text = b
    doc.add_heading("3. Limitations", level=1)
    doc.add_paragraph("Post-model adjustments are not covered in this report.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def results_workbook_xlsx() -> bytes:
    """A current-period results workbook, including the two awkward cases:
    a hidden sheet, and a formula with no cached value."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "ECL"
    for row in [
        ["Scenario", "ECL, SAR million", "Weight"],
        ["Base", 19.20, 0.60],
        ["Upturn", 15.00, 0.15],
        ["Downturn", 36.00, 0.25],
    ]:
        ws.append(row)
    # Written as a formula and never calculated, so no cached value exists.
    ws["B6"] = "=SUMPRODUCT(B2:B4,C2:C4)"
    ws["A6"] = "Weighted"

    exposure = wb.create_sheet("Exposure")
    exposure.append(["Period", "Exposure, SAR million"])
    exposure.append(["Q1 2026", 1000])
    exposure.append(["Q2 2026", 1050])

    working = wb.create_sheet("Working")
    working.append(["scratch", 1])
    working.sheet_state = "hidden"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def methodology_docx() -> bytes:
    """A methodology document naming topics a report should cover."""
    from docx import Document

    doc = Document()
    doc.add_heading("IFRS 9 ECL Methodology", level=1)
    for topic in [
        "Scenario design and weighting",
        "Staging criteria and SICR",
        "Post-model adjustments",
        "Model monitoring and validation",
    ]:
        doc.add_heading(topic, level=2)
        doc.add_paragraph(f"Requirements for {topic.lower()}.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def simple_deck_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Q2 2026 ECL"
    slide.placeholders[1].text = "Weighted ECL rose to SAR 22.77 million."
    slide.notes_slide.notes_text_frame.text = "Mention the coverage ratio."

    blank = prs.slides.add_slide(prs.slide_layouts[6])
    table = blank.shapes.add_table(2, 2, Inches(1), Inches(1),
                                   Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Scenario"
    table.cell(0, 1).text = "ECL"
    table.cell(1, 0).text = "Base"
    table.cell(1, 1).text = "19.20"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture
def text_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.drawString(72, 760, "Behavioural Scorecard Validation — Q2 2026")
    c.drawString(72, 740, "Gini was 0.58 against a limit of 0.45.")
    c.showPage()
    c.drawString(72, 760, "Population stability index was 0.08.")
    c.showPage()
    c.save()
    return buf.getvalue()


@pytest.fixture
def image_only_pdf() -> bytes:
    """A PDF with no text layer at all — a scan, in effect."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.rect(100, 100, 300, 300, fill=1)
    c.showPage()
    c.save()
    return buf.getvalue()


@pytest.fixture
def db():
    """A session against the configured platform database, or a skip.

    Follows the repository's own convention: `database_available()` from the
    root conftest decides, and a missing database skips rather than failing —
    except that Playbook's persistence tests are the point of the suite, so a
    skip here is reported as a skip and never counted as a pass.
    """
    from tests.conftest import database_available

    if not database_available():
        pytest.skip("Playbook persistence needs the platform database")

    from sqlalchemy.orm import sessionmaker

    from backend.db.engine import engine

    # Each test runs inside a transaction that is rolled back afterwards, so
    # one test's workspaces cannot become another's search results. The
    # alternative — committing and cleaning up — is what filled this
    # repository's development database with 2,079 identically named Projects
    # (see backend/demo/workspace.py).
    connection = engine.connect()
    transaction = connection.begin()
    session = sessionmaker(bind=connection, expire_on_commit=False)()
    try:
        yield session
    finally:
        session.close()
        transaction.rollback()
        connection.close()


@pytest.fixture
def scope():
    from backend.playbook.repository import Scope

    return Scope(tenant="test-tenant", user_id=None)


@pytest.fixture
def workspace(db, scope):
    from backend.playbook import repository as repo

    ws = repo.create_workspace(db, scope, title="IFRS 9 Committee Report",
                               document_family="ifrs9_committee_report")
    db.flush()
    return ws


@pytest.fixture(autouse=True)
def _configured_author_model(request, monkeypatch):
    """Give every Playbook test a configured assistant model.

    The conversational runtime checks for one before it does anything, which
    the old authoring path did not — it went straight to a patched
    `provider.author`. Without this, a test that patches the provider but
    never mentions a model now fails with a configuration error instead of
    exercising what it is about.

    `tests/conftest.py`'s `_offline_ai` still forces the offline provider, so
    naming a model here cannot produce a network call. The one test that is
    ABOUT a missing model overrides this deliberately.
    """
    from backend.llm import roles as role_config

    # The real Role, with a model set. A hand-rolled stub drifts from the
    # dataclass the moment a field is added, which is how this first went
    # wrong: telemetry reads `role.inherited` and the stub had never heard
    # of it.
    configured = role_config.Role(name=role_config.AUTHOR,
                                  model="scripted-author",
                                  effort="standard", inherited=False)
    monkeypatch.setattr("backend.playbook.assistant.role_config.role",
                        lambda _name: configured, raising=False)

    # And a client that cannot reach a network, whatever else a test forgets
    # to patch.
    #
    # The first attempt here supplied a fake credential instead, so that the
    # product's own configuration check would run for real. It did — and then
    # a test that had not scripted `provider._call` built a genuine client and
    # sent a genuine request, which came back 401. Nothing was generated and
    # nothing was spent, but a test suite that can reach the provider at all is
    # one outbound call away from spending money. So the boundary is closed
    # here rather than guarded downstream: a test that wants a call must script
    # `_call`, and one that forgets fails loudly against this object rather
    # than quietly leaving the building.
    if request.node.get_closest_marker("real_provider_client"):
        # A test ABOUT the client builder has to get the real one. Marked
        # rather than detected, so opting out is a decision in the test file
        # and not a coincidence of its name.
        return

    class _NoNetwork:
        def __getattr__(self, name):
            raise AssertionError(
                "This test reached the provider client. Script "
                "`provider._call` (see the `scripted_author` fixture) rather "
                "than letting a request leave the process.")

    monkeypatch.setattr("backend.playbook.provider._client",
                        lambda: _NoNetwork())


@pytest.fixture
def scripted_author(monkeypatch):
    """Stand in for the provider, so the pipeline around it can be tested.

    Returns a setter: the test supplies the Markdown the author "wrote" and,
    optionally, files it "produced". Nothing here reaches a network, and results
    obtained this way are never reported as live verification.
    """
    from backend.playbook import provider

    state = {"text": "", "files": [], "model": "scripted-author",
             "chunk": 40, "pause": 0.0}

    def fake_author(*, system, messages, formats, purpose="playbook_authoring",
                    container_id="", on_milestone=None, on_delta=None,
                    is_cancelled=None, document_tools=None):
        if is_cancelled and is_cancelled():
            raise provider.Cancelled("stopped")
        if on_milestone:
            on_milestone("drafting", "scripted")
        # Delivered in pieces, like the real one. A stand-in that hands over
        # the whole answer at once would let a streaming bug through every
        # test that uses it.
        if on_delta:
            text = state["text"]
            size = state.get("chunk", 40)
            for i in range(0, len(text), size):
                if is_cancelled and is_cancelled():
                    raise provider.Cancelled("stopped")
                on_delta(text[i:i + size])
                if state.get("pause"):
                    time.sleep(state["pause"])
        result = provider.AuthoringResult(
            text=state["text"],
            files=list(state["files"]),
            model_requested=state["model"],
            model_served=state["model"],
            request_ids=["req_scripted"],
            turns=1,
        )
        state["last_system"] = system
        state["last_user"] = messages[0]["content"]
        return result

    monkeypatch.setattr(provider, "author", fake_author)
    monkeypatch.setattr("backend.playbook.service.provider.author", fake_author)

    # ---- and the conversational runtime the real message path now uses ----
    #
    # A user's message reaches `assistant.converse`, which decides whether the
    # turn needs a file by calling a tool. So the stand-in has to script that
    # DECISION as well as the document: `makes_document` says whether the
    # scripted assistant reaches for `create_document`, which is the one thing
    # a test cannot infer from the Markdown it supplied.
    #
    # It defaults to True so that every journey written when authoring was the
    # only path still means what it meant. A test for an ordinary question —
    # the case that had no way to exist before — passes False.

    class _Block:
        def __init__(self, type_, **kw):
            self.type = type_
            for k, v in kw.items():
                setattr(self, k, v)

    class _Response:
        def __init__(self, content, stop_reason, model):
            self.content = content
            self.stop_reason = stop_reason
            self.model = model
            self._request_id = "req_scripted_chat"
            self.usage = type("U", (), {"input_tokens": 0, "output_tokens": 0})()

    def fake_call(client, *, model, system, messages, tools, container,
                  purpose, role, on_delta=None, is_cancelled=None,
                  deadline=None, with_tools=False, tool_choice=None):
        if is_cancelled and is_cancelled():
            raise provider.Cancelled("stopped")
        state.setdefault("chat_calls", []).append(
            {"system": system, "messages": [dict(m) for m in messages],
             "tools": [t["name"] for t in (tools or [])],
             "tool_choice": tool_choice})

        # Whether a tool has already run IN THIS conversation, read from the
        # messages rather than from a counter. A counter is global across
        # generations, so the second message in a workspace would never reach
        # the tool — which is how this first went wrong, silently turning a
        # timeout test into a test of a chat reply.
        already_used_a_tool = any(
            isinstance(m.get("content"), list)
            and any(isinstance(b, dict) and b.get("type") == "tool_result"
                    for b in m["content"])
            for m in messages)
        wants = state.get("makes_document", True) and any(
            t["name"] == "create_document" for t in (tools or []))
        if wants and not already_used_a_tool:
            # A real assistant says what it is about to do before it does it,
            # and the interface shows that text while the document is being
            # written. A stand-in that goes straight to the tool call would
            # make the thread look frozen and would hide a streaming defect.
            opening = "I will prepare that now."
            if on_delta:
                on_delta(opening)
            return _Response(
                [_Block("text", text=opening),
                 _Block("tool_use", name="create_document", id="tu_chat",
                        input={"instruction": "as asked",
                               "formats": list(state.get("chat_formats")
                                               or ["docx", "pdf"])})],
                "tool_use", state["model"])

        said = state.get("chat_text") or _plain(state["text"])
        if on_delta:
            size = state.get("chunk", 40)
            for i in range(0, len(said), size):
                if is_cancelled and is_cancelled():
                    raise provider.Cancelled("stopped")
                on_delta(said[i:i + size])
                if state.get("pause"):
                    time.sleep(state["pause"])
        return _Response([_Block("text", text=said)], "end_turn", state["model"])

    def _plain(markdown: str) -> str:
        """What the assistant would say, as opposed to what it wrote."""
        lines = [ln for ln in (markdown or "").splitlines()
                 if ln.strip() and not ln.lstrip().startswith("#")]
        return " ".join(lines)[:600] or "Done."

    _stub_matches(fake_call, provider._call)
    monkeypatch.setattr(provider, "_call", fake_call)
    monkeypatch.setattr(provider, "_client", lambda: object())

    # The autouse `_configured_author_model` fixture already installs a real
    # Role with a model set, and this fixture runs inside it.

    def configure(text: str, files=None, model: str = "scripted-author",
                  chunk: int = 40, pause: float = 0.0,
                  makes_document: bool = True, chat_text: str = "",
                  chat_formats=None):
        state["text"] = text
        state["files"] = files or []
        state["model"] = model
        state["chunk"] = chunk
        state["pause"] = pause
        state["makes_document"] = makes_document
        state["chat_text"] = chat_text
        state["chat_formats"] = chat_formats
        state["chat_calls"] = []
        return state

    configure.state = state
    return configure


def pytest_configure(config):
    config.addinivalue_line(
        "markers",
        "real_provider_client: needs the real provider._client builder rather "
        "than the no-network stand-in; for tests ABOUT how the client is built")
