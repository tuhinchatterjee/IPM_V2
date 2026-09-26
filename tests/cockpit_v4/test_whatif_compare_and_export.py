""""Compare two frozen ledgers" and "reconcile the workbook to the chat".

R12 and R14 of section 16.3. Both were previously reported honestly and
neither was blocked by anything: R12 was a library function nobody had
written, and R14 needed a script rather than a protected-core route.

NO DATABASE · NO MODEL. Ledgers are built from `RowResult`s here, so the
arithmetic is the subject and a query is not.
"""

from __future__ import annotations

import json
from decimal import Decimal as D
from pathlib import Path

import pytest

from backend.cockpit_v4.scenario import delta as dl
from backend.cockpit_v4.scenario import ledger as lg
from backend.cockpit_v4.scenario.errors import ScenarioError

HASH = "h" * 64


def ledger(scenario: dict[str, str], *, domain_id: str = "corporate",
           period: str = "2026Q2", membership_hash: str = HASH,
           baseline: str = "100",
           book_baseline: str = "1000") -> lg.Ledger:
    return lg.build(
        domain_id=domain_id, period=period, membership_hash=membership_hash,
        labels=None, book_baseline=D(book_baseline),
        results={key: dl.RowResult(D(baseline), D(value), dl.SCALED)
                 for key, value in scenario.items()})


# ---- R12: two frozen ledgers ------------------------------------------

def test_two_runs_over_one_cohort_are_comparable() -> None:
    """The case it exists for: two methods, one population, one period."""
    delta = ledger({"f1": "120", "f2": "110"})
    emulator = ledger({"f1": "130", "f2": "110"})
    got = lg.compare(delta, emulator, left_name="Delta",
                     right_name="Emulator")
    assert got.baseline_gap == 0, (
        "one cohort at one period has one baseline, so a non-zero gap here "
        "means the two runs read different data for the same rows.")
    assert got.scenario_gap == D("10")
    assert got.change_gap == D("10")
    assert got.relative_gap() == D("5")
    assert got.moved() == [("f1", D("120"), D("130"))]


def test_a_row_only_one_side_carries_is_reported_separately() -> None:
    """An entity present on one side is a cohort mismatch, not a value that
    changed, and folding it into the differences would read as the latter."""
    left = ledger({"f1": "120"})
    right = ledger({"f1": "120", "f2": "110"})
    got = lg.compare(left, right)
    assert got.moved() == []
    assert got.only_in_one()["second"] == ["f2"]


def test_a_disposition_disagreement_is_surfaced() -> None:
    """The kind that hides in a total: an ineligible row keeps its baseline,
    so two runs that disagree about what could be calculated can differ by
    very little and mean something quite different."""
    left = lg.build(
        domain_id="corporate", period="2026Q2", membership_hash=HASH,
        labels=None, book_baseline=D("1000"),
        results={"f1": dl.RowResult(D("100"), D("120"), dl.SCALED)})
    right = lg.build(
        domain_id="corporate", period="2026Q2", membership_hash=HASH,
        labels=None, book_baseline=D("1000"),
        results={"f1": dl.RowResult(D("100"), D("100"), dl.INELIGIBLE,
                                    "no undrawn to convert")})
    got = lg.compare(left, right)
    rows = got.dispositions()
    assert len(rows) == 1
    assert rows[0]["first"] == dl.SCALED
    assert rows[0]["second"] == dl.INELIGIBLE
    assert rows[0]["reason_right"] == "no undrawn to convert"


@pytest.mark.parametrize("field,other", [
    ("domain_id", {"domain_id": "retail"}),
    ("period", {"period": "2026Q1"}),
    ("membership_hash", {"membership_hash": "z" * 64}),
])
def test_incomparable_ledgers_are_refused_not_caveated(field, other) -> None:
    """A difference across books, periods or cohorts is a POPULATION effect,
    and publishing it beside a scenario would let it read as one.

    The membership hash is the one that would be tempting to relax: two runs
    over "construction" a quarter apart have different hashes and look like
    the same cohort. Refused, with both values named.
    """
    left = ledger({"f1": "120"})
    right = ledger({"f1": "130"}, **other)
    with pytest.raises(ScenarioError) as caught:
        lg.compare(left, right)
    assert caught.value.code == "RECONCILIATION_FAILED"
    assert field in caught.value.message


def test_a_comparison_cannot_invent_a_number() -> None:
    """It subtracts; it does not re-run. Every figure it publishes is one of
    the two ledgers' own, or a difference of them."""
    left = ledger({"f1": "120"})
    right = ledger({"f1": "130"})
    got = lg.compare(left, right)
    values = {row["left"] for row in got.rows()} | {
        row["right"] for row in got.rows()}
    for value in values:
        if not value:
            continue
        assert value in {str(left.baseline), str(right.baseline),
                         str(left.scenario), str(right.scenario),
                         str(left.change), str(right.change),
                         "120", "130", dl.SCALED, dl.INELIGIBLE}


# ---- R14: the offline workbook ----------------------------------------

def workbook_module():
    import importlib.util

    path = Path("scripts/whatif/build_workbook.py")
    spec = importlib.util.spec_from_file_location("build_workbook", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def result_rows(**over):
    """A result whose arithmetic closes, so a test can break exactly one
    thing."""
    def row(section, item, **kw):
        base = {c: "" for c in workbook_module().HEADERS}
        base.update({"section": section, "item": item})
        base.update(kw)
        return base

    rows = [
        row("cohort", "Total ECL", baseline_sar_mn="100",
            scenario_sar_mn="120", change_sar_mn="20"),
        row("book", "Total ECL", baseline_sar_mn="1000",
            scenario_sar_mn="1020", change_sar_mn="20"),
        row("book", "Outside the cohort", baseline_sar_mn="900",
            scenario_sar_mn="900", change_sar_mn="0"),
        row("method", "Delta (proportional)", baseline_sar_mn="100",
            scenario_sar_mn="120", change_sar_mn="20"),
        row("attribution_economic", "pd_pit_12m", change_sar_mn="20"),
        row("attribution_economic", "Reconciles to", change_sar_mn="20"),
    ]
    for item, changes in over.items():
        for entry in rows:
            if entry["item"] == item:
                entry.update(changes)
    return rows


def test_the_workbook_reconciles_before_it_is_written(tmp_path) -> None:
    module = workbook_module()
    checks = module.reconcile(result_rows())
    assert len(checks) == 4
    assert any("book identity" in c for c in checks)
    assert any("exactly zero" in c for c in checks)


def test_a_broken_book_identity_fails_the_build() -> None:
    """Not a warning in a cell. A workbook is the artifact people forward and
    reconcile against a ledger six months later."""
    module = workbook_module()
    broken = result_rows()
    for row in broken:
        if row["item"] == "Outside the cohort":
            # The remainder no longer adds up with the cohort to the book.
            # Mutated in place rather than appended: `reconcile` takes the
            # FIRST row matching a section and item, so an extra row would
            # have been ignored and this test would have passed by accident.
            row["scenario_sar_mn"] = "950"
    with pytest.raises(module.Mismatch) as caught:
        module.reconcile(broken)
    assert "identity does not hold" in str(caught.value)


def test_a_method_whose_change_does_not_add_up_fails_the_build() -> None:
    module = workbook_module()
    broken = result_rows()
    for row in broken:
        if row["section"] == "method":
            row["change_sar_mn"] = "25"
    with pytest.raises(module.Mismatch):
        module.reconcile(broken)


def test_an_unavailable_method_is_not_reconciled_against_zero() -> None:
    """Section 12: an unavailable method leaves its cells EMPTY. Reading an
    empty cell as zero here would reintroduce the exact substitution the
    reporting rule forbids, inside the check meant to catch it."""
    module = workbook_module()
    rows = result_rows() + [
        {**{c: "" for c in module.HEADERS}, "section": "method",
         "item": "Emulator", "baseline_sar_mn": "100",
         "status": "UNAVAILABLE",
         "note": "this emulator did not pass its gates"}]
    checks = module.reconcile(rows)
    assert not any("Emulator" in c for c in checks)


def test_a_formula_in_a_value_is_stored_as_text(tmp_path) -> None:
    """A sector or borrower name is attacker-influenced text as far as a
    spreadsheet is concerned, and Excel evaluates a cell that starts with
    `=`, `+`, `-`, `@`, a tab or a carriage return."""
    from openpyxl import load_workbook

    module = workbook_module()
    rows = result_rows()
    rows[0]["note"] = '=HYPERLINK("http://example.invalid","click")'
    rows[0]["scope"] = "@SUM(A1:A9)"
    rows[0]["status"] = "-1+1"
    out = tmp_path / "book.xlsx"
    module.write(rows, {"origin": "SYNTHETIC_DEMO"},
                 module.reconcile(rows), out)
    loaded = load_workbook(out)
    cells = [c.value for row in loaded["Summary"].iter_rows()
             for c in row if isinstance(c.value, str)]
    for dangerous in ("=HYPERLINK", "@SUM", "-1+1"):
        matching = [c for c in cells if dangerous in c]
        assert matching, f"{dangerous} did not reach the workbook at all"
        for cell in matching:
            assert not cell.startswith(("=", "@", "+", "-")), (
                f"{cell!r} would be evaluated as a formula when this "
                f"workbook is opened.")


def test_the_workbook_says_what_it_is_measured_on(tmp_path) -> None:
    """A spreadsheet outlives its conversation. The label travels with it."""
    from openpyxl import load_workbook

    module = workbook_module()
    rows = result_rows()
    out = tmp_path / "book.xlsx"
    module.write(rows, {"origin": "SYNTHETIC_DEMO"},
                 module.reconcile(rows), out)
    loaded = load_workbook(out)
    cover = " ".join(str(c.value) for row in loaded["Provenance"].iter_rows()
                     for c in row if c.value)
    assert "GENERATED BOOK" in cover.upper()
    assert "bank-validated" in cover


def test_the_ml_sheet_says_it_is_not_a_decomposition(tmp_path) -> None:
    """The two attributions are one sheet apart and answer different
    questions. A reader who adds them has been misled by the layout."""
    from openpyxl import load_workbook

    module = workbook_module()
    rows = result_rows() + [
        {**{c: "" for c in module.HEADERS}, "section": "ml_explanation",
         "item": "dpd_days", "unit": "mean |SHAP|, model rate space",
         "note": "0.037 over 248 rows"}]
    out = tmp_path / "book.xlsx"
    module.write(rows, {}, module.reconcile(rows), out)
    loaded = load_workbook(out)
    text = " ".join(str(c.value) for row in loaded["ML explanation"].iter_rows()
                    for c in row if c.value)
    assert "NOT A DECOMPOSITION" in text
    assert "Do not add the two" in text


def test_the_workbook_round_trips_a_real_result(tmp_path) -> None:
    """End to end from the shape the scenario step actually publishes."""
    module = workbook_module()
    payload = {"rows": result_rows(),
               "provenance": {"whatif_scenario_id": "sc-1",
                              "origin": "SYNTHETIC_DEMO"}}
    source = tmp_path / "result.json"
    source.write_text(json.dumps(payload), encoding="utf-8")
    out = tmp_path / "book.xlsx"
    written = module.write(
        payload["rows"], payload["provenance"],
        module.reconcile(payload["rows"]), out)
    assert written.exists() and written.stat().st_size > 4000
