"""U15, O06, M09, J11 and the API's access rules."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile

import pytest
from conftest import QUESTION, make_service, run

from backend.model_lab import export


def test_O06_formula_injection_and_html_are_neutralised():
    assert export.safe_cell("=HYPERLINK(\"x\")").startswith("'")
    assert export.safe_cell("@SUM(A1)").startswith("'")
    assert export.safe_cell("-12.5") == "-12.5"
    assert export.safe_cell(3) == 3
    assert export.safe_name("../../etc/passwd") == "etc_passwd"
    html = export._answer_html({"display_name": "<b>x</b>",
                                "execution_state": "OK", "fixture": False,
                                "answer": {"narrative": "<script>alert(1)"
                                                        "</script>"}})
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert "Content-Security-Policy" in html


def test_U15_M09_pack_is_complete_consistent_and_keeps_unknowns(demo):
    svc, cid, ev = demo
    exp = svc.coord.store.latest_export(cid, svc.cfg.tenant_id)
    assert exp["state"] == "READY"
    z = zipfile.ZipFile(exp["path"])
    names = set(z.namelist())
    for need in ("README.html", "manifest.json", "summary.csv", "stages.csv",
                 "calls.csv", "checks.csv", "claims.csv", "failures.csv",
                 "resource_samples.csv", "comparison.xlsx", "events.jsonl",
                 "reviews.jsonl", "limitations.md", "checksums.sha256",
                 "opus_match.csv"):
        assert need in names, need
    assert all(not n.startswith("/") and ".." not in n for n in names)
    for line in z.read("checksums.sha256").decode().splitlines():
        digest, name = line.split("  ", 1)
        assert hashlib.sha256(z.read(name)).hexdigest() == digest
    rows = list(csv.DictReader(io.StringIO(z.read("summary.csv").decode())))
    opus = next(r for r in rows if r["profile_id"] == "opus-frozen")
    assert opus["service_ms"] == "" and opus["execution_state"] == "BLOCKED"
    ref = next(r for r in rows if r["profile_id"] == "fixture-reference")
    assert ref["first_protocol_event_ms"] == ""
    assert ref["first_protocol_event_ms_status"] == "UNAVAILABLE"
    # UI evaluation and CSV come from one calculation source.
    ui = next(k for k in ev["children"]
              if k["profile_id"] == "fixture-reference")
    assert float(ref["service_ms"]) == pytest.approx(
        ui["metrics"]["service_ms"]["value"])
    m = json.loads(z.read("manifest.json"))
    assert m["frozen_source_id"].startswith("245c50e")
    assert m["release_claim"] == "EXPERIMENTAL"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(z.read("comparison.xlsx")))
    for sheet in ("Read Me", "Models & Profiles", "Summary", "Four Stages",
                  "Calls", "Validation Checks", "Claim Review",
                  "Failure Diagnosis", "Costs & Resources",
                  "Reproducibility"):
        assert sheet in wb.sheetnames
    assert wb["Summary"].freeze_panes == "A2"


def test_O06_no_credentials_in_the_pack(tmp_path):
    svc = make_service(tmp_path, env={"COCKPIT_ANTHROPIC_API_KEY":
                                      "sk-ant-SECRET-123456789"})
    cid, _ = run(svc, ["fixture-reference", "opus-frozen"])
    exp = svc.coord.store.latest_export(cid, svc.cfg.tenant_id)
    z = zipfile.ZipFile(exp["path"])
    for n in z.namelist():
        assert b"sk-ant-SECRET" not in z.read(n), n


def test_J11_export_is_rebuilt_without_inference(demo):
    svc, cid, _ = demo
    before = svc.coord.runs._connect().execute(
        "SELECT COUNT(*) FROM runs").fetchone()[0]
    res = svc.export(cid)
    assert res["state"] == "READY" and res["revision"] >= 2
    after = svc.coord.runs._connect().execute(
        "SELECT COUNT(*) FROM runs").fetchone()[0]
    assert before == after


# ---- API ------------------------------------------------------------------

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("MODEL_LAB_RUNTIME_DIR", str(tmp_path / "lab"))
    monkeypatch.setenv("MODEL_LAB_FIXTURE_APP", "true")
    monkeypatch.setenv("COCKPIT_V4_RUNTIME_DIR", str(tmp_path / "frozen"))
    monkeypatch.setenv("COCKPIT_V4_STATE_DATABASE",
                       str(tmp_path / "frozen" / "s.sqlite3"))
    monkeypatch.setenv("COCKPIT_V4_LOCAL_DEMO_AUTH", "true")
    monkeypatch.setenv("COCKPIT_V4_MEMORY_ENABLED", "false")
    monkeypatch.setenv("COCKPIT_ANTHROPIC_API_KEY", "sk-ant-NEVER-SENT")
    from fastapi.testclient import TestClient

    from backend.model_lab.app import create_lab_app
    return TestClient(create_lab_app())


def test_api_profiles_never_carry_secrets(client):
    body = client.get("/api/v1/model-lab/model-profiles").text
    assert "sk-ant-NEVER-SENT" not in body
    assert "api_key_env" in body            # the NAME only


def test_api_preflight_does_no_inference(client):
    r = client.post("/api/v1/model-lab/comparisons/preflight", json={
        "question": QUESTION, "profile_ids": ["fixture-reference"],
        "comparator_id": "fixture-reference"})
    assert r.status_code == 200 and r.json()["no_inference_performed"]
    assert client.get("/api/v1/model-lab/comparisons").json()[
        "comparisons"] == []


def test_api_guessing_an_id_reveals_nothing(client):
    for path in ("", "/events", "/export"):
        assert client.get(f"/api/v1/model-lab/comparisons/cmp-000000000000"
                          f"{path}").status_code == 404


def test_api_other_tenant_gets_404(client):
    import backend.model_lab.api as api
    svc = api._SVC["svc"]
    original = svc.cfg.tenant_id
    svc.cfg.tenant_id = "someone-else"
    try:
        assert client.get("/api/v1/model-lab/comparisons").status_code == 404
    finally:
        svc.cfg.tenant_id = original


def test_api_replay_is_refused_without_approval(client):
    r = client.post("/api/v1/model-lab/comparisons", json={
        "question": QUESTION, "profile_ids": ["fixture-reference"],
        "comparator_id": "fixture-reference"},
        headers={"Idempotency-Key": "k"})
    cid = r.json()["comparison_id"]
    r = client.post(f"/api/v1/model-lab/comparisons/{cid}/replays",
                    json={"kind": "S4"})
    assert r.status_code == 409
    assert r.json()["detail"]["status"] == "INSUFFICIENT_REFERENCE"


def test_lab_routes_are_not_on_the_frozen_router():
    from backend.cockpit_v4 import routes
    assert not [r for r in routes.router.routes
                if "model-lab" in getattr(r, "path", "")]
