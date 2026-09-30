"""Saved-Opus agreement reference: explicit, separate, read-only.

Three concepts stay apart: the LIVE comparator (a child of this comparison),
the SAVED agreement reference (another comparison's stored Opus child) and
the independent ORACLE (correctness). A fixture registered under the profile
id `opus-frozen` stands in for the saved Opus run; nothing here calls a
model.
"""

from __future__ import annotations

import csv
import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from conftest import MANIFEST_CHECK, QUESTION, ROOT, child, make_service, run

from backend.model_lab import evaluate, registry


def _as_opus(fixture_id: str) -> registry.Profile:
    raw = registry.load_profiles()[fixture_id].raw | {
        "profile_id": "opus-frozen", "display_name": "Saved Opus (fixture)"}
    return registry._validate(raw, Path("opus-frozen.json"))


def _svc(tmp: Path, opus_fixture: str = "fixture-reference"):
    return make_service(tmp, profiles=registry.load_profiles()
                        | {"opus-frozen": _as_opus(opus_fixture)})


def _snapshot(svc, cid: str) -> str:
    st, t = svc.coord.store, svc.cfg.tenant_id
    return json.dumps({
        "evaluations": st.evaluation_history(cid, t),
        "events": len(st.events(cid, t, 0, 10 ** 7)),
        "children": st.children(cid),
        "runs": [st.child_runs(c["child_run_id"]) for c in st.children(cid)],
    }, sort_keys=True, default=str)


@pytest.fixture(scope="module")
def linked(tmp_path_factory):
    svc = _svc(tmp_path_factory.mktemp("lab"))
    ref, _ = run(svc, ["opus-frozen"], comparator="opus-frozen")
    cand, ev = run(svc, ["fixture-reference"], comparator="",
                   reference_comparison_id=ref)
    return svc, ref, cand, ev


# ---- READY, and the live comparator stays separate ------------------------

def test_saved_opus_reference_is_ready_and_separate_from_comparator(linked):
    _, ref, _, ev = linked
    r = ev["saved_reference"]
    assert r["reference_status"] == "READY"
    assert r["reference_source"] == "saved_comparison"
    assert r["reference_comparison_id"] == ref
    assert r["reference_profile_id"] == "opus-frozen"
    assert r["reference_link_source"] == "spec"
    assert isinstance(r["reference_evaluation_revision"], int)
    assert r["reference_evaluator_version"] in \
        evaluate.COMPATIBLE_REFERENCE_VERSIONS
    assert "agreement" in r["note"] and "not truth" in r["note"]
    # The live comparator is NOT dressed up as the saved run.
    assert ev["comparator"]["profile_id"] == ""
    assert ev["comparator"]["status"] == "COMPARATOR_UNAVAILABLE"
    k = child(ev, "fixture-reference")
    assert all(v["display"] == "N/A" for v in k["opus_match"].values())
    assert k["reference_match"]["S1"]["pct"] is not None
    assert k["reference_match"]["S2"]["pct"] == 100.0
    assert k["reference_match"]["S4"]["pct"] is not None


def test_oracle_is_unchanged_by_the_reference(linked):
    svc, _, _, ev = linked
    _, plain = run(svc, ["fixture-reference"], comparator="")
    a, b = child(ev, "fixture-reference"), child(plain, "fixture-reference")
    assert plain["saved_reference"]["reference_status"] == "NO_REFERENCE"
    assert child(plain, "fixture-reference")["reference_match"] is None
    assert [(c["check_id"], c["outcome"], json.dumps(c["actual"],
                                                     sort_keys=True))
            for c in a["checks"]] == \
        [(c["check_id"], c["outcome"], json.dumps(c["actual"],
                                                  sort_keys=True))
         for c in b["checks"]]
    assert [c["verification_status"] for c in a["claims"]] == \
        [c["verification_status"] for c in b["claims"]]
    assert [s["status"] for s in a["stages"].values()] == \
        [s["status"] for s in b["stages"].values()]


# ---- read-only, no model call ---------------------------------------------

def test_reevaluation_reads_the_reference_and_calls_no_model(linked,
                                                            monkeypatch):
    svc, ref, cand, _ = linked
    import anthropic

    import backend.model_lab.adapters as adapters

    def boom(*a, **k):
        raise AssertionError("a model provider was built during "
                             "re-evaluation")
    monkeypatch.setattr(adapters, "build_provider", boom)
    monkeypatch.setattr(anthropic, "Anthropic", boom)
    monkeypatch.setattr(svc.coord, "provider_factory", boom)
    ref_before, cand_runs = _snapshot(svc, ref), [
        svc.coord.store.child_runs(c["child_run_id"])
        for c in svc.coord.store.children(cand)]
    total = svc.coord.runs._connect().execute(
        "SELECT COUNT(*) FROM runs").fetchone()[0]
    body = svc.evaluate(cand)
    assert body["saved_reference"]["reference_status"] == "READY"
    assert _snapshot(svc, ref) == ref_before           # reference untouched
    assert [svc.coord.store.child_runs(c["child_run_id"])
            for c in svc.coord.store.children(cand)] == cand_runs
    assert svc.coord.runs._connect().execute(
        "SELECT COUNT(*) FROM runs").fetchone()[0] == total


# ---- explicit failures, never a silent fallback ---------------------------

def _resolve(svc, cand: str, override: str | None = None) -> dict:
    return evaluate.evaluate_comparison(svc.coord, cand,
                                        reference_override=override)


def test_missing_reference_is_explicit(linked):
    svc, _, cand, _ = linked
    ev = _resolve(svc, cand, override="cmp-doesnotexist")
    r = ev["saved_reference"]
    assert r["reference_status"] == "REFERENCE_COMPARISON_UNAVAILABLE"
    assert r["reference_link_source"] == "operator --reference"
    k = child(ev, "fixture-reference")
    assert all(v["reason"] == "REFERENCE_COMPARISON_UNAVAILABLE"
               for v in k["reference_match"].values())


def test_reference_without_an_opus_child_is_a_blocker(linked):
    svc, _, cand, _ = linked
    no_opus, _ = run(svc, ["fixture-reference"])
    r = _resolve(svc, cand, override=no_opus)["saved_reference"]
    assert r["reference_status"] == "REFERENCE_HAS_NO_OPUS_CHILD"
    assert "no opus-frozen child" in r["remedy"]


def test_incompatible_saved_evaluation_is_a_blocker(tmp_path):
    svc = _svc(tmp_path)
    ref, _ = run(svc, ["opus-frozen"], comparator="opus-frozen")
    cand, _ = run(svc, ["fixture-reference"], comparator="",
                  reference_comparison_id=ref)
    st, t = svc.coord.store, svc.cfg.tenant_id
    old = st.current_evaluation(ref, t)["body"] | {
        "evaluator_version": "lab-eval-2"}
    st.add_evaluation(ref, t, "lab-eval-2", old)     # test setup only
    r = _resolve(svc, cand)["saved_reference"]
    assert r["reference_status"] == "REFERENCE_EVALUATION_INCOMPATIBLE"
    assert r["reference_evaluator_version"] == "lab-eval-2"
    assert f"reevaluate.py --comparison {ref}" in r["remedy"]


def test_reference_opus_that_did_not_complete_gives_no_agreement(tmp_path):
    svc = _svc(tmp_path, opus_fixture="fixture-no-tool")
    ref, _ = run(svc, ["opus-frozen"], comparator="opus-frozen")
    cand, ev = run(svc, ["fixture-reference"], comparator="",
                   reference_comparison_id=ref)
    assert ev["saved_reference"]["reference_status"] == \
        "REFERENCE_OPUS_NOT_COMPLETED"
    k = child(ev, "fixture-reference")
    assert all(v["pct"] is None and v["reason"] ==
               "REFERENCE_OPUS_NOT_COMPLETED"
               for v in k["reference_match"].values())


def test_a_different_question_is_not_comparable(linked):
    svc, ref, _, _ = linked
    other, ev = run(svc, ["fixture-reference"], comparator="",
                    question=QUESTION + " Please answer briefly.",
                    reference_comparison_id=ref)
    assert ev["saved_reference"]["reference_status"] == \
        "REFERENCE_NOT_COMPARABLE"


# ---- export carries the saved reference -----------------------------------

def test_export_contains_the_saved_reference_metadata(linked):
    svc, ref, cand, _ = linked
    res = svc.export(cand)
    assert res["state"] == "READY"
    z = zipfile.ZipFile(res["path"])
    man = json.loads(z.read("manifest.json"))
    assert man["comparator"]["status"] == "COMPARATOR_UNAVAILABLE"
    sr = man["saved_reference"]
    assert sr["reference_status"] == "READY"
    assert sr["reference_comparison_id"] == ref
    assert sr["reference_profile_id"] == "opus-frozen"
    assert sr["reference_source"] == "saved_comparison"
    assert "reference_answer" not in sr
    assert isinstance(man["evaluation_revision"], int)
    readme = z.read("README.html").decode()
    assert "Live comparator: none (COMPARATOR_UNAVAILABLE)" in readme
    assert f"Saved Opus agreement reference: READY — {ref}" in readme
    om = list(csv.DictReader(io.StringIO(z.read("opus_match.csv")
                                         .decode())))
    saved = [r for r in om if r["match_source"] == "saved_comparison"]
    live = [r for r in om if r["match_source"] == "live_comparator"]
    assert len(saved) == 4 and len(live) == 4
    assert all(r["reference_status"] == "READY" and
               r["reference_comparison_id"] == ref for r in saved)
    assert next(r for r in saved if r["stage"] == "S2")["match_pct"] == \
        "100.0"
    summ = list(csv.DictReader(io.StringIO(z.read("summary.csv").decode())))
    assert summ[0]["reference_status"] == "READY"
    assert summ[0]["reference_profile_id"] == "opus-frozen"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(z.read("comparison.xlsx")))
    readme_cells = [c for row in wb["Read Me"].iter_rows(values_only=True)
                    for c in row if c]
    assert any(str(c).startswith("Saved Opus agreement reference: READY")
               for c in readme_cells)
    sheet = list(wb["Opus Match"].iter_rows(values_only=True))
    col = sheet[0].index("match_source")
    assert sum(1 for r in sheet[1:] if r[col] == "saved_comparison") == 4
    lim = z.read("limitations.md").decode()
    assert "SAVED separate Opus run" in lim


# ---- CLI --------------------------------------------------------------------

def test_cli_reports_the_reference_and_rebuilds_the_pack(linked, capsys):
    svc, ref, cand, _ = linked
    sys.path.insert(0, str(ROOT / "scripts" / "model_lab"))
    import reevaluate

    rc = reevaluate.main(["--comparison", cand, "--runtime-dir",
                          str(svc.cfg.runtime_dir), "--export"])
    out = capsys.readouterr().out
    assert rc == 0, out
    assert "live comparator: none (COMPARATOR_UNAVAILABLE)" in out
    assert f"-> READY ({ref}, opus-frozen, evaluation r" in out
    assert f"reference {ref} untouched: True" in out
    assert "export READY:" in out


def test_protected_manifest_is_clean():
    out = subprocess.run([sys.executable,
                          "scripts/model_lab/protected_manifest.py",
                          MANIFEST_CHECK], cwd=ROOT, capture_output=True,
                         text=True)
    assert out.returncode == 0, out.stdout + out.stderr
