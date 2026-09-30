"""Full Model I/O Trace: observation only, complete, and credential-free.

Offline. Mock HTTP servers (httpx.MockTransport) stand in for OpenAI-
compatible and Anthropic endpoints, the real Anthropic SDK builds real
requests, and fixtures drive the real frozen engine. No model is called.
"""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path
from typing import Any

import httpx
import pytest
from conftest import QUESTION, child, make_service, run
from test_adapters import TOOLS, _history, _openai_reference_server
from test_observer_neutrality import Recorder, _drive, _norm

from backend.cockpit_v4.provider import ProviderFailure
from backend.model_lab import io_trace, model_io, registry
from backend.model_lab.adapters import build_provider
from backend.model_lab.adapters.openai_compat import OpenAICompatProvider

CANARY_KEY = "sk-canary-LABTRACE-0123456789abcdef"
CANARY_SSH = "CANARY-SSH-PRIVATE-KEY-b3BlbnNzaC1rZXktdjEAAAAABG5vbmU"
CANARY_COOKIE = "canary_session_cookie_ZZZ999"


def _sse(events: list[dict]) -> str:
    return "".join(f"data: {json.dumps(e)}\n\n" for e in events) + \
        "data: [DONE]\n\n"


def _tool_stream(reasoning: str = "") -> str:
    ev = []
    if reasoning:
        ev.append({"choices": [{"delta": {"reasoning": reasoning}}]})
    ev += [{"id": "r1", "model": "m", "choices": [{"delta": {"tool_calls": [
        {"index": 0, "id": "c1", "function": {
            "name": "execute_analysis", "arguments": '{"objective": '}}]}}]},
        {"choices": [{"delta": {"tool_calls": [{"index": 0, "function": {
            "arguments": '"x", "steps": []}'}}]}}]},
        {"choices": [{"delta": {}, "finish_reason": "tool_calls"}]},
        {"choices": [], "usage": {"prompt_tokens": 11,
                                  "completion_tokens": 7}}]
    return _sse(ev)


def _result_view(r) -> dict[str, Any]:
    return {k: getattr(r, k) for k in (
        "assistant_blocks", "text", "tool_calls", "stop_reason", "model",
        "input_tokens", "output_tokens")}


# ---- 1. the wire recorder is byte-neutral ---------------------------------

@pytest.mark.parametrize("stream", [True, False])
def test_recording_transport_sends_and_returns_identical_bytes(stream):
    seen: list[tuple[bytes, dict]] = []

    def handler(req: httpx.Request) -> httpx.Response:
        seen.append((req.content, dict(req.headers)))
        if stream:
            return httpx.Response(200, text=_tool_stream("let me think"),
                                  headers={"content-type":
                                           "text/event-stream"})
        return httpx.Response(200, json={
            "id": "r", "model": "m", "choices": [{"message": {
                "role": "assistant", "content": None,
                "reasoning": "hidden thoughts",
                "tool_calls": [{"id": "c1", "type": "function", "function": {
                    "name": "execute_analysis",
                    "arguments": '{"objective": "x", "steps": []}'}}]},
                "finish_reason": "tool_calls"}],
            "usage": {"prompt_tokens": 11, "completion_tokens": 7}})

    def call(provider):
        return provider.converse(
            system=[{"type": "text", "text": "S"}], messages=_history(),
            tools=TOOLS, max_tokens=512, model="m",
            tool_choice={"type": "any"})

    plain = OpenAICompatProvider(base_url="http://127.0.0.1:9/v1", model="m",
                                 transport=httpx.MockTransport(handler),
                                 stream=stream)
    traced = OpenAICompatProvider(base_url="http://127.0.0.1:9/v1",
                                  model="m",
                                  transport=httpx.MockTransport(handler),
                                  stream=stream)
    rec = io_trace.attach_wire_recorder(traced)
    a, b = call(plain), call(traced)
    assert len(seen) == 2                      # no extra call, no retry
    assert seen[0] == seen[1]                  # identical bytes and headers
    assert _result_view(a) == _result_view(b)  # identical assembly
    assert b.tool_calls[0]["id"] == "c1"
    (ex,) = rec.exchanges
    view = io_trace.exchange_view(ex, [])
    assert view["request"]["body_bytes"] == len(seen[0][0])
    assert view["request"]["body_sha256"] == io_trace.sha256_text(seen[0][0])
    text = json.dumps(view)
    assert "let me think" not in text and "hidden thoughts" not in text
    assert io_trace.HIDDEN in text
    if stream:
        resp = view["response"]
        assert resp["format"] == "sse" and resp["chunk_count"] >= 1
        assembled = resp["assembled"]["message"]["tool_calls"][0]
        assert json.loads(assembled["function"]["arguments"]) == \
            b.tool_calls[0]["input"]
        assert [e.get("done") for e in resp["events"]][-1] is True


# ---- 2. engine behaviour is identical with the trace on --------------------

def test_trace_on_is_neutral_through_the_frozen_engine():
    from backend.model_lab.adapters.fixture import FixtureProvider
    from backend.model_lab.observe import observe

    bare = Recorder(FixtureProvider("repair", "lab/fixture-repair"))
    out1, (s1, r1), _ = _drive(bare)
    traces: list[dict] = []
    inner = Recorder(FixtureProvider("repair", "lab/fixture-repair"))
    out2, (s2, r2), _ = _drive(observe(inner, lambda r: None,
                                       trace_sink=traces.append))
    assert out1.state == out2.state == "COMPLETED"
    assert len(bare.sent) == len(inner.sent) == len(traces) == 4
    for a, b in zip(bare.sent, inner.sent, strict=True):
        assert _norm(a) == _norm(b)
    assert [(e.event_type, e.status) for e in s1.events_since(r1)] == \
        [(e.event_type, e.status) for e in s2.events_since(r2)]
    f1, f2 = s1.get_run(r1).final_response, s2.get_run(r2).final_response
    assert _norm(f1) == _norm(f2)
    # The trace holds the engine request exactly as sent (normalised the
    # same way), including the full system prompt.
    for sent, tr in zip(inner.sent, traces, strict=True):
        assert _norm(tr["engine_request"]) == _norm(sent)
        assert tr["dispatch_status"] == "IN_PROCESS_NO_NETWORK"
    assert len(json.dumps(traces[0]["engine_request"]["system"])) > 10_000


def _ready(p):
    import dataclasses
    return dataclasses.replace(p, declared_status="READY_E2E", raw=p.raw | {
        "requires_probe": False, "status": "READY_E2E"})


def _openai_run(tmp_path, *, trace: bool, extra_env: dict | None = None,
                handler=None):
    handler_, seen = _openai_reference_server()
    handler = handler or handler_
    loaded = registry.load_profiles()
    prof = _ready(loaded["qwen3.5-4b"])
    raw = prof.raw | {"endpoint": prof.raw["endpoint"] | {
        "api_key_env": "LAB_FAKE_API_KEY"}}
    prof = registry._validate(raw, Path("qwen3.5-4b.json"))
    bodies: list[bytes] = []
    headers: list[dict] = []

    def factory(profile):
        p = build_provider(profile, env={"LAB_FAKE_API_KEY": CANARY_KEY})
        if profile.profile_id == "qwen3.5-4b":
            def h(req):
                bodies.append(req.content)
                headers.append(dict(req.headers))
                return handler(req)
            p._transport = httpx.MockTransport(h)
        return p

    svc = make_service(tmp_path, profiles=loaded | {"qwen3.5-4b": prof},
                       provider_factory=factory,
                       env={"LAB_FAKE_API_KEY": CANARY_KEY}
                       | (extra_env or {}))
    svc.cfg.full_io_trace = trace
    cid, ev = run(svc, ["qwen3.5-4b"], comparator="")
    return svc, cid, ev, bodies, headers, seen


def test_trace_on_vs_off_sends_identical_payloads_and_answers(tmp_path):
    on = _openai_run(tmp_path / "on", trace=True)
    off = _openai_run(tmp_path / "off", trace=False)
    b_on, b_off = on[3], off[3]
    assert len(b_on) == len(b_off) >= 2        # no extra call or retry
    for x, y in zip(b_on, b_off, strict=True):
        assert _norm(json.loads(x)) == _norm(json.loads(y))
    k_on, k_off = child(on[2], "qwen3.5-4b"), child(off[2], "qwen3.5-4b")
    assert k_on["execution_state"] == k_off["execution_state"] == "COMPLETED"
    assert [c["tool_names"] for c in k_on["calls"]] == \
        [c["tool_names"] for c in k_off["calls"]]
    assert _norm(k_on["answer"]) == _norm(k_off["answer"])
    assert [(c["check_id"], c["outcome"]) for c in k_on["checks"]] == \
        [(c["check_id"], c["outcome"]) for c in k_off["checks"]]
    t_on = model_io.build(on[0].coord, on[1])
    t_off = model_io.build(off[0].coord, off[1])
    assert t_on["children"][0]["traced_calls"] == len(b_on)
    assert t_off["children"][0]["traced_calls"] == 0


# ---- 3. the trace is complete and exact -----------------------------------

@pytest.fixture(scope="module")
def traced(tmp_path_factory):
    import os
    os.environ["LAB_FAKE_API_KEY"] = CANARY_KEY
    os.environ["LAB_TEST_SSH_PRIVATE_KEY"] = CANARY_SSH
    os.environ["LAB_TEST_COOKIE_SECRET"] = CANARY_COOKIE
    base, _ = _openai_reference_server()

    def leaky(req):
        # A hostile endpoint: sets a cookie and echoes secrets in a body.
        r = base(req)
        return httpx.Response(r.status_code, text=r.text, headers={
            "content-type": "text/event-stream",
            "set-cookie": f"sid={CANARY_COOKIE}",
            "x-echo-key": CANARY_KEY})
    try:
        yield _openai_run(tmp_path_factory.mktemp("lab"), trace=True,
                          handler=leaky)
    finally:
        for k in ("LAB_FAKE_API_KEY", "LAB_TEST_SSH_PRIVATE_KEY",
                  "LAB_TEST_COOKIE_SECRET"):
            os.environ.pop(k, None)


def test_every_call_has_four_views_and_exact_tool_results(traced):
    svc, cid, _, bodies, headers, _ = traced
    tr = model_io.build(svc.coord, cid, include_bodies=True)
    ch = tr["children"][0]
    calls = [i for i in ch["timeline"] if i["kind"] == "call"]
    tools = [i for i in ch["timeline"] if i["kind"] == "tool"]
    assert len(calls) == len(bodies) and len(tools) >= len(calls)
    for c, sent in zip(calls, bodies, strict=True):
        v = c["views"]
        assert c["dispatch_status"] == "SENT"
        assert set(v) == set(model_io.VIEWS) | {
            "assistance_appended_to_request"}
        assert v["assistance_appended_to_request"] is None    # baseline
        er = v["engine_request"]
        assert er["request"]["system"] and er["request"]["tools"]
        assert er["metadata"]["context_capacity_tokens"] == 32_768
        (wr,) = v["wire_request"]
        assert wr["body_bytes"] == len(sent)
        assert wr["body_sha256"] == io_trace.sha256_text(sent)
        assert wr["body"] == json.loads(sent)          # nothing to redact
        assert wr["url"] == "http://127.0.0.1:11434/v1/chat/completions"
        (raw,) = v["wire_response_raw"]
        assert raw["status_code"] == 200 and raw["format"] == "sse"
        assert v["normalized_response"]["response"]["tool_calls"]
        assert c["stage_tags"]
    # The tool_result each round trip reports is byte-for-byte the tool
    # message the NEXT wire request carried.
    for t in tools:
        if t["tool_result_source"].startswith("next model request"):
            nxt = json.loads(bodies[t["call_n"]])        # call_n is 1-based
            carried = [m for m in nxt["messages"] if m["role"] == "tool"
                       and m.get("tool_call_id") ==
                       t["model_tool_call"]["id"]]
            assert carried and carried[0]["content"] == \
                t["tool_result_returned"]["content"]
    assert tools[-1]["is_final_response"]
    assert any(t["validation"] for t in tools)
    assert any(t["execution"] for t in tools)
    links = ch["stage_links"]
    assert links["S1"]["calls"] or links["S1"]["shared_calls"]
    assert links["S4"]["tool_roundtrips"]


def test_no_secret_reaches_the_trace_store_or_the_export(traced):
    svc, cid, _, _, headers, _ = traced
    assert headers[0]["authorization"] == f"Bearer {CANARY_KEY}"  # it was sent
    canaries = (CANARY_KEY, CANARY_SSH, CANARY_COOKIE)
    for e in svc.coord.store.events(cid, svc.cfg.tenant_id, 0, 10 ** 7):
        if e.get("payload_ref"):
            blob = svc.coord.store.get_blob(e["payload_ref"]).decode()
            assert not any(c in blob for c in canaries), e["event_type"]
    tr = json.dumps(model_io.build(svc.coord, cid, include_bodies=True))
    assert not any(c in tr for c in canaries)
    low = tr.lower()
    for key in ('"authorization":', '"x-api-key":', '"set-cookie":',
                '"cookie":', '"x-echo-key":', '"api-key":'):
        assert key not in low, key
    res = svc.export(cid)
    with zipfile.ZipFile(res["path"]) as z:
        for n in z.namelist():
            data = z.read(n)
            blobs = [data]
            if n.endswith(".xlsx"):
                with zipfile.ZipFile(io.BytesIO(data)) as x:
                    blobs += [x.read(m) for m in x.namelist()]
            for b in blobs:
                assert not any(c.encode() in b for c in canaries), n
                assert b"Bearer " not in b, n


def test_export_contains_every_trace_artifact(traced):
    svc, cid, _, bodies, _, _ = traced
    res = svc.export(cid)
    z = zipfile.ZipFile(res["path"])
    names = set(z.namelist())
    for n in range(1, len(bodies) + 1):
        for f in ("engine_request", "wire_request", "wire_response_raw",
                  "normalized_response", "metadata"):
            assert f"model_io/call_{n:03d}/{f}.json" in names
    assert "tool_roundtrips/tool_001.json" in names
    for f in ("model_io.jsonl", "tool_roundtrips.jsonl",
              "MODEL_IO_TRACE.html", "MODEL_IO_MANIFEST.json",
              "MODEL_IO_CHECKSUMS.sha256"):
        assert f in names, f
    for line in z.read("MODEL_IO_CHECKSUMS.sha256").decode().splitlines():
        digest, name = line.split("  ", 1)
        assert io_trace.sha256_text(z.read(name)) == digest, name
    wire = json.loads(z.read("model_io/call_001/wire_request.json"))
    assert wire[0]["body_bytes"] == len(bodies[0])
    readme = z.read("README.html").decode()
    assert model_io.README_LINE in readme
    assert "full system prompts" not in json.loads(
        z.read("manifest.json"))["omissions"].__str__()
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(z.read("comparison.xlsx")))
    for sheet in ("Model Calls", "Request Summary", "Tool Round Trips"):
        rows = list(wb[sheet].iter_rows(values_only=True))
        assert len(rows) > 1, sheet
        assert all(len(str(c or "")) < 32_767 for r in rows for c in r)


def test_saved_trace_opens_without_any_model_call(traced, monkeypatch):
    svc, cid, _, _, _, _ = traced
    import backend.model_lab.adapters as adapters

    def boom(*a, **k):
        raise AssertionError("a provider was built while reading a trace")
    monkeypatch.setattr(adapters, "build_provider", boom)
    monkeypatch.setattr(svc.coord, "provider_factory", boom)
    tr = model_io.build(svc.coord, cid)
    first = tr["children"][0]["timeline"][0]
    detail = model_io.call_detail(svc.coord, cid, first["call_id"])
    assert detail["views"]["engine_request"]["request"]["messages"]


# ---- 4. Anthropic route: real SDK, credential-free wire record ------------

def test_anthropic_wire_is_recorded_without_the_key():
    from backend.llm.anthropic_provider import AnthropicProvider

    seen = []

    def handler(req: httpx.Request) -> httpx.Response:
        seen.append(req)
        return httpx.Response(200, json={
            "id": "msg_1", "type": "message", "role": "assistant",
            "model": "claude-opus-5", "stop_reason": "tool_use",
            "content": [
                {"type": "thinking", "thinking": "secret chain of thought",
                 "signature": "sig-abc"},
                {"type": "tool_use", "id": "toolu_1",
                 "name": "execute_analysis", "input": {"objective": "x"}}],
            "usage": {"input_tokens": 10, "output_tokens": 5}},
            headers={"request-id": "req_1", "set-cookie": "a=b"})

    p = AnthropicProvider(api_key=CANARY_KEY)
    rec = io_trace.attach_wire_recorder(p)
    rec.inner = httpx.MockTransport(handler)
    r = p.converse(system="S", messages=[{"role": "user", "content": "Q"}],
                   tools=TOOLS, max_tokens=256, model="claude-opus-5",
                   allow_retry=False)
    assert r.tool_calls[0]["id"] == "toolu_1"
    assert seen[0].headers["x-api-key"] == CANARY_KEY      # it was sent
    view = io_trace.exchange_view(rec.exchanges[0],
                                  io_trace.secret_values(
                                      {"ANTHROPIC_KEY": CANARY_KEY}))
    text = json.dumps(view)
    assert CANARY_KEY not in text and "x-api-key" not in text
    assert "secret chain of thought" not in text and "sig-abc" not in text
    assert view["request"]["url"] == "https://api.anthropic.com/v1/messages"
    assert view["request"]["body"]["model"] == "claude-opus-5"
    assert view["response"]["body"]["content"][0]["thinking"]["reason"] == \
        io_trace.HIDDEN
    assert view["response"]["headers"]["request-id"] == "req_1"
    assert "set-cookie" not in view["response"]["headers"]


# ---- 5. failures are traced ----------------------------------------------

def test_http_error_and_timeout_are_traced(tmp_path):
    from backend.model_lab.observe import observe

    def run_one(handler):
        p = OpenAICompatProvider(base_url="http://127.0.0.1:9/v1", model="m",
                                 transport=httpx.MockTransport(handler))
        rec = io_trace.attach_wire_recorder(p)
        traces: list[dict] = []
        o = observe(p, lambda r: None, trace_sink=traces.append,
                    recorder=rec)
        with pytest.raises(ProviderFailure):
            o.converse(system="S", messages=[{"role": "user",
                                              "content": "Q"}],
                       tools=TOOLS, max_tokens=64, model="m")
        return traces[0]

    t = run_one(lambda req: httpx.Response(400, json={
        "error": "unknown field reasoning_effort"}))
    assert t["dispatch_status"] == "SENT"
    assert t["error"]["code"] == "PROVIDER_REQUEST_INVALID"
    resp = t["wire_exchanges"][0]["response"]
    assert resp["status_code"] == 400 and "unknown field" in \
        json.dumps(resp["body"])
    assert t["engine_request"]["messages"][0]["content"] == "Q"

    def slow(req):
        raise httpx.ReadTimeout("slow", request=req)
    t = run_one(slow)
    assert t["dispatch_status"] == "SENT"
    assert "ReadTimeout" in t["wire_exchanges"][0]["response"][
        "transport_error"]
    assert t["error"]["retry_class"] == "transport"


def test_a_call_refused_before_dispatch_is_not_sent(tmp_path):
    base = registry.load_profiles()["fixture-reference"]
    small = registry._validate(base.raw | {"profile_id": "fx-small",
                                           "context_tokens": 4_000},
                               Path("fx-small.json"))
    svc = make_service(tmp_path, profiles=registry.load_profiles()
                       | {"fx-small": small})
    cid, ev = run(svc, ["fx-small"], comparator="")
    tr = model_io.build(svc.coord, cid)
    items = tr["children"][0]["timeline"]
    assert items and items[0]["dispatch_status"] == "NOT_SENT"
    assert items[0]["dispatch_reason"] == "INPUT_CONTEXT_LIMIT"
    assert tr["children"][0]["traced_calls"] == 0


# ---- 6. offline multi-tool fixture: the trace summary ---------------------

def test_multi_tool_fixture_trace_summary(tmp_path, capsys):
    svc = make_service(tmp_path)
    cid, _ = run(svc, ["fixture-repair"], comparator="")
    tr = model_io.build(svc.coord, cid, include_bodies=True)
    items = tr["children"][0]["timeline"]
    calls = [i for i in items if i["kind"] == "call"]
    tools = [i for i in items if i["kind"] == "tool"]
    assert len(calls) == 4 and len(tools) == 4
    with capsys.disabled():
        print("\nMODEL I/O TRACE SUMMARY (fixture-repair, offline)")
        for it in items:
            if it["kind"] == "call":
                print(f"CALL {it['n']}: stage {'+'.join(it['stage_tags'])} "
                      f"engine request "
                      f"{it['sizes']['engine_request']['bytes']} B, wire "
                      f"{it['dispatch_status']}, response "
                      f"{it['sizes']['normalized_response']['bytes']} B, "
                      f"tool calls {[t['name'] for t in it['tool_calls']]}")
            else:
                print(f"  TOOL {it['n']}: {it['model_tool_call']['name']} "
                      f"validation "
                      f"{[e['status'] for e in it['validation']]} result "
                      f"{it['tool_result_bytes']} B")
    assert QUESTION in json.dumps(calls[0]["views"]["engine_request"])
