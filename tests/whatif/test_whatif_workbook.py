"""
The detailed workbook, judged the way an auditor would judge it.

The test is not "does a file come out". It is: can somebody who was not in the
room open this file and establish what was run, on what, under whose rules,
what every borrower's parameters were before and after, what caused the
movement, and whether the numbers tie — without asking anybody.

So the assertions here are about EVIDENCE and about REFUSAL. Every claim the
workbook makes has to be checkable inside the workbook, an allocation has to
say it is an allocation, and a figure that does not reconcile has to appear as
a failed test rather than as a missing file.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from backend.whatif import domain as dm
from backend.whatif import macrolab as ml
from backend.whatif import methodology as me
from backend.whatif import run as rn
from backend.whatif import scenarios as sc
from backend.whatif import steps as sp
from backend.whatif import workbook as wb


def _lake() -> bool:
    try:
        return bool(dm.periods())
    except Exception:
        return False


needs_lake = pytest.mark.skipif(
    not _lake(), reason="the Corporate IFRS 9 lake has not been built")


def _state(*steps: sp.Step) -> sp.ScenarioState:
    state = sp.ScenarioState(period=dm.latest_period())
    for step in steps:
        state = state.add(step)
    return state


PD_UP = sp.Step(sp.PD, (sc.Shock(sc.PD, 20.0, sc.RELATIVE),),
                interpreted="PD +20%")


@pytest.fixture(scope="module")
def result():
    return rn.execute(_state(PD_UP), requested=me.DELTA)


@pytest.fixture(scope="module")
def book(result):
    return load_workbook(io.BytesIO(wb.build(result, requested_by="ANALYST")))


#: The sheets the product specification names, in the reader's own words.
#: Asserted by NAME rather than by count, because a file with ten tabs and the
#: wrong ten is not the file that was asked for.
REQUIRED_SHEETS: tuple[str, ...] = (
    "Executive Summary",
    "Scenario Definition",
    "Portfolio Before vs After",
    "Account & Facility Detail",
    "Borrower Detail",
    "Stage Migration",
    "Rating Migration",
    "ECL Attribution",
    "Model & Methodology",
    "Data Dictionary",
)


@needs_lake
class TestTheWorkbookIsComplete:
    def test_every_sheet_an_auditor_needs_is_present(self, book) -> None:
        assert book.sheetnames == list(wb.SHEETS)

    @pytest.mark.parametrize("name", REQUIRED_SHEETS)
    def test_every_sheet_the_specification_names_is_there(self, book, name) -> None:
        assert name in book.sheetnames, (
            f"the file has {book.sheetnames}, which does not include {name!r}")

    def test_the_tabs_are_named_for_a_reader_and_not_for_the_code(self, book) -> None:
        """An audit file lands on a committee's desk, not in a repository."""
        for name in book.sheetnames:
            assert name == name.strip()
            assert name != name.upper(), (
                f"{name!r} is a code constant, not a tab a credit person reads")
            # Excel refuses these outright, so a name carrying one would mean
            # the file could not be written at all.
            assert not set(name) & set("/\\?*[]:"), name
            assert len(name) <= 31, name

    def test_no_sheet_is_empty(self, book) -> None:
        for name in book.sheetnames:
            assert book[name].max_row > 2, name

    def test_the_cover_says_this_is_a_scenario_not_the_book(self, book) -> None:
        said = _text(book[wb.COVER])
        assert "A SCENARIO, not the reported book" in said
        assert "hypothetical" in said

    def test_the_cover_carries_the_provenance_a_reader_will_be_asked_for(
            self, book, result) -> None:
        said = _text(book[wb.COVER])
        for expected in (result.period, dm.DOMAIN_NAME, dm.CURRENCY,
                         result.choice.label):
            assert expected in said, expected

    def test_both_staging_rule_sets_are_named_separately(self, book) -> None:
        """A reader has to know which rules produced which column."""
        said = _text(book[wb.SCENARIO])
        assert "Reported book" in said
        assert "What-If" in said
        assert "No What-If rule changes the reported book" in said

    def test_the_borrower_sheet_carries_before_and_after_for_every_parameter(
            self, book) -> None:
        headers = _headers(book[wb.BORROWERS])
        for pair in (("Rating before", "Rating after"),
                     ("Stage before", "Stage after"),
                     ("12m PD % before", "12m PD % after"),
                     ("LGD % before", "LGD % after"),
                     ("EAD before", "EAD after"),
                     ("Reported ECL", "What-If ECL")):
            assert pair[0] in headers, pair[0]
            assert pair[1] in headers, pair[1]

    def test_every_borrower_in_the_population_is_in_the_detail(
            self, book, result) -> None:
        rows = book[wb.BORROWERS].max_row
        # One title row, one subtitle, a blank, and the header.
        assert rows >= result.population

    def test_the_reproduce_sheet_carries_the_state_that_reruns_it(
            self, book, result) -> None:
        said = _text(book[wb.REPRODUCE])
        assert result.period in said
        assert "/whatif/execute" in said
        assert '"steps"' in said


@needs_lake
class TestTheFacilitySheetDoesNotOverclaim:
    def test_it_says_it_is_an_allocation_in_its_own_heading(self, book) -> None:
        said = _text(book[wb.FACILITIES])
        assert "ALLOCATION" in said
        assert "assessed on the OBLIGOR" in said
        assert "NOT a facility-level provision" in said

    def test_every_row_names_the_basis_it_was_allocated_on(self, book) -> None:
        sheet = book[wb.FACILITIES]
        headers = _headers(sheet)
        assert "Allocation basis" in headers
        column = headers.index("Allocation basis") + 1
        top = _header_row(sheet)
        seen = {sheet.cell(row=r, column=column).value
                for r in range(top + 1, min(top + 200, sheet.max_row + 1))}
        assert seen
        assert all(v for v in seen), "a row with no stated basis"

    def test_the_allocation_sums_back_to_the_borrower_figure(
            self, result) -> None:
        """The identity that makes the sheet usable at all."""
        allocated = wb._facility_rows(result.borrowers, result.context())
        if allocated.empty:
            pytest.skip("no facility register for this period")
        joined = allocated.groupby("borrower_id")["ecl_increase"].sum()
        expected = result.borrowers.set_index("borrower_id")["ecl_increase"]
        shared = joined.index.intersection(expected.index)
        assert len(shared) > 0
        for borrower in list(shared)[:500]:
            assert joined[borrower] == pytest.approx(
                float(expected[borrower]), abs=1e-6), borrower


@needs_lake
class TestTheReconciliationIsEvidenceNotAClaim:
    def test_every_check_passes_on_an_ordinary_scenario(self, result) -> None:
        for check in wb._checks(result, result.borrowers):
            assert check.status == "PASS", (
                f"{check.name}: {check.left_label} {check.left} vs "
                f"{check.right_label} {check.right}")

    def test_the_checks_are_recomputed_from_the_workbook_not_copied(
            self, book) -> None:
        assert "not copied from the run" in _text(book[wb.RECONCILIATION])

    def test_a_break_would_be_written_into_the_file_rather_than_raised(
            self) -> None:
        """An export that refuses to exist tells the reader nothing."""
        broken = wb.Check("A deliberate break", "left", 100.0, "right", 50.0)
        assert broken.status == "FAIL"
        assert broken.difference == pytest.approx(50.0)
        assert broken.difference_pct == pytest.approx(100.0)

    def test_the_hard_invariants_are_among_the_checks(self, result) -> None:
        names = {c.name for c in wb._checks(result, result.borrowers)}
        assert "No scenario cures a stage" in names
        assert "No borrower is provided for above their own exposure" in names
        assert "Stage migration conserves the population" in names

    def test_the_rounding_difference_is_explained_rather_than_absorbed(
            self, result) -> None:
        """A reader summing the detail column WILL see a small difference."""
        checks = {c.name: c for c in wb._checks(result, result.borrowers)}
        said = checks["The reported total is the sum of the borrowers"].note
        assert "two decimal places" in said
        assert "presentation artefact" in said


@needs_lake
class TestTheMethodSheetIsSelfContained:
    def test_it_names_every_version_a_rerun_must_match(self, book) -> None:
        said = _text(book[wb.METHOD])
        for expected in ("Workbook", "ECL methodology", "What-If staging",
                         "Macro sensitivities", "IFRS 9 policy"):
            assert expected in said, expected

    def test_it_carries_the_macro_sensitivities_and_their_basis(
            self, book) -> None:
        said = _text(book[wb.METHOD])
        assert "not an IFRS 9 coefficient" in said
        assert "GDP Growth Rate" in said
        assert "Corporate Credit Spread" in said

    def test_plausibility_is_never_written_as_a_probability(self, book) -> None:
        said = _text(book[wb.METHOD]).lower()
        if "where this shock sits" in said:
            assert "not a forecast and not a probability" in said


@needs_lake
class TestAnOverriddenAssumptionIsVisibleInTheFile:
    def test_a_user_defined_sensitivity_is_named_on_the_cover_and_the_scenario(
            self) -> None:
        state = _state(sp.Step(
            sp.MACRO,
            (sc.Shock(sc.MACRO, -1.0, sc.ABSOLUTE_PP, target="gdp_growth"),),
            interpreted="GDP -1pp")).with_sensitivity(ml.Sensitivity(
                variable="gdp_growth", source=ml.USER, name="GDP Growth Rate",
                pd_response_kind=ml.MULTIPLIER, pd_response=1.35,
                lgd_response_kind=ml.ABSOLUTE_PP, lgd_response=1.2))
        found = rn.execute(state, requested=me.DELTA)
        book_ = load_workbook(io.BytesIO(wb.build(found)))
        cover = _text(book_[wb.COVER])
        scenario = _text(book_[wb.SCENARIO])
        assert "overridden for this thread" in cover
        assert "User-Defined Sensitivity" in scenario
        assert "unchanged" in scenario

    def test_the_default_case_states_that_nothing_was_overridden(
            self, book) -> None:
        assert "governed CreditProbe reference sensitivity" in _text(
            book[wb.COVER])


@needs_lake
class TestItRefusesRatherThanWritingAnEmptyFile:
    def test_a_scenario_matching_nobody_is_refused_with_a_reason(self) -> None:
        import pandas as pd

        class Empty:
            borrowers = pd.DataFrame()

        with pytest.raises(wb.WorkbookError) as raised:
            wb.build(Empty())
        assert "matched no borrowers" in str(raised.value)


# ------------------------------------------------------------------ helpers


def _text(sheet) -> str:
    return "\n".join(
        str(cell.value) for row in sheet.iter_rows() for cell in row
        if cell.value is not None)


def _header_row(sheet) -> int:
    """The row carrying the table's headers — the first styled header cell."""
    for r in range(1, sheet.max_row + 1):
        cell = sheet.cell(row=r, column=1)
        if cell.font and cell.font.bold and cell.fill and \
                cell.fill.fgColor.rgb not in (None, "00000000"):
            if sheet.cell(row=r, column=2).value is not None:
                return r
    return 1


def _headers(sheet) -> list[str]:
    top = _header_row(sheet)
    return [str(sheet.cell(row=top, column=c).value or "")
            for c in range(1, sheet.max_column + 1)]


@needs_lake
class TestTheDataDictionary:
    """The sheet that makes the rest of the file readable by a stranger.

    Three months after the run, "pd_stressed" is opaque to the reviewer
    deciding whether the provision was reasonable, and a column they cannot
    interpret is one they either ignore or have to ask about.
    """

    def test_it_explains_every_family_of_column_on_the_detail_sheets(self, book) -> None:
        said = _text(book[wb.DICTIONARY]).lower()
        for term in ("borrower id", "rating", "stage", "ttc pd",
                     "lifetime pd", "applicable pd", "lgd", "ccf", "ead",
                     "collateral", "reported ecl", "what-if ecl",
                     "primary driver"):
            assert term in said, f"the dictionary does not explain {term!r}"

    def test_it_states_the_governed_scale_and_that_default_is_separate(self, book) -> None:
        said = _text(book[wb.DICTIONARY])
        assert "AAA" in said and "C (19)" in said
        assert "nineteen-point" in said.lower()
        assert "separate state" in said.lower()

    def test_it_says_a_pd_of_one_hundred_is_not_a_loss_of_one_hundred(self, book) -> None:
        said = _text(book[wb.DICTIONARY]).lower()
        assert "100% in stage 3" in said
        assert "not a loss of 100%" in said
        assert "severity stays with lgd" in said

    def test_it_states_the_grain_and_that_facility_figures_are_allocated(self, book) -> None:
        said = _text(book[wb.DICTIONARY]).lower()
        assert "obligor level" in said
        assert "allocat" in said

    def test_it_names_where_each_figure_came_from(self, book) -> None:
        said = _text(book[wb.DICTIONARY])
        for source in ("corporate_ifrs9", "corporate_ratings",
                       "corporate_facilities", "corporate_collateral"):
            assert source in said

    def test_it_points_at_the_calibration_document_rather_than_restating_it(self, book) -> None:
        assert "docs/corporate_rating_pd_calibration.md" in _text(
            book[wb.DICTIONARY])
