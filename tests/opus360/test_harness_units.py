"""Unit tests for the Opus360 harness. No provider, no paid call."""

from __future__ import annotations

import csv
import gzip
import hashlib
import json
import time
from pathlib import Path

import pytest
from cert import evaluate as ev
from cert import oracles as orc
from cert import protected, telemetry
from cert.ledger import Experiment
from cert.observe import UNKNOWN, ObservingProvider, Recorder
from cert.paths import BANK_PATH, BANK_SHA_PATH
from cert.runner import Guards, Price, StopRun
from cert.safe_io import append_jsonl, cell, read_jsonl, redact, rows_to_xlsx, write_csv

# ---- question bank -------------------------------------------------------------------

def test_bank_has_exactly_250_core_turns(bank):
    rows = bank["rows"]
    assert len(rows) == 250 == bank["core_turns"]
    counts = {}
    for r in rows:
        counts[r["test_class"]] = counts.get(r["test_class"], 0) + 1
    assert counts == {"A_BASELINE": 50, "B_PARAPHRASE": 50, "C_COMPLEX": 40, "D_THREAD": 50,
                      "E_CLARIFICATION": 20, "F_GOVERNANCE": 20, "G_REPEAT": 20}
    assert len({r["case_id"] for r in rows}) == 250


def test_bank_rows_have_every_required_field(bank):
    required = {"case_id", "family_id", "domain", "question", "test_class", "thread_id", "turn_number",
                "fresh_thread", "oracle_type", "expected_behavior", "difficulty", "paraphrase_group",
                "repeatability_group", "tags", "notes"}
    for r in bank["rows"]:
        assert required <= set(r), r["case_id"]
        assert r["domain"] in ("corporate", "retail")


def test_bank_group_structure(bank):
    rows = bank["rows"]
    a = [r for r in rows if r["test_class"] == "A_BASELINE"]
    assert sum(r["domain"] == "corporate" for r in a) == 25
    c = [r for r in rows if r["test_class"] == "C_COMPLEX"]
    assert sum(r["domain"] == "corporate" for r in c) == 20
    threads = {}
    for r in rows:
        if r["test_class"] == "D_THREAD":
            threads.setdefault(r["thread_id"], []).append(r["turn_number"])
    assert len(threads) == 10 and all(sorted(t) == [1, 2, 3, 4, 5] for t in threads.values())
    e = [r for r in rows if r["test_class"] == "E_CLARIFICATION"]
    assert {r["clarification_class"] for r in e} == {"CLARIFICATION_REQUIRED", "CLARIFICATION_AVOIDABLE"}
    assert all(r.get("clarification_followup") for r in e)
    b = [r for r in rows if r["test_class"] == "B_PARAPHRASE"]
    assert len({r["paraphrase_group"] for r in b}) == 25
    g = [r for r in rows if r["test_class"] == "G_REPEAT"]
    assert len({r["repeatability_group"] for r in g}) == 10
    switching = {r["thread_id"] for r in rows if (r.get("transport") or {}).get("on_domain_pinned")}
    assert len(switching) >= 2


def test_bank_hash_is_frozen():
    stored = BANK_SHA_PATH.read_text().split()[0]
    assert hashlib.sha256(BANK_PATH.read_bytes()).hexdigest() == stored


def test_every_oracle_in_the_bank_computes(bank):
    for r in bank["rows"]:
        ref = orc.compute(r["oracle"], r["domain"])
        assert ref.domain == r["domain"]
        if r.get("followup_oracle"):
            orc.compute(r["followup_oracle"], r["domain"])


# ---- protected core ------------------------------------------------------------------------

def test_protected_manifest_verifies():
    v = protected.verify()
    assert v.ok, v.problems
    assert v.checked_files >= 300


def test_protected_detects_a_changed_manifest(tmp_path):
    from cert.paths import PROTECTED_MANIFEST

    data = json.loads(PROTECTED_MANIFEST.read_text())
    first = sorted(data["files"])[0]
    data["files"][first]["sha256"] = "0" * 64
    fake = tmp_path / "m.json"
    fake.write_text(json.dumps(data))
    v = protected.verify(fake)
    assert not v.ok and any(first in p for p in v.problems)


def test_protected_detects_a_swapped_manifest_file():
    v = protected.verify(expected_manifest_sha="f" * 64)
    assert not v.ok and any("manifest itself changed" in p for p in v.problems)


def test_protected_refuses_to_rebuild_over_drift(monkeypatch):
    monkeypatch.setattr(protected, "worktree_blobs", lambda paths: {p: "deadbeef" for p in paths})
    with pytest.raises(RuntimeError, match="refusing"):
        protected.build(Path("/nonexistent/should-not-write.json"))


def test_runtime_closure_is_covered_by_the_manifest():
    """Every repository module a real scripted run loads (clean interpreter, no test
    conftest) is in the protected manifest."""
    import subprocess
    import sys

    from cert.paths import PROTECTED_MANIFEST, ROOT

    out = subprocess.run([sys.executable, str(ROOT / "scripts/opus360/runtime_closure.py")], cwd=ROOT,
                         capture_output=True, text=True, timeout=600)
    loaded = json.loads(out.stdout)
    covered = set(json.loads(PROTECTED_MANIFEST.read_text())["files"])
    unprotected = [p for p in loaded if p.startswith("backend/") and p not in covered]
    assert loaded and not unprotected, unprotected


# ---- observation is pass-through ------------------------------------------------------------

class _Inner:
    def __init__(self):
        self.result = object()
        self.kwargs = None

    def converse(self, **kwargs):
        self.kwargs = kwargs
        return self.result

    def count_tokens(self, **kwargs):
        return 42

    something_else = "passthrough"


def test_observing_provider_returns_the_identical_object(tmp_path):
    inner, rec = _Inner(), Recorder(tmp_path)
    obs = ObservingProvider(inner, rec)
    args = {"system": [{"type": "text", "text": "s"}], "messages": [{"role": "user", "content": "q"}],
            "tools": [{"name": "t"}], "max_tokens": 10, "model": "m", "purpose": "p", "allow_retry": False}
    assert obs.converse(**args) is inner.result
    assert inner.kwargs == args
    assert obs.count_tokens(model="m") == 42
    assert obs.something_else == "passthrough"
    calls = rec.calls()
    assert [c.kind for c in calls] == ["converse", "count_tokens"]
    assert calls[0].input_tokens == UNKNOWN          # never a fabricated zero
    assert Path(calls[0].payload_path).exists()


def test_observing_provider_reraises_the_same_exception(tmp_path):
    class Boom(Exception):
        pass

    class Inner:
        def converse(self, **kw):
            raise Boom("429 rate limit sk-ant-abcdefghijklmnop")

    rec = Recorder(tmp_path)
    with pytest.raises(Boom):
        ObservingProvider(Inner(), rec).converse(messages=[])
    c = rec.calls()[0]
    assert c.status == "error" and "sk-ant" not in c.error_text and "[REDACTED]" in c.error_text


def test_sdk_and_http_shims_are_pass_through(tmp_path):
    import httpx
    from cert.observe import install_sdk_shims

    rec = Recorder(tmp_path)
    install_sdk_shims(rec)
    transport = httpx.MockTransport(lambda request: httpx.Response(200, json={"ok": True},
                                                                   headers={"request-id": "req_1"}))
    client = httpx.Client(transport=transport)
    active = rec.begin("converse")
    response = client.get("https://api.anthropic.com/v1/messages")
    other = client.get("https://example.com/")
    rec.end(active)
    assert response.json() == {"ok": True} and other.status_code == 200
    assert len(active.http_attempts) == 1 and active.http_attempts[0]["status"] == 200
    assert active.http_attempts[0]["request_id"] == "req_1"


# ---- evaluation ---------------------------------------------------------------------------------

def test_compare_tiers_and_scales():
    assert ev.compare(1016613.64, 1016613.64, "amount") == "EXACT"
    assert ev.compare(1016613.64, 1016614.0, "amount") == "ROUNDED"
    assert ev.compare(1016613.64, 1016.61364, "amount") == "EXACT"     # SAR billion
    assert ev.compare(0.2419, 24.19, "ratio") == "EXACT"                # percent
    assert ev.compare(0.2419, 24.2, "ratio") == "ROUNDED"
    assert ev.compare(0.2419, 0.30, "ratio") == ""
    assert ev.compare(None, 1.0, "amount") == "" and ev.compare(1.0, None, "amount") == ""
    assert ev.compare(12, 12, "count") == "EXACT" and ev.compare(12, 14, "count") == ""


def _final_with(rows, claims=None):
    return {"numeric_claims": claims or [], "tables": [{"artifact_id": "a1", "rows": [
        {"row_id": f"r{i}", "canonical": r} for i, r in enumerate(rows)]}], "narrative": ""}


def test_oracle_exact_match_and_mismatch():
    ref = orc.compute({"fn": "by_dim", "dim": "sector", "metrics": ["ead_s2"], "rank_by": "ead_s2", "topn": 2},
                      "corporate")
    good = [{"sector": k["sector"], "ead_s2": f.value} for k in ref.ranking
            for f in ref.required() if f.key == k]
    pub = ev.published_cells(_final_with(good), [])
    res = ev.match_required(ref, pub, [])
    assert all(r.status == "PUBLISHED_EXACT" for r in res)
    bad = [dict(r, ead_s2=r["ead_s2"] * 1.5) for r in good]
    res = ev.match_required(ref, ev.published_cells(_final_with(bad), []), [])
    assert all(r.status == "VALUE_MISMATCH" for r in res)
    rk = ev.check_ranking(ref, _final_with(list(reversed(good))), "")
    assert rk["required"] and not rk["ok"] and rk["set_ok"]


def test_wrong_period_is_named():
    ref = orc.compute({"fn": "scalar", "metrics": ["ead_s2"]}, "corporate")
    prev = orc.compute({"fn": "scalar", "metrics": ["ead_s2"], "period": "prev"}, "corporate").required()[0].value
    res = ev.match_required(ref, ev.published_cells(_final_with([{"ead_s2": prev}]), []), [])
    assert res[0].status == "WRONG_PERIOD"


def test_unsupported_numbers_in_narrative_are_detected():
    text = ("Stage 2 EAD is {{claim.a}}. It rose by 12.7% versus 2026Q1, and the top 3 sectors hold "
            "SAR 2,999,999 million. Stage 3 and 90+ days past due are shown in 2026.")
    found = [u["raw"] for u in ev.unbound_numbers(text)]
    assert "12.7%" in found and any("2,999,999" in f for f in found)
    assert not any(f.strip() in ("2026", "3", "90", "2") for f in found)


def test_causal_heuristic():
    assert ev.causal_check("ECL rose because oil prices fell.")["unhedged_causality"]
    assert not ev.causal_check("ECL rose; the data cannot establish why, because no macro field exists.")["unhedged_causality"]


def test_root_cause_provider_and_budget_rules():
    base = {"case_id": "x", "domain": "corporate", "turn_number": 1}
    v = ev.Verdict("x", "ANSWER", "FAILED", False, failure_reasons=["f"])
    rc = ev.root_cause(base, v, {"record": {"error_code": "PROVIDER_UNAVAILABLE"}}, {}, [], [])
    assert rc["primary"] == "PROVIDER_TRANSIENT"
    rc = ev.root_cause(base, v, {"record": {"error_code": "COST_LIMIT"}}, {}, [], [])
    assert rc["primary"] == "ARCH_BUDGET"
    rc = ev.root_cause(base, v, {"record": {"error_code": "EXECUTION_LIMIT"}},
                       {"any_success": False, "submissions": 5, "statuses": ["rejected"] * 5}, [], [])
    assert rc["primary"] == "MODEL_REPAIR"
    ok = ev.Verdict("x", "ANSWER", "ANSWER", True, passed=True)
    assert ev.root_cause(base, ok, {}, {}, [], [])["primary"] == "PASS"


def test_validator_replay_detects_rejected_correct_sql():
    ref = orc.compute({"fn": "scalar", "metrics": ["ead"]}, "corporate")
    sql = "SELECT SUM(ead_sar_mn) AS ead FROM corp_facility_quarter WHERE reporting_quarter = '2026Q2'"
    subs = [{"submission_id": "s1", "status": "rejected", "payload": {"steps": [{"language": "sql", "code": sql}]}}]
    out = ev.validator_replay(ref, subs, [], {})
    assert out and out[0]["would_have_been_correct"] is True
    subs[0]["payload"]["steps"][0]["code"] = "DROP TABLE corp_facility_quarter"
    assert ev.validator_replay(ref, subs, [], {})[0]["replayed"] is False


# ---- telemetry ---------------------------------------------------------------------------------

def test_token_decomposition_splits_history_and_tool_results(tmp_path):
    first = ("USER REQUEST (original wording, unmodified):\nOnly the top three.\n\n"
             "RECENT COMPLETED TURNS IN THIS THREAD (exact records; these outrank any summary):\n"
             + json.dumps([{"turn_id": "t", "question": "Stage 2 by sector?", "answer": "A long answer " * 20}])
             + "\n\nDecide what this request is, who owns it, and take your next action now.")
    payload = {"system": [{"type": "text", "text": "I" * 400}, {"type": "text", "text": "K" * 800},
                          {"type": "text", "text": "V" * 100}],
               "tools": [{"name": "x", "input_schema": {"a": "b" * 300}}],
               "messages": [{"role": "user", "content": first},
                            {"role": "assistant", "content": [{"type": "tool_use", "id": "u1", "name": "execute_analysis", "input": {}}]},
                            {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "u1", "content": "R" * 500}]}]}
    parts = telemetry.decompose_payload(payload)
    assert parts["history_user"] > 0 and parts["history_assistant"] > parts["history_user"]
    assert parts["current_question"] > 0 and parts["in_run_tool_results_current"] >= 500
    assert parts["system_static_knowledge"] == 800 and parts["tool_schema"] > 300
    path = tmp_path / "p.json.gz"
    with gzip.open(path, "wt") as fh:
        json.dump(payload, fh)
    d = telemetry.decompose_call({"payload_path": str(path), "input_tokens": 1000})
    assert d["method"] == "ESTIMATED_BYTE_SHARE" and abs(d["estimated_sum"] - 1000) < 1
    unk = telemetry.decompose_call({"payload_path": str(path), "input_tokens": UNKNOWN})
    assert all(v == UNKNOWN for v in unk["tokens"].values())


def test_token_summary_never_fabricates_zero():
    calls = [{"kind": "converse", "input_tokens": UNKNOWN, "output_tokens": UNKNOWN, "started_at": ""}]
    s = telemetry.token_summary(calls, [], None, None)
    assert s["native_input_tokens"] == UNKNOWN and s["native_total_tokens"] == UNKNOWN


# ---- cost and guards ------------------------------------------------------------------------------

def test_cost_accounting_matches_the_price_card():
    card = json.loads((Path(__file__).resolve().parents[2] / "config/cockpit_v4/price_card.claude-opus-5.json").read_text())
    p = card["models"]["claude-opus-5"]["price"]
    price = Price(p["input_usd_per_mtok"], p["output_usd_per_mtok"], p["cache_write_usd_per_mtok"], p["cache_read_usd_per_mtok"])
    assert price.cost(1200, 300) == pytest.approx(0.0135)
    assert price.cost(1_000_000, 0) == pytest.approx(5.0)
    assert price.cost(UNKNOWN, 1) is None


def test_spend_cap_stops_before_it_could_be_exceeded():
    g = Guards(live=True, max_usd=10.0, max_hours=None)
    g.check(8.49)
    with pytest.raises(StopRun) as e:
        g.check(8.51)
    assert e.value.condition == "NOT_RUN_SPEND_CAP"
    with pytest.raises(StopRun) as e:
        Guards(live=True, max_usd=None, max_hours=None).check(0)
    assert e.value.condition == "NO_SPEND_CAP"


def test_wall_clock_cap_and_stop_file(tmp_path):
    g = Guards(live=False, max_usd=None, max_hours=0.05, started_mono=time.monotonic() - 3600)
    with pytest.raises(StopRun) as e:
        g.check(0)
    assert e.value.condition == "NOT_RUN_TIME_CAP"
    stop = tmp_path / "STOP"
    stop.write_text("x")
    with pytest.raises(StopRun) as e:
        Guards(live=False, max_usd=None, max_hours=None, stop_file=stop).check(0)
    assert e.value.condition == "STOPPED_BY_OPERATOR"


# ---- ledger: resume, duplicates, immutability, checksums ---------------------------------------------

def test_experiment_is_immutable_and_resume_verifies_checksums(tmp_path):
    exp = Experiment(tmp_path / "e")
    exp.create({"experiment_id": "e"})
    with pytest.raises(FileExistsError):
        Experiment(tmp_path / "e").create({"experiment_id": "e"})
    d, sha = exp.write_evidence("core", "A-1", 1, {"x": {"v": 1}})
    exp.record_result({"record_type": "turn", "phase": "core", "case_id": "A-1", "evidence_dir": str(d),
                       "evidence_sha256": sha})
    d2, sha2 = exp.write_evidence("core", "A-2", 1, {"x": {"v": 2}})
    exp.record_result({"record_type": "turn", "phase": "core", "case_id": "A-2", "evidence_dir": str(d2),
                       "evidence_sha256": sha2})
    assert set(exp.completed("core")) == {"A-1", "A-2"}
    (d2 / "x.json").write_text('{"v": 999}')           # tamper
    assert set(exp.completed("core")) == {"A-1"}         # re-run, never trusted
    with open(exp.results_path, "a") as fh:            # torn final line after a crash
        fh.write('{"record_type": "turn", "case_id": "A-3"')
    assert len(read_jsonl(exp.results_path)) == 2
    assert exp.write_checksums().exists()


def test_runner_skips_completed_turns(harness, tmp_path):
    """Duplicate prevention: a verified completed turn is not run again on resume."""
    from cert import scripted
    from cert.runner import Runner

    row = {"case_id": "DUP-1", "family_id": "d", "domain": "corporate", "question": "What is total EAD?",
           "test_class": "A_BASELINE", "thread_id": "T-DUP-1", "turn_number": 1, "fresh_thread": True,
           "oracle_type": "EXACT", "expected_behavior": "ANSWER", "acceptable_behaviours": ["ANSWER"],
           "oracle": {"fn": "scalar", "metrics": ["ead"]}, "tags": []}
    refs = {"DUP-1:primary": orc.compute(row["oracle"], "corporate")}
    exp = Experiment(tmp_path / "e")
    exp.create({"experiment_id": "e"})

    def make():
        return Runner(exp, harness, harness.recorder, bank_rows=[row], refs=refs, price=Price(5, 25, 6.25, .5),
                      guards=Guards(live=False, max_usd=None, max_hours=None), live=False, echo=False,
                      script_factory=lambda c, r, d, b, a: scripted.analyst_for(c, r, domain=d, behaviour=b),
                      protected_manifest_sha=protected.verify().manifest_sha256)

    make().run_phase("core", [row])
    make().run_phase("core", [row])
    turns = [r for r in exp.results() if r.get("record_type") == "turn"]
    assert len(turns) == 1 and turns[0]["passed"]


# ---- safe IO -------------------------------------------------------------------------------------

def test_formula_injection_is_neutralised(tmp_path):
    assert cell("=HYPERLINK(\"x\")") == "'=HYPERLINK(\"x\")"
    assert cell("+cmd") == "'+cmd" and cell("@SUM(A1)") == "'@SUM(A1)"
    assert cell("-1.5") == "-1.5" and cell(-1.5) == -1.5
    write_csv(tmp_path / "x.csv", [{"q": "=1+1", "n": 3}])
    rows = list(csv.DictReader(open(tmp_path / "x.csv")))
    assert rows[0]["q"] == "'=1+1"
    rows_to_xlsx(tmp_path / "x.xlsx", [("S", [{"q": "=1+1"}])])
    from openpyxl import load_workbook

    ws = load_workbook(tmp_path / "x.xlsx")["S"]
    assert ws["A2"].value == "'=1+1" and ws["A2"].data_type == "s"


def test_secrets_are_redacted_everywhere(tmp_path):
    secret = "sk-ant-api03-" + "A" * 40
    assert secret not in redact(f"key={secret}")
    assert "Bearer abcdefghijkl" not in redact("Authorization: Bearer abcdefghijkl")
    append_jsonl(tmp_path / "e.jsonl", {"detail": f"x-api-key: {secret}"})
    write_csv(tmp_path / "e.csv", [{"detail": secret}])
    for f in ("e.jsonl", "e.csv"):
        assert secret not in (tmp_path / f).read_text()


def test_null_values_survive_every_writer(tmp_path):
    write_csv(tmp_path / "n.csv", [{"a": None, "b": UNKNOWN, "c": 0}])
    row = list(csv.DictReader(open(tmp_path / "n.csv")))[0]
    assert row == {"a": "", "b": "UNKNOWN", "c": "0"}
