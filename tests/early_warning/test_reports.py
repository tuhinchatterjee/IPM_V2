"""Early Warning V2 Word reports (spec Sections AS/AT): every report is
generated from live data, never a static sample — and a multi-borrower
report gets one substantive section per borrower, not a summary table.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from backend.early_warning import reports
from backend.early_warning.v2_service import EarlyWarningDataNotBuilt, segment_summary, top_high_risk


def _skip_if_not_built():
    try:
        top_high_risk(limit=1)
    except EarlyWarningDataNotBuilt:
        pytest.skip("Early Warning V2 domain not built")


def test_borrower_report_is_a_real_docx_with_live_figures():
    _skip_if_not_built()
    row = top_high_risk(limit=1)[0]
    data = reports.generate_docx("borrower", row["customer_id"])
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert row["customer_name"] in text
    assert f"{row['ews_score']:.1f}"[:4] in text  # score appears, not a placeholder
    assert len(doc.tables) >= 1


def test_borrower_report_unknown_customer_raises_keyerror():
    _skip_if_not_built()
    with pytest.raises(KeyError):
        reports.generate_docx("borrower", "NOT-A-REAL-CUSTOMER")


def test_portfolio_report_sections_match_live_summary():
    _skip_if_not_built()
    data = reports.generate_docx("portfolio")
    doc = Document(BytesIO(data))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("Executive Summary" in h for h in headings)
    assert any("Portfolio Position" in h for h in headings)
    assert any("Watchlist" in h for h in headings)


def test_segment_report_only_includes_that_segment():
    _skip_if_not_built()
    seg = segment_summary()[0]["segment"]
    data = reports.generate_docx("segment", seg)
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert seg in text


def test_segment_report_unknown_segment_raises_keyerror():
    _skip_if_not_built()
    with pytest.raises(KeyError):
        reports.generate_docx("segment", "Not A Real Segment Name")


def test_multi_borrower_report_has_one_section_per_borrower_not_a_summary_table():
    """Spec Section AT: selecting several borrowers must not collapse to a
    tiny summary table — each gets its own substantive section."""
    _skip_if_not_built()
    rows = top_high_risk(limit=5)
    ids = [r["customer_id"] for r in rows]
    data = reports.generate_docx("multi_borrower", customer_ids=ids)
    doc = Document(BytesIO(data))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading") and p.style.name.endswith("1")]
    # one heading per borrower name, plus "Population Comparison" and "The Early Warning Model"
    for row in rows:
        assert any(row["customer_name"] in h for h in headings), row["customer_name"]
    assert len(doc.tables) >= len(rows)  # comparison table + one history table per borrower


def test_multi_borrower_report_requires_at_least_one_valid_customer():
    with pytest.raises(KeyError):
        reports.generate_docx("multi_borrower", customer_ids=["NOT-REAL-1", "NOT-REAL-2"])


def test_generate_docx_rejects_unknown_scope():
    with pytest.raises(ValueError):
        reports.generate_docx("not_a_real_scope")
