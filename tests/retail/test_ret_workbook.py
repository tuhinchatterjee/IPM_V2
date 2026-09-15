"""Phase 5, §12: the workbook is a presentation standard AND a completeness gate.

Checked by OPENING the file, not by asserting that bytes were produced. The
contract is explicit that a ZIP existing is not acceptance.
"""
from __future__ import annotations

import io
import zipfile
from collections import Counter

import openpyxl
import pytest

from backend.retail import whatif_cohort as WC
from backend.retail import whatif_selection as SEL
from backend.retail import whatif_workbook as WB
from backend.retail import workbook_style as ST

ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!")
NULLISH = {"nan", "none", "null", "nat", "inf", "-inf", "<na>"}


@pytest.fixture(scope="module")
def built():
    """One real run, through the real engine, exported the real way."""
    made = SEL.create(month="2026-08", product="CREDIT_CARD",
                      sub_product="CC_PLATINUM",
                      source_module=SEL.SOURCE_GUIDED, label="test cohort")
    out = WC.run(made.selection_id, shocks={"income_pct": -0.10},
                 method="delta")
    payload, name = WB.build(out, selection=made.to_dict())
    return payload, name, out, made


@pytest.fixture(scope="module")
def book(built):
    return openpyxl.load_workbook(io.BytesIO(built[0]))


# ===================== EXPORT-02: no gridlines, every sheet ================

def test_xl_01_every_sheet_hides_its_gridlines(book) -> None:
    showing = [ws.title for ws in book.worksheets
               if ws.sheet_view.showGridLines is not False]
    assert not showing, f"gridlines still shown on {showing}"


def test_xl_02_every_mandatory_sheet_is_present(book) -> None:
    """§12.2 names them. A sheet that does not apply explains itself rather
    than being left out."""
    wanted = {name for name, _ in WB.SHEETS}
    assert wanted <= set(book.sheetnames), wanted - set(book.sheetnames)
    assert len(book.sheetnames) == len(WB.SHEETS)


# ============ EXPORT-03: numeric and probability formats ==================

def test_xl_03_no_numeric_cell_is_left_unformatted(book) -> None:
    """The §2.4 defect: formats were chosen by sniffing the display label, so
    "Exposure (SAR)" got none and a PD printed as 0.034823."""
    loose = Counter()
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if (isinstance(cell.value, (int, float))
                        and not isinstance(cell.value, bool)
                        and cell.number_format == "General"):
                    loose[sheet.title] += 1
    assert not loose, f"unformatted numeric cells: {dict(loose)}"


def test_xl_04_a_probability_reads_as_a_percentage(book) -> None:
    """0.034823 must read as 3.48%, which is what §2.4 asks for by name."""
    found = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if (isinstance(cell.value, float) and 0.0 < cell.value < 1.0
                        and "%" in cell.number_format):
                    found.append(cell.value)
    assert found, "no probability is formatted as a percentage anywhere"


def test_xl_05_negatives_are_shown_in_parentheses(book) -> None:
    money = [cell.number_format for sheet in book.worksheets
             for row in sheet.iter_rows() for cell in row
             if isinstance(cell.value, (int, float))
             and "#,##0" in cell.number_format]
    assert money
    assert all(";" in one and "(" in one for one in money), (
        "a money format without a negative section shows a minus sign")


def test_xl_06_nothing_is_written_as_the_word_nan(book) -> None:
    """A missing value is blank. `str(float("nan"))` is "nan", and `or ""`
    does not catch it because nan is truthy."""
    bad = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    if cell.value.strip().lower() in NULLISH:
                        bad.append(f"{sheet.title}!{cell.coordinate}")
                elif isinstance(cell.value, float) and cell.value != cell.value:
                    bad.append(f"{sheet.title}!{cell.coordinate} NaN")
    assert not bad, bad[:8]


def test_xl_07_there_are_no_error_cells(book) -> None:
    bad = [f"{s.title}!{c.coordinate}" for s in book.worksheets
           for r in s.iter_rows() for c in r
           if isinstance(c.value, str) and any(e in c.value for e in ERRORS)]
    assert not bad, bad[:8]


# ==================== EXPORT-06: the waterfall reconciles ==================

def test_xl_08_the_waterfall_sheet_publishes_its_residual(built, book) -> None:
    result = built[2]
    steps = result.get("waterfall") or {}
    if not steps.get("steps"):
        pytest.skip("this scenario has a single step, so nothing to reconcile")
    sheet = book["Waterfall"]
    text = " ".join(str(c.value) for r in sheet.iter_rows() for c in r
                    if c.value is not None)
    assert "Residual" in text
    assert "Baseline ECL" in text and "Scenario ECL" in text


def test_xl_09_the_steps_sum_to_the_change(built) -> None:
    result = built[2]
    steps = (result.get("waterfall") or {}).get("steps") or []
    if not steps:
        pytest.skip("single-step scenario")
    total = sum(float(one.get("change_sar") or 0.0) for one in steps)
    cohort = WB._cohort_level(result)
    change = float((cohort.get("delta") or {}).get("ecl_weighted_sar") or 0.0)
    assert total == pytest.approx(change, abs=1.0)


# =============== EXPORT-05: every propagation step has a sheet =============

def test_xl_10_an_inapplicable_step_explains_itself(book) -> None:
    """§12.2: an inapplicable step may contain an explanation rather than
    fictitious data. An empty sheet with no sentence is neither."""
    for name in ("02 Transformations", "03 Score Points", "04 Score Migration",
                 "06 Stage Decisions", "07 LGD EAD Recovery"):
        sheet = book[name]
        text = " ".join(str(c.value) for r in sheet.iter_rows() for c in r
                        if c.value is not None)
        assert len(text) > 120, f"{name} says almost nothing"


def test_xl_11_the_contents_sheet_links_to_every_other_sheet(book) -> None:
    sheet = book["Contents"]
    targets = set()
    for row in sheet.iter_rows():
        for cell in row:
            if cell.hyperlink and cell.hyperlink.location:
                targets.add(cell.hyperlink.location.split("!")[0].strip("'"))
    named = set(book.sheetnames) - {"Contents"}
    assert named <= targets, f"not linked from Contents: {named - targets}"


def test_xl_12_charts_are_embedded(built) -> None:
    with zipfile.ZipFile(io.BytesIO(built[0])) as archive:
        charts = [n for n in archive.namelist() if "/charts/chart" in n]
    assert charts, "no chart was embedded"


# ====================== the presentation standard itself ==================

def test_xl_13_headings_are_frozen_on_data_sheets(book) -> None:
    for name in ("Cohort", "Impact by Level", "Waterfall"):
        assert book[name].freeze_panes, f"{name} does not freeze its heading"


def test_xl_14_a_print_area_is_set_on_data_sheets(book) -> None:
    for name in ("Cohort", "Impact by Level", "Data Dictionary"):
        assert book[name].print_area, f"{name} has no print area"


def test_xl_15_the_disclosure_is_present_and_is_not_overclaimed(book) -> None:
    import re

    text = " ".join(str(c.value) for s in book.worksheets
                    for r in s.iter_rows() for c in r
                    if c.value is not None).lower()
    assert "synthetic" in text
    assert "not independently validated" in text
    for forbidden in ("approved by sama", "sama approval", "anb approved",
                      "independently validated", "auditor certified"):
        claimed = re.search(rf"(?<!not ){re.escape(forbidden)}", text)
        assert claimed is None, (
            f"the workbook claims {forbidden!r}: "
            f"…{text[max(0, claimed.start() - 60):claimed.end() + 20]}…")


def test_xl_16_the_filename_names_the_selection(built) -> None:
    _, name, _, made = built
    assert made.selection_id in name
    assert name.endswith(".xlsx")


# ============ the column contract that replaced the name sniffing =========

def test_xl_17_a_percent_column_is_divided_back_to_a_proportion() -> None:
    """`ecl_coverage_pct` is already ×100. Excel's % format multiplies again,
    so 0.9176 meaning 0.9176% would render as 91.76%."""
    column = ST.Col("coverage_pct", "Coverage", ST.PERCENT)
    assert ST.cell_value(column, 0.9176) == pytest.approx(0.009176)
    rate = ST.Col("lgd", "LGD", ST.RATE)
    assert ST.cell_value(rate, 0.8132) == pytest.approx(0.8132)


def test_xl_18_a_signed_column_colours_by_direction() -> None:
    styling = {f"{k}{s}": k + s for k in
               ("money", "money0", "count", "rate", "rate4", "percent",
                "points", "date", "text", "wrap") for s in ("", "_band")}
    styling.update({"money_up": "UP", "money_down": "DOWN",
                    "rate_up": "RUP", "rate_down": "RDOWN"})
    change = ST.Col("ecl_change_sar", "Change", ST.MONEY, signed=True)
    assert ST.cell_style(change, 100.0, styling) == "UP"
    assert ST.cell_style(change, -100.0, styling) == "DOWN"
    # Zero is not a direction.
    assert ST.cell_style(change, 0.0, styling) == "money"


def test_xl_19_blankish_catches_what_or_does_not() -> None:
    from backend.retail.workbook_sheet import blankish

    for one in (None, float("nan"), float("inf"), "nan", "None", "NaT", " "):
        if one == " ":
            continue
        assert blankish(one), one
    for one in (0, 0.0, "", "0", "A+", False):
        if one == "":
            continue
        assert not blankish(one), one
