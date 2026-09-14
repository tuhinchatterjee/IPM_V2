"""
The retail What-If rebuild: migrations, the waterfall, the cohort vocabulary,
the thread's reading, and the workbook.

What this covers
----------------
Four things that did not exist before, and one that existed and was wrong.

New: migrations — a share of a delinquency bucket, a score band or an IFRS 9
stage moved into another; a waterfall that is COMPUTED rather than allocated;
an Excel workbook built from the result the screen showed; and a parser that
understands the Early Warning cohort vocabulary, so a sentence naming a
sub-product, a classification, a severity band or the forward-risk split
resolves to the same population the Early Warning card was showing.

Wrong: the What-If engine refused to filter on any of those dimensions,
because they are derived by the Early Warning model and are not columns of the
canonical book. "The canonical dataset has no such column" was true and
useless — the cohort exists, it is just not a column of the table being
filtered.

The defects this file holds the line on
---------------------------------------
1.  **A migration that picked at random.** Which facilities move decides the
    answer, so it is worst-first and deterministic: the facilities closest to
    the edge are the ones that go over it, the same ones every run, and they
    can be named. A migration that sampled would give a different number each
    time and could not be reconciled to a customer list.

2.  **A share that was a share of nothing in particular.** "20% of the 30-59
    bucket" is twenty per cent of its money or twenty per cent of its
    facilities, and those are different populations. The caller says which.

3.  **A landing value invented rather than read.** A facility moved to 90+
    takes the median days past due and the median PD anchor its own product
    shows in that bucket IN THIS BOOK. Nothing is assumed and nothing is
    carried across products.

4.  **A waterfall that allocated a total across the shocks asked for.** The
    IFRS 9 identity multiplies its inputs, so splitting the total invents
    steps that do not exist. Every step here is a real recomputation, and they
    sum to the total exactly because each is a difference between consecutive
    prefixes.

5.  **A second parser in the browser.** The thread carried six regular
    expressions and understood six shocks while the composer on the next
    screen understood the whole vocabulary. There is one reader.

6.  **A sentence that widened the cohort it was meant to narrow.** "Stress
    only the forward-risk customers in this selection" read as a fresh filter
    would stress every forward-risk customer in the book.

Every figure here is SYNTHETIC demonstration data from a synthetic
demonstration engine. Not an ANB model, not an ANB policy, not a SAMA
requirement, not independently validated.
"""

from __future__ import annotations

import numpy as np
import pytest

from backend.retail import ecl as ecl_mod
from backend.retail import ews_model as M
from backend.retail import ews_score as S
from backend.retail import profile
from backend.retail import whatif as W
from backend.retail import whatif_cohort as WC
from backend.retail import whatif_language as L
from backend.retail import whatif_selection as WS
from backend.retail.config import load_config

pytestmark = pytest.mark.skipif(
    not profile.is_retail(), reason="the corporate profile is active")


@pytest.fixture(scope="module")
def month() -> str:
    found = S.panel_months()
    if not found:
        pytest.skip("the Early Warning Score domain has not been built")
    return found[-1]


@pytest.fixture(scope="module")
def book(month: str):
    return S._read_book(month)


@pytest.fixture(scope="module")
def cards(book):
    return book[book["product_code"] == "CREDIT_CARD"]


@pytest.fixture(scope="module")
def weights():
    cfg = load_config()
    return {s: float(cfg.scenarios.weights[s]) for s in ecl_mod.SCENARIOS}


def _scenario(frame, **shocks):
    return W.Scenario(
        name="test", dataset_version="test",
        snapshot_date=str(frame["snapshot_date"].iloc[0]),
        filters={}, shocks=shocks)


# ======================================================= MIG: migrations =====

def test_mig_01_every_migration_is_a_declared_methodology() -> None:
    for one in ("dpd_migration", "score_band_migration", "stage_migration",
                "expense_pct"):
        assert one in W.SUPPORTED_METHODOLOGIES, one
        assert len(W.SUPPORTED_METHODOLOGIES[one]) > 40, one


def test_mig_02_a_delinquency_migration_moves_the_worst_first(cards,
                                                              weights) -> None:
    """Defect 1. Deterministic, and the ones nearest the edge go over it."""
    out = W._recompute(cards, _scenario(
        cards, dpd_migration={"from": "30-59", "to": "90-179",
                              "share": 0.5, "basis": "exposure"}), weights)
    picked = out["moved"]["dpd_migration"]
    assert picked.any(), "nothing moved"

    source = (cards["dpd_bucket"].astype(str) == "30-59").to_numpy()
    assert (picked & ~source).sum() == 0, "something outside the bucket moved"
    days = cards["dpd"].to_numpy(dtype=float)
    assert days[picked].min() >= days[source & ~picked].max(), (
        "a facility with fewer days past due moved while a worse one stayed")

    # Same inputs, same facilities: a migration is not a sample.
    again = W._recompute(cards, _scenario(
        cards, dpd_migration={"from": "30-59", "to": "90-179",
                              "share": 0.5, "basis": "exposure"}), weights)
    assert np.array_equal(picked, again["moved"]["dpd_migration"])


def test_mig_03_the_share_is_a_share_of_what_it_says(cards, weights) -> None:
    """Defect 2. Exposure and accounts are different populations."""
    source = cards[cards["dpd_bucket"].astype(str) == "30-59"]
    if len(source) < 10:
        pytest.skip("too few facilities in this bucket to split")

    by_money = W._recompute(cards, _scenario(
        cards, dpd_migration={"from": "30-59", "to": "90-179", "share": 0.5,
                              "basis": "exposure"}), weights)
    by_count = W._recompute(cards, _scenario(
        cards, dpd_migration={"from": "30-59", "to": "90-179", "share": 0.5,
                              "basis": "accounts"}), weights)

    money = cards["gross_carrying_amount_sar"].to_numpy(dtype=float)
    moved_money = by_money["moved"]["dpd_migration"]
    moved_count = by_count["moved"]["dpd_migration"]

    # Half the money, and at least half of it — the facility that crosses the
    # line is included so the share asked for is reached, not approached.
    whole = float(money[(cards["dpd_bucket"].astype(str) == "30-59").to_numpy()].sum())
    assert float(money[moved_money].sum()) >= 0.5 * whole
    # Half the accounts, counted.
    assert moved_count.sum() == round(0.5 * len(source))


def test_mig_04_a_migrated_facility_lands_where_this_book_puts_them(
        cards, weights) -> None:
    """Defect 3. Read from the data, never assumed."""
    out = W._recompute(cards, _scenario(
        cards, dpd_migration={"from": "30-59", "to": "90-179",
                              "share": 1.0}), weights)
    picked = out["moved"]["dpd_migration"]
    if not picked.any():
        pytest.skip("nothing in the source bucket")

    profile_of = W._bucket_profile(cards)
    products = cards["product_code"].astype(str).to_numpy()
    for at in np.flatnonzero(picked)[:20]:
        wanted = profile_of.get((products[at], "90-179"))
        if wanted is None:
            continue
        assert out["dpd"][at] == pytest.approx(wanted[0])
        assert out["pd_anchor"][at] == pytest.approx(wanted[1], rel=1e-9)


def test_mig_05_deterioration_raises_the_allowance(cards, weights) -> None:
    base = W._recompute(cards, _scenario(cards), weights)["ecl_final"].sum()
    for shock in (
        {"dpd_migration": {"from": "CURRENT", "to": "30-59", "share": 0.05}},
        {"stage_migration": {"from": 1, "to": 2, "share": 0.15}},
        {"score_band_migration": {"from": "B", "to": "D", "share": 0.25,
                                  "basis": "accounts"}},
        {"expense_pct": 0.10},
        {"income_pct": -0.10},
    ):
        after = W._recompute(cards, _scenario(cards, **shock),
                             weights)["ecl_final"].sum()
        assert after > base, shock


def test_mig_06_a_stage_migration_moves_only_the_stage(cards, weights) -> None:
    out = W._recompute(cards, _scenario(
        cards, stage_migration={"from": 1, "to": 2, "share": 0.15}), weights)
    picked = out["moved"]["stage_migration"]
    assert picked.any()
    assert (out["stage"][picked] == 2).all()
    # Nothing else was inferred from the move.
    untouched = W._recompute(cards, _scenario(cards), weights)
    assert np.allclose(out["pd_anchor"], untouched["pd_anchor"])
    assert np.allclose(out["gca"], untouched["gca"])


def test_mig_07_a_cure_is_a_migration_the_other_way(cards, weights) -> None:
    base = W._recompute(cards, _scenario(cards), weights)["ecl_final"].sum()
    cured = W._recompute(cards, _scenario(
        cards, stage_migration={"from": 2, "to": 1, "share": 0.5}),
        weights)["ecl_final"].sum()
    assert cured < base, "curing Stage 2 back to Stage 1 did not lower the ECL"


def test_mig_08_a_score_band_migration_reaches_pd_through_the_mapping(
        cards, weights) -> None:
    out = W._recompute(cards, _scenario(
        cards, score_band_migration={"from": "B", "to": "E", "share": 1.0,
                                     "basis": "accounts"}), weights)
    picked = out["moved"].get("score_band_migration")
    if picked is None or not picked.any():
        pytest.skip("no facilities in band B at this month")
    base = W._recompute(cards, _scenario(cards), weights)
    assert (out["pd_anchor"][picked] > base["pd_anchor"][picked]).mean() > 0.9
    # Everyone else is untouched.
    assert np.allclose(out["pd_anchor"][~picked], base["pd_anchor"][~picked])


def test_mig_09_a_migration_to_nowhere_leaves_the_facility_alone(
        cards, weights) -> None:
    """A share of a bucket this product has nobody in cannot land."""
    out = W._recompute(cards, _scenario(
        cards, dpd_migration={"from": "30-59", "to": "NOT_A_BUCKET",
                              "share": 1.0}), weights)
    base = W._recompute(cards, _scenario(cards), weights)
    assert np.allclose(out["dpd"], base["dpd"])
    assert np.allclose(out["ecl_final"], base["ecl_final"])


# ========================================================= WF: waterfall =====

@pytest.fixture(scope="module")
def many(cards, weights):
    scenario = _scenario(
        cards, income_pct=-0.10,
        dpd_migration={"from": "CURRENT", "to": "1-29", "share": 0.05},
        stage_migration={"from": 1, "to": 2, "share": 0.10},
        pd_relative=0.20, lgd_relative=0.05)
    return W.waterfall(cards, scenario, weights)


def test_wf_01_the_steps_sum_exactly_to_the_total(many) -> None:
    """Defect 4. Because each is a difference between consecutive runs."""
    assert many["available"]
    summed = sum(one["change_sar"] for one in many["steps"])
    assert summed == pytest.approx(many["total_change_sar"], abs=0.02)


def test_wf_02_every_shock_asked_for_gets_a_step(many) -> None:
    keys = {one["key"] for one in many["steps"]}
    assert keys == {"income_pct", "dpd_migration", "stage_migration",
                    "pd_relative", "lgd_relative"}


def test_wf_03_the_steps_chain_end_to_end(many) -> None:
    previous = many["baseline_sar"]
    for one in many["steps"]:
        assert one["from_sar"] == pytest.approx(previous, abs=0.02)
        previous = one["to_sar"]
    assert previous == pytest.approx(many["final_sar"], abs=0.02)


def test_wf_04_it_runs_in_the_published_causal_order(many) -> None:
    order = [one for one in W.WATERFALL_ORDER
             if one in {step["key"] for step in many["steps"]}]
    assert [one["key"] for one in many["steps"]] == order


def test_wf_05_a_migration_step_names_how_many_moved(many) -> None:
    moved = {one["key"]: one["facilities_moved"] for one in many["steps"]}
    assert moved["dpd_migration"] and moved["dpd_migration"] > 0
    assert moved["stage_migration"] and moved["stage_migration"] > 0
    # A shock that applies to everybody reports None rather than a count that
    # would read as a subset.
    assert moved["pd_relative"] is None


def test_wf_06_it_says_what_a_step_is_and_is_not(many) -> None:
    assert "sum exactly" in many["basis"]
    assert "not independent" in many["basis"]


def test_wf_07_a_whole_book_run_declines_rather_than_waits(book) -> None:
    found = W.run(book, W.Scenario(
        name="wide", dataset_version="", filters={},
        snapshot_date=str(book["snapshot_date"].iloc[0]),
        shocks={"pd_relative": 0.1}), load_config())
    steps = found.get("waterfall") or {}
    if len(book) > W.WATERFALL_MAX_FACILITIES:
        assert steps.get("available") is False
        assert "Narrow it" in steps.get("because", "")
    else:
        assert steps.get("available")


# ============================================ SEL: the cohort vocabulary =====

def test_sel_01_the_early_warning_dimensions_are_declared() -> None:
    for one in ("classification", "sub_product", "ews_severity",
                "current_bad_flag", "forward_risk_flag"):
        assert one in W.EARLY_WARNING_DIMENSIONS, one


def test_sel_02_a_derived_dimension_resolves_to_a_population(book) -> None:
    """Defect: the engine refused every Early Warning cut."""
    whole = len(book)
    for wanted in (
        {"classification": "NON_SALARIED"},
        {"sub_product": "CC_PLATINUM"},
        {"ews_severity": ["HIGH", "CRITICAL"]},
        {"forward_risk_flag": True},
        {"current_bad_flag": True},
    ):
        found = W.select(book, wanted)
        assert 0 < len(found) < whole, wanted


def test_sel_03_derived_and_book_dimensions_compose(book) -> None:
    narrow = W.select(book, {"product_code": "CREDIT_CARD",
                             "classification": "SALARIED",
                             "sub_product": "CC_PLATINUM"})
    wider = W.select(book, {"product_code": "CREDIT_CARD",
                            "classification": "SALARIED"})
    assert 0 < len(narrow) <= len(wider) < len(book)
    assert set(narrow["facility_id"]) <= set(wider["facility_id"])


def test_sel_04_an_unknown_column_is_still_refused(book) -> None:
    with pytest.raises(KeyError):
        W.select(book, {"not_a_column": 1})


# ================================================== PAR: the parser reads ====

@pytest.mark.parametrize("said,wanted", [
    ("Stress non-salaried credit card customers",
     {"product_code": "CREDIT_CARD", "classification": "NON_SALARIED"}),
    ("Stress Platinum Credit Card salaried customers",
     {"sub_product": "CC_PLATINUM", "classification": "SALARIED"}),
    ("Stress only high and critical EWS customers",
     {"ews_severity": ["HIGH", "CRITICAL"]}),
    ("Run the shock on forward-risk customers only",
     {"forward_risk_flag": True}),
    ("Show impact only for already bad customers", {"current_bad_flag": True}),
    ("Apply the shock to customers with behavioural score band D and below",
     {"behavioural_score_band": ["E", "D"]}),
    ("Stress Stage 2 non-salaried borrowers in Credit Card",
     {"product_code": "CREDIT_CARD", "classification": "NON_SALARIED",
      "ifrs9_stage": 2}),
])
def test_par_01_a_cohort_sentence_resolves(said: str, wanted: dict) -> None:
    found = L.read(said, months=S.panel_months()).filters
    for column, value in wanted.items():
        assert found.get(column) == value, (said, column, found)


@pytest.mark.parametrize("said,key,wanted", [
    ("Move 20% of 30-59 DPD exposure to 90+", "dpd_migration",
     {"from": "30-59", "to": "90-179", "share": 0.2, "basis": "exposure"}),
    ("Move 100% of 1-29 DPD to 60-89", "dpd_migration",
     {"from": "1-29", "to": "60-89", "share": 1.0, "basis": "exposure"}),
    ("Move 15% of Stage 1 exposure to Stage 2", "stage_migration",
     {"from": 1, "to": 2, "share": 0.15, "basis": "exposure"}),
    ("Move 10% of Stage 2 accounts to Stage 3", "stage_migration",
     {"from": 2, "to": 3, "share": 0.1, "basis": "accounts"}),
    ("Cure 20% of Stage 2 back to Stage 1", "stage_migration",
     {"from": 2, "to": 1, "share": 0.2, "basis": "exposure"}),
    ("Move half the Stage 1 exposure to Stage 2", "stage_migration",
     {"from": 1, "to": 2, "share": 0.5, "basis": "exposure"}),
])
def test_par_02_a_migration_sentence_resolves(said: str, key: str,
                                              wanted: dict) -> None:
    found = L.read(said, months=S.panel_months()).shocks
    assert key in found, (said, found)
    for field, value in wanted.items():
        assert found[key][field] == value, (said, field, found[key])


def test_par_03_a_migration_does_not_leave_its_source_as_a_filter() -> None:
    """Read as a filter as well, "move 15% of Stage 1" stresses Stage 1 only."""
    ask = L.read("Move 15% of Stage 1 exposure to Stage 2",
                 months=S.panel_months())
    assert "ifrs9_stage" not in ask.filters
    ask = L.read("Move 20% of 30-59 DPD exposure to 90+",
                 months=S.panel_months())
    assert "dpd_bucket" not in ask.filters


def test_par_04_expenses_are_read_apart_from_income() -> None:
    assert L.read("Increase household expense burden by 10%",
                  months=S.panel_months()).shocks == {"expense_pct": 0.10}
    assert L.read("Reduce verified income by 10%",
                  months=S.panel_months()).shocks == {"income_pct": -0.10}


def test_par_05_every_sub_product_in_the_taxonomy_is_readable() -> None:
    phrases = L._sub_products()
    for code in M.SUB_PRODUCT_LABELS:
        assert code in set(phrases.values()), code


def test_par_06_a_compound_sentence_reads_all_of_its_parts() -> None:
    ask = L.read(
        "Shift 10% of current accounts into 1-29 DPD for non-salaried "
        "credit card customers", months=S.panel_months())
    assert ask.shocks["dpd_migration"]["from"] == "CURRENT"
    assert ask.shocks["dpd_migration"]["to"] == "1-29"
    assert ask.shocks["dpd_migration"]["basis"] == "accounts"
    assert ask.filters["product_code"] == "CREDIT_CARD"
    assert ask.filters["classification"] == "NON_SALARIED"


# ================================================ THR: the thread reads ======

@pytest.fixture(scope="module")
def selection(month: str):
    made = WS.create(month=month, level="sub_product", product="CREDIT_CARD",
                     classification="SALARIED", sub_product="CC_PLATINUM",
                     created_by="tests")
    if not made.selected_facility_ids:
        pytest.skip("nothing in this cohort at this month")
    WS.save(made)
    return made


def test_thr_01_the_thread_reads_with_the_governed_parser(selection) -> None:
    """Defect 5. One reader, and it is the one that gets improved."""
    found = WC.read("Increase PIT 12-month PD by 20%", selection)
    assert found["understood"]
    assert found["shocks"] == {"pd_relative": 0.20}


def test_thr_02_an_unreadable_sentence_is_a_question_not_a_crash(
        selection) -> None:
    found = WC.read("make it worse somehow", selection)
    assert found["understood"] is False
    assert found["question"]


def test_thr_03_the_thread_reads_every_chip_it_offers(selection) -> None:
    """Every chip the composer shows has to run, or it is a broken promise."""
    chips = [
        "Increase PIT 12-month PD by 20%",
        "Increase LGD by 5 percentage points",
        "Reduce verified income by 10%",
        "Increase household expense burden by 10%",
        "Move 20% of 30-59 DPD exposure to 90+",
        "Move 15% of Stage 1 exposure to Stage 2",
        "Move 20% of behavioural score band B to C band",
        "Increase card utilisation by 10 percentage points",
        "Increase CCF by 10 percentage points",
        "Reduce collateral value by 15%",
        "Stress only the forward-risk customers in this selection",
        "Stress only the customers who are already bad",
    ]
    unreadable = []
    for chip in chips:
        found = WC.read(chip, selection)
        if not found["understood"] and not found.get("within"):
            unreadable.append(chip)
    # The last two name only a cohort; they narrow rather than shock, and the
    # thread answers them by asking what to apply.
    assert unreadable == [
        "Stress only the forward-risk customers in this selection",
        "Stress only the customers who are already bad",
    ], unreadable


def test_thr_04_a_narrowing_never_widens_the_cohort(selection) -> None:
    """Defect 6."""
    held = set(selection.selected_facility_ids)
    kept, said = WC.narrow(selection, {"forward_risk_flag": True})
    assert set(kept) <= held
    if len(kept) != len(held):
        assert said and "of the selection's" in said


def test_thr_05_a_narrowing_that_matches_nobody_is_refused(
        selection) -> None:
    out = WC.run(selection.selection_id, shocks={"pd_relative": 0.2},
                 within={"classification": "NON_SALARIED"})
    # This selection is Salaried, so narrowing it to Non-Salaried is empty.
    assert out.get("available") is False
    assert "nothing to stress" in out["because"]


@pytest.mark.slow
def test_thr_06_a_result_carries_its_waterfall_and_its_baseline(
        selection) -> None:
    out = WC.run(selection.selection_id,
                 shocks={"pd_relative": 0.20, "lgd_relative": 0.05})
    assert out["available"]
    assert (out["waterfall"] or {}).get("available")
    assert (out["baseline"] or {}).get("available")
    assert out["stressed_facilities"] == selection.selected_account_count


# ============================== STAMP: a rebuilt book is not skipped =========

def test_stamp_01_a_derived_domain_records_the_book_it_came_from() -> None:
    """Every derived domain knows which book it was built from."""
    from pathlib import Path

    from backend.config import settings
    from backend.retail import domains as D
    from backend.retail import source_stamp

    root = Path(settings.analytics_dir)
    current = source_stamp.book_hash()
    assert current, "the published book has no manifest hash"
    for domain in (S.DOMAIN, *(one.dataset for one in D.DERIVED)):
        if not (root / domain).exists():
            continue
        assert not source_stamp.stale(root, domain), (
            f"{domain} was built from a different book than the one "
            "published, so it is serving data that no longer exists")


def test_stamp_02_a_different_book_reads_as_stale(tmp_path) -> None:
    """The mechanism itself, at the level it failed.

    The period names survive a regeneration, so nothing in the file layout can
    notice. This compares the books.
    """
    import json

    from backend.retail import source_stamp

    domain = "a_domain"
    (tmp_path / domain).mkdir(parents=True)
    metadata = tmp_path / "metadata"
    metadata.mkdir()

    (metadata / source_stamp.MANIFEST).write_text(
        json.dumps({"manifest_hash": "first"}))
    source_stamp.record(tmp_path, domain, metadata)
    assert not source_stamp.stale(tmp_path, domain, metadata)

    (metadata / source_stamp.MANIFEST).write_text(
        json.dumps({"manifest_hash": "second"}))
    assert source_stamp.stale(tmp_path, domain, metadata)


def test_stamp_03_an_unstamped_installation_is_not_called_stale(
        tmp_path) -> None:
    """A missing stamp is not evidence, and must not force a rebuild."""
    import json

    from backend.retail import source_stamp

    (tmp_path / "unstamped").mkdir(parents=True)
    metadata = tmp_path / "metadata"
    metadata.mkdir()
    (metadata / source_stamp.MANIFEST).write_text(
        json.dumps({"manifest_hash": "something"}))
    assert not source_stamp.stale(tmp_path, "unstamped", metadata)


# ===================================================== XL: the workbook ======

@pytest.mark.slow
def test_xl_01_the_workbook_builds_and_opens(selection) -> None:
    import io

    from openpyxl import load_workbook

    from backend.retail import whatif_workbook as book

    out = WC.run(selection.selection_id, shocks={"pd_relative": 0.20},
                 method="both")
    payload, filename = book.build(out, selection=selection.to_dict())
    assert filename.endswith(".xlsx")
    assert len(payload) > 10_000

    opened = load_workbook(io.BytesIO(payload))
    assert len(opened.sheetnames) >= 12, opened.sheetnames
    for wanted in ("Summary", "Waterfall", "Cohort", "Impact by level",
                   "Affected customers", "Method"):
        assert any(wanted in one for one in opened.sheetnames), wanted


@pytest.mark.slow
def test_xl_02_every_sheet_says_something(selection) -> None:
    """A sheet that does not apply says so rather than being blank."""
    import io

    from openpyxl import load_workbook

    from backend.retail import whatif_workbook as book

    out = WC.run(selection.selection_id, shocks={"pd_relative": 0.20})
    payload, _ = book.build(out, selection=selection.to_dict())
    opened = load_workbook(io.BytesIO(payload))
    for name in opened.sheetnames:
        page = opened[name]
        text = "".join(str(cell.value or "")
                       for row in page.iter_rows(max_row=12)
                       for cell in row)
        assert len(text) > 20, f"{name} is empty"


@pytest.mark.slow
def test_xl_03_the_workbook_agrees_with_the_result_it_came_from(
        selection) -> None:
    """Nothing is recomputed, so nothing can disagree."""
    import io

    from openpyxl import load_workbook

    from backend.retail import whatif_workbook as book

    out = WC.run(selection.selection_id, shocks={"pd_relative": 0.20})
    payload, _ = book.build(out, selection=selection.to_dict())
    opened = load_workbook(io.BytesIO(payload))

    steps = out["waterfall"]["steps"]
    page = opened[next(one for one in opened.sheetnames if "Waterfall" in one)]
    found = {str(row[0].value): row for row in page.iter_rows(min_row=1)
             if row and row[0].value}
    for one in steps:
        row = found.get(one["label"])
        assert row is not None, one["label"]
        assert float(row[3].value) == pytest.approx(one["change_sar"], abs=0.02)


@pytest.mark.slow
def test_xl_04_it_claims_nothing_it_cannot_support(selection) -> None:
    import io

    from openpyxl import load_workbook

    from backend.retail import whatif_workbook as book

    out = WC.run(selection.selection_id, shocks={"pd_relative": 0.20})
    payload, _ = book.build(out, selection=selection.to_dict())
    opened = load_workbook(io.BytesIO(payload))
    text = " ".join(
        str(cell.value or "")
        for name in opened.sheetnames
        for row in opened[name].iter_rows()
        for cell in row).lower()
    for forbidden in ("approved by sama", "sama approval", "anb approved",
                      "independently validated", "auditor certified"):
        assert forbidden not in text, forbidden
    assert "synthetic" in text
