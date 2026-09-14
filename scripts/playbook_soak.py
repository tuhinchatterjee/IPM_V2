"""
Run the Playbook journeys again and again, and check they land in the same
place every time. §20, PB-042.

Why a soak harness and not another pytest run
---------------------------------------------
A single pass proves a journey works once. It cannot see the failures that
only appear on repetition: state leaking from one run into the next, a second
version quietly appearing, a selection lost on reopen, formatting drifting a
little further each cycle, a source degraded by being read again. Those are
the defects that reach a user, because a user does the same thing every
quarter.

So every journey runs N times (default 10) in its own workspace, and the
governed state each cycle produces is fingerprinted and compared. A journey
that lands somewhere different on cycle 7 than on cycle 1 fails here, and the
report says which field moved.

    .venv/bin/python scripts/playbook_soak.py
    .venv/bin/python scripts/playbook_soak.py --cycles 3 --only A,B
    .venv/bin/python scripts/playbook_soak.py --list

Exit 0 if every cycle of every journey held, 1 if any did not, 2 if it could
not run. **No provider call is made.** The authoring provider is replaced with
a deterministic stand-in for the whole run, and any attempt to reach the real
one fails the harness rather than spending money — §20 is explicit that paid
credits are not to be burned repeating deterministic checks.

What the harness does NOT cover
-------------------------------
Anything that needs a browser (`scripts/acceptance/playbook_browser_acceptance.py`
does that) and anything that needs the live provider
(`scripts/playbook_live_slice.py`). A run of this script is never reported as
live verification.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import pathlib
import re
import sys
import time
import traceback
import zipfile

REPO = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

EVIDENCE = REPO / "docs" / "playbook" / "soak_results.json"

REPORT_MD = """# {title}

## 1. Executive summary

Weighted ECL rose to SAR 22.77 million from SAR 20.90 million, an increase of
SAR 1.87 million.

## 2. Scenario results

| Scenario | ECL, SAR million |
| --- | --- |
| Base | 19.20 |
| Downturn | 36.00 |

## 3. Limitations

Post-model adjustments are outside the scope of this report.
"""

EDITED_SUMMARY = """Weighted ECL rose to SAR 22.77 million from SAR 20.90
million. The increase of SAR 1.87 million is concentrated in Contracting."""


# --------------------------------------------------------------------------
# The deterministic stand-in
# --------------------------------------------------------------------------


class ProviderReached(AssertionError):
    """The real provider was called. The soak must never spend money."""


def install_stand_in():
    """Replace the authoring provider for the whole run.

    Returns a setter for what the author "writes". Nothing here reaches a
    network, and results obtained this way are never reported as live
    verification.
    """
    from backend.playbook import provider

    state = {"text": ""}

    def author(*, system, messages, formats, purpose="playbook_authoring",
               container_id="", on_milestone=None, on_delta=None,
               is_cancelled=None):
        if on_milestone:
            on_milestone("drafting", "soak")
        if on_delta:
            on_delta(state["text"])
        return provider.AuthoringResult(
            text=state["text"], files=[], model_requested="soak-stand-in",
            model_served="soak-stand-in", request_ids=["req_soak"], turns=1)

    def refuse(*args, **kwargs):
        raise ProviderReached(
            "the soak harness reached the real provider; no cycle may spend "
            "money")

    provider.author = author
    provider._stream_once = refuse
    provider._client = refuse

    def write(text: str) -> None:
        state["text"] = text

    return write


# --------------------------------------------------------------------------
# Fixtures, built in code so every cycle gets identical bytes
# --------------------------------------------------------------------------


def workbook(*, base=19.20, downturn=36.00) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet.title = "ECL"
    sheet.append(["Scenario", "ECL, SAR million", "Weight"])
    sheet.append(["Base", base, 0.60])
    sheet.append(["Downturn", downturn, 0.25])
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "0.00"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def word_report() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("IFRS 9 Committee Report — Q1 2026", level=1)
    doc.add_heading("1. Executive summary", level=1)
    doc.add_paragraph("Weighted ECL for the quarter was SAR 20.90 million.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# The fingerprint a cycle has to reproduce
# --------------------------------------------------------------------------


#: Two renders of the same document are not byte-identical, and should not be
#: expected to be. Measured, not assumed: for DOCX every ZIP member's CONTENT
#: is identical and only the members' timestamps move; for PDF the bytes are
#: identical once reportlab's `/ID` and `/CreationDate` are normalised.
#:
#: Those are the file's own record of when it was made. Forcing them to a
#: constant would make every generated document claim a fictional creation
#: time, which is a worse thing to ship than a hash that has to say what it
#: excludes. So the digest below excludes exactly those fields and nothing
#: else — which makes it a real test of "no progressive formatting drift"
#: rather than a test of the clock.
_PDF_ID = re.compile(rb"/ID\s*\[[^\]]*\]")
_PDF_DATE = re.compile(rb"/(?:Creation|Mod)Date\s*\([^)]*\)")


def stable_digest(content: bytes, fmt: str) -> str:
    """A hash of what a rendered file SAYS, ignoring when it was made."""
    if fmt in ("docx", "pptx", "xlsx"):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            payload = b"".join(
                name.encode() + b"\0" + archive.read(name)
                for name in sorted(archive.namelist()))
    elif fmt == "pdf":
        payload = _PDF_DATE.sub(b"/Date()", _PDF_ID.sub(b"/ID[]", content))
    else:
        payload = content
    return hashlib.sha256(payload).hexdigest()


def fingerprint(session, workspace_id: int) -> dict:
    """The governed state, in a form two cycles can be compared by.

    Ids, timestamps and workspace names are excluded on purpose: they differ
    between cycles by construction. Everything that describes WHAT the journey
    produced is included, so a drift in content, counts, statuses or
    formatting shows up as a difference.
    """
    from backend.playbook import reparse, store
    from backend.playbook import repository as repo
    from backend.playbook.intelligence import service as intel

    dash = intel.dashboard(session, workspace_id).as_dict()
    artifacts = repo.artifacts(session, workspace_id)

    documents = []
    for artifact in artifacts:
        for version in repo.versions(session, artifact.id):
            documents.append({
                "kind": artifact.kind,
                "version": version.version,
                "content_hash": version.content_hash,
                "files": sorted(
                    (f.format, f.size_bytes,
                     stable_digest(store.read(f.bytes_path), f.format))
                    for f in repo.files(session, version.id)),
            })

    return {
        "documents": documents,
        "sections": [(s["heading"], s["status"], s["word_count"])
                     for s in dash["sections"]],
        "metrics": {k: dash["metrics"].get(k) for k in
                    ("detected", "confirmed", "suggested", "coverage_pct")},
        "findings": {k: dash["findings"].get(k)
                     for k in ("total", "open", "blocking")},
        "readiness": dash["readiness"].get("completion_pct"),
        "sources": [(s["filename"], s["status"], s["chunk_count"], s["stale"])
                    for s in reparse.summary(session, workspace_id)["items"]],
        "since_last_time": [
            (r["metric_id"], r["then"], r["now"], r.get("change", ""))
            for r in dash["since_last_time"].get("rows", [])],
    }


def difference(first: dict, other: dict) -> str:
    """The first field that differs, named. "They are not equal" is not a
    diagnosis anybody can act on."""
    for key in first:
        if first[key] != other.get(key):
            return f"{key}: cycle 1 {first[key]!r} != this cycle {other.get(key)!r}"
    for key in other:
        if key not in first:
            return f"{key}: present this cycle, absent on cycle 1"
    return ""


# --------------------------------------------------------------------------
# Journeys
# --------------------------------------------------------------------------


class Cycle:
    """One run of one journey, in its own workspace."""

    def __init__(self, session, scope, write, number: int):
        self.session = session
        self.scope = scope
        self.write = write
        self.number = number
        self.checks: list[tuple[bool, str]] = []
        self.workspace = None

    def check(self, ok: bool, what: str, detail: str = "") -> bool:
        self.checks.append((bool(ok), what if ok else f"{what} — {detail}"))
        return bool(ok)

    def new_workspace(self, title: str):
        from backend.playbook import repository as repo

        self.workspace = repo.create_workspace(
            self.session, self.scope,
            title=f"{title} (soak {self.number})",
            document_family="ifrs9_committee_report")
        self.session.flush()
        return self.workspace

    def upload(self, filename: str, content: bytes, role="results"):
        from backend.playbook import service

        return service.add_source(self.session, self.scope, self.workspace.id,
                                  filename=filename, content=content,
                                  source_role=role)

    def generate(self, text: str, **kwargs):
        from backend.playbook import service

        self.write(text)
        return service.author_document(
            self.session, self.scope, self.workspace.id,
            instruction=kwargs.pop("instruction", "Write the report."),
            ledger=kwargs.pop("ledger", None) or self.ledger(),
            title="IFRS 9 Committee Report",
            formats=kwargs.pop("formats", ["docx", "pdf"]), **kwargs)

    def ledger(self):
        from backend.playbook import service

        return service.ledger_for(self.session, self.scope, self.workspace.id)


def journey_a(cycle: Cycle) -> None:
    """A. Create a report from an uploaded workbook."""
    cycle.new_workspace("Create")
    cycle.upload("results.xlsx", workbook())
    outcome = cycle.generate(REPORT_MD.format(title="IFRS 9 Committee Report"))

    cycle.check(outcome.version_id is not None, "a version was written")
    cycle.check(outcome.grounding_final and outcome.grounding_final.ok,
                "the saved report is grounded",
                str(outcome.grounding_final.findings
                    if outcome.grounding_final else "no result"))
    cycle.check(set(outcome.files) == {"docx", "pdf"},
                "both formats were produced", str(sorted(outcome.files)))
    cycle.check(outcome.adoption.get("sections") == 3,
                "the dashboard records three sections",
                str(outcome.adoption))


def journey_b(cycle: Cycle) -> None:
    """B. Edit one section and leave the rest byte-identical."""
    from backend.playbook import merge
    from backend.playbook import repository as repo

    cycle.new_workspace("Scoped edit")
    cycle.upload("results.xlsx", workbook())
    first = cycle.generate(REPORT_MD.format(title="IFRS 9 Committee Report"))

    before = {s.heading: merge.section_hash(s)
              for s in first.document.sections}

    edited = REPORT_MD.format(title="IFRS 9 Committee Report").replace(
        "Weighted ECL rose to SAR 22.77 million from SAR 20.90 million, an "
        "increase of\nSAR 1.87 million.", EDITED_SUMMARY)
    second = cycle.generate(edited, instruction="Tighten the summary.",
                            artifact_id=first.artifact_id, task_kind="edit",
                            task_scope="1. Executive summary")

    after = {s.heading: merge.section_hash(s)
             for s in second.document.sections}
    unrelated = [h for h in before
                 if h != "1. Executive summary" and before[h] != after.get(h)]
    cycle.check(not unrelated, "no unrelated section changed", str(unrelated))
    cycle.check(before["1. Executive summary"] != after["1. Executive summary"],
                "the requested section did change")
    cycle.check(len(repo.versions(cycle.session, first.artifact_id)) == 2,
                "exactly two versions exist — no duplicate")


def journey_c(cycle: Cycle) -> None:
    """C. Re-read a stale source."""
    from backend.playbook import reparse
    from backend.playbook import repository as repo

    cycle.new_workspace("Re-read")
    source = cycle.upload("results.xlsx", workbook())
    chunks_before = [(c.locator, c.text, c.data)
                     for c in repo.chunks(cycle.session, source.id)]

    reparse.latest(cycle.session, source.id).parser_version = "0"
    cycle.session.flush()
    cycle.check(reparse.state(cycle.session, source).stale,
                "an older reading is reported stale")

    reading = reparse.reread(cycle.session, cycle.scope, source.id)
    cycle.check(not reading.stale, "the re-read brought it up to date")
    cycle.check(reading.revision == 2, "a new revision was written",
                str(reading.revision))

    chunks_after = [(c.locator, c.text, c.data)
                    for c in repo.chunks(cycle.session, source.id)]
    # No source corruption: the same bytes, read by the same reader, produce
    # the same chunks. A re-read that drifted would poison every later report.
    cycle.check(chunks_before == chunks_after,
                "re-reading the same bytes produced the same chunks")
    cycle.check(len(reparse.history(cycle.session, source.id)) == 2,
                "the earlier reading is kept, not deleted")


def journey_d(cycle: Cycle) -> None:
    """D. New data arrives, and Since Last Time says what moved.

    The real journey, and the order matters: a version is written, the data
    then moves, and the comparison is made WITHOUT regenerating. THEN is what
    the paper in front of the reader relied on; NOW is what the data says
    today. Regenerating first would freeze the new reading and leave nothing
    to compare — which is what makes this worth running ten times.
    """
    from backend.exports import playbook_contract as contract
    from backend.models.playbook import PlaybookMetricBinding
    from backend.playbook.intelligence import binding as bind
    from backend.playbook.intelligence import compare

    cycle.new_workspace("Since last time")
    cycle.upload("results.xlsx", workbook())
    bind.apply(cycle.session, cycle.workspace.id, bind.from_export(
        # `value` is exact and `display_value` is how it is shown. Both
        # travel: the comparison is made on the exact value, and a fixture
        # that supplies only the display would be comparing presentation.
        [contract.Metric(metric_id="ecl.weighted", label="Weighted ECL",
                         value="20.90", display_value="20.90",
                         unit="currency")]))
    cycle.session.flush()
    cycle.generate(REPORT_MD.format(title="IFRS 9 Committee Report"))

    [binding] = (cycle.session.query(PlaybookMetricBinding)
                 .filter(PlaybookMetricBinding.workspace_id ==
                         cycle.workspace.id).all())
    binding.value_in_document = "22.77"
    binding.raw_value = "22.77"
    binding.display_value = "22.77"
    binding.freshness = bind.NEW_AVAILABLE
    cycle.session.flush()

    table = compare.since_last_time(cycle.session, cycle.workspace.id)
    cycle.check(table.get("available"), "there is something to compare",
                table.get("reason", ""))
    rows = {r["metric_id"]: r for r in table.get("rows", [])}
    row = rows.get("ecl.weighted", {})
    cycle.check(row.get("then", {}).get("display") == "20.90",
                "THEN is the reading the saved version relied on",
                str(row.get("then")))
    cycle.check(row.get("now", {}).get("display") == "22.77",
                "NOW is the current reading", str(row.get("now")))
    cycle.check(row.get("comparable") is not False,
                "the two readings are comparable", str(row.get("reason", "")))

    # And the stale reading is never quietly treated as the document's own:
    # the section that relied on it is named for refresh, not rewritten.
    from backend.playbook.intelligence import context as ctx

    built = ctx.build(cycle.session, cycle.workspace.id, kind="stale_metrics")
    cycle.check(built.action == ctx.REFRESH_AFFECTED,
                "the newer reading offers a governed refresh")
    cycle.check(any("Nothing is rewritten" in c for c in built.caveats),
                "and says plainly that it rewrites nothing")


def journey_e(cycle: Cycle) -> None:
    """E. Read the dashboard repeatedly; it must not change by being read."""
    from backend.playbook.intelligence import service as intel

    cycle.new_workspace("Dashboard")
    cycle.upload("results.xlsx", workbook())
    cycle.generate(REPORT_MD.format(title="IFRS 9 Committee Report"))

    reads = [intel.dashboard(cycle.session, cycle.workspace.id).as_dict()
             for _ in range(5)]
    cycle.check(all(r == reads[0] for r in reads),
                "five reads of the dashboard agree")
    cycle.check(reads[0]["available"] is True, "the dashboard is available")

    provider_free = _provider_is_untouched(cycle)
    cycle.check(provider_free, "reading the dashboard called no provider")


def _provider_is_untouched(cycle: Cycle) -> bool:
    from backend.playbook import provider
    from backend.playbook.intelligence import service as intel

    saved, provider.author = provider.author, _explode
    try:
        intel.dashboard(cycle.session, cycle.workspace.id)
        return True
    except ProviderReached:
        return False
    finally:
        provider.author = saved


def _explode(*args, **kwargs):
    raise ProviderReached("the dashboard must never call a provider")


def journey_f(cycle: Cycle) -> None:
    """F. Reopen after every state change, from a session that never saw it.

    The closest a harness gets to a browser refresh: everything is read back
    through a NEW session, so anything held only in the identity map fails.
    """
    from backend.db.engine import get_session
    from backend.playbook import repository as repo

    cycle.new_workspace("Reopen")
    cycle.upload("results.xlsx", workbook())
    outcome = cycle.generate(REPORT_MD.format(title="IFRS 9 Committee Report"))
    workspace_id, artifact_id = cycle.workspace.id, outcome.artifact_id
    cycle.session.commit()

    with get_session() as fresh:
        versions = repo.versions(fresh, artifact_id)
        cycle.check(len(versions) == 1, "the version survived the reopen")
        cycle.check(len(repo.sources(fresh, workspace_id)) == 1,
                    "the source survived the reopen")
        state = fingerprint(fresh, workspace_id)
        cycle.check(bool(state["sections"]),
                    "the dashboard reads back from a session that never wrote")


def journey_g(cycle: Cycle) -> None:
    """G. Persistence across a restart, simulated by a new engine connection.

    A new session from the pool is not a restart. Disposing the engine forces
    a genuinely new connection, which is what the process would get.
    """
    from backend.db.engine import engine, get_session
    from backend.playbook import repository as repo

    cycle.new_workspace("Restart")
    cycle.upload("results.xlsx", workbook())
    outcome = cycle.generate(REPORT_MD.format(title="IFRS 9 Committee Report"))
    workspace_id, artifact_id = cycle.workspace.id, outcome.artifact_id
    before = fingerprint(cycle.session, workspace_id)
    cycle.session.commit()

    engine.dispose()

    with get_session() as fresh:
        cycle.check(fingerprint(fresh, workspace_id) == before,
                    "the governed state is identical after a restart",
                    difference(before, fingerprint(fresh, workspace_id)))
        files = repo.files(fresh,
                           repo.versions(fresh, artifact_id)[0].id)
        from backend.playbook import store

        for row in files:
            content = store.read(row.bytes_path)
            cycle.check(store.sha256(content) == row.sha256,
                        f"the stored {row.format} bytes are unchanged")


def journey_h(cycle: Cycle) -> None:
    """H. A double send is one generation."""
    from backend.playbook import repository as repo
    from backend.playbook import service

    cycle.new_workspace("Duplicate send")
    cycle.upload("results.xlsx", workbook())
    key = f"soak-{cycle.number}-double"

    first = service.begin_generation(cycle.session, cycle.scope,
                                     cycle.workspace.id, text="Write it.",
                                     idempotency_key=key)
    second = service.begin_generation(cycle.session, cycle.scope,
                                      cycle.workspace.id, text="Write it.",
                                      idempotency_key=key)

    cycle.check(second["duplicate"] is True, "the second send is a duplicate")
    cycle.check(first["job_id"] == second["job_id"],
                "it found the running job rather than starting another")
    messages = repo.messages(cycle.session, cycle.workspace.id)
    cycle.check(len([m for m in messages if m.role == "user"]) == 1,
                "the thread carries one question, not two")


def journey_i(cycle: Cycle) -> None:
    """I. Restore an earlier version, forwards."""
    from backend.playbook import repository as repo
    from backend.playbook import service

    cycle.new_workspace("Restore")
    cycle.upload("results.xlsx", workbook())
    first = cycle.generate(REPORT_MD.format(title="IFRS 9 Committee Report"))
    v1_hash = repo.versions(cycle.session, first.artifact_id)[0].content_hash

    cycle.generate(
        REPORT_MD.format(title="IFRS 9 Committee Report").replace(
            "Post-model adjustments are outside the scope of this report.",
            "Post-model adjustments are outside scope. Nothing else changed."),
        instruction="Revise.", artifact_id=first.artifact_id)

    restored = service.restore_version(cycle.session, cycle.scope,
                                       first.artifact_id, 1)
    versions = repo.versions(cycle.session, first.artifact_id)
    cycle.check(len(versions) == 3, "restoring moved forward to v3",
                str(len(versions)))
    cycle.check(versions[-1].content_hash == v1_hash,
                "v3 carries v1's content exactly")
    cycle.check(versions[0].content_hash == v1_hash,
                "v1 is still there and unchanged")
    cycle.check(restored is not None, "the restore returned its version")


def journey_j(cycle: Cycle) -> None:
    """J. A malformed source is refused, and the workspace still works."""
    from backend.playbook import ingest
    from backend.playbook import repository as repo

    cycle.new_workspace("Malformed source")

    refused = 0
    for filename, content in (
            ("truncated.docx", b"PK\x03\x04not really a zip"),
            ("empty.xlsx", b""),
            ("binary.exe", b"MZ\x90\x00\x03\x00\x00\x00"),
    ):
        try:
            cycle.upload(filename, content)
        except (ingest.RejectedUpload, ingest.UnreadableSource):
            refused += 1
        except Exception as exc:                      # noqa: BLE001
            cycle.check(False, f"{filename} was refused cleanly",
                        f"{type(exc).__name__}: {exc}")
    cycle.check(refused == 3, "every malformed file was refused", str(refused))
    cycle.check(len(repo.sources(cycle.session, cycle.workspace.id)) == 0,
                "no half-made source row was left behind")

    # And the workspace still works afterwards: a refusal is not a wound.
    cycle.upload("results.xlsx", workbook())
    outcome = cycle.generate(REPORT_MD.format(title="IFRS 9 Committee Report"))
    cycle.check(outcome.version_id is not None,
                "a good file still generates after the refusals")


JOURNEYS = {
    "A": ("create a report", journey_a),
    "B": ("scoped edit", journey_b),
    "C": ("source re-read", journey_c),
    "D": ("metric update and Since Last Time", journey_d),
    "E": ("dashboard refresh", journey_e),
    "F": ("reopen after a state change", journey_f),
    "G": ("restart persistence", journey_g),
    "H": ("duplicate send", journey_h),
    "I": ("version restore", journey_i),
    "J": ("malformed source", journey_j),
}


# --------------------------------------------------------------------------
# Running them
# --------------------------------------------------------------------------


def discard(workspace_id: int) -> None:
    """Remove one soak workspace and everything hanging off it.

    By id, so nothing else in the database is touched. The foreign keys
    cascade, so deleting the workspace takes its messages, jobs, sources,
    artifacts, versions and governance rows with it.
    """
    from backend.db.engine import get_session
    from backend.models.playbook import PlaybookWorkspace

    with get_session() as session:
        row = session.get(PlaybookWorkspace, workspace_id)
        if row is not None:
            session.delete(row)
            session.commit()


def run_journey(letter: str, cycles: int, write, scope) -> dict:
    from backend.db.engine import get_session

    label, journey = JOURNEYS[letter]
    print(f"\n-- {letter}. {label} " + "-" * max(0, 46 - len(label)))

    passed = failed = 0
    failures: list[str] = []
    prints: list[dict] = []
    started = time.monotonic()

    for number in range(1, cycles + 1):
        with get_session() as session:
            cycle = Cycle(session, scope, write, number)
            try:
                journey(cycle)
                if cycle.workspace is not None:
                    prints.append(fingerprint(session, cycle.workspace.id))
            except Exception as exc:                  # noqa: BLE001
                cycle.check(False, f"cycle {number} completed",
                            f"{type(exc).__name__}: {exc}")
                failures.append(traceback.format_exc(limit=6))
            finally:
                # Each cycle is its own experiment. Rolling back is what makes
                # "no state leakage" a property of the harness rather than a
                # hope about the code.
                #
                # Some journeys commit on purpose — a reopen or a restart is
                # not testable inside one open transaction — so what they
                # committed is removed by id afterwards. Removing by id and
                # never by truncation matters: this runs against a developer
                # database that also holds the seeded demonstration, and a
                # harness that wiped it would be a worse problem than the one
                # it was solving.
                session.rollback()
                if cycle.workspace is not None:
                    discard(cycle.workspace.id)

        for ok, what in cycle.checks:
            if ok:
                passed += 1
            else:
                failed += 1
                failures.append(f"cycle {number}: {what}")

    drift = ""
    if len(prints) > 1:
        for number, state in enumerate(prints[1:], start=2):
            moved = difference(prints[0], state)
            if moved:
                drift = f"cycle {number} differs from cycle 1 — {moved}"
                failed += 1
                failures.append(drift)
                break
        if not drift:
            passed += 1

    elapsed = time.monotonic() - started
    mark = "ok" if not failed else "FAILED"
    print(f"   {cycles} cycles, {passed} checks passed, {failed} failed "
          f"({elapsed:.1f}s) {mark}")
    for line in failures[:6]:
        print(f"   ! {line.strip().splitlines()[-1]}")

    return {"journey": letter, "label": label, "cycles": cycles,
            "passed": passed, "failed": failed, "seconds": round(elapsed, 1),
            "same_state_every_cycle": not drift and len(prints) > 1,
            "failures": failures}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cycles", type=int, default=10,
                        help="how many times to run each journey (default 10)")
    parser.add_argument("--only", default="",
                        help="comma-separated journey letters, e.g. A,B,C")
    parser.add_argument("--list", action="store_true",
                        help="list the journeys and exit")
    args = parser.parse_args()

    if args.list:
        for letter, (label, _) in JOURNEYS.items():
            print(f"{letter}. {label}")
        return 0

    from backend.config import settings

    if not settings.has_database:
        print("CANNOT RUN: no database configured. This is NOT a pass.")
        return 2
    if args.cycles < 1:
        print("CANNOT RUN: --cycles must be at least 1.")
        return 2

    letters = ([x.strip().upper() for x in args.only.split(",") if x.strip()]
               or list(JOURNEYS))
    unknown = [x for x in letters if x not in JOURNEYS]
    if unknown:
        print(f"CANNOT RUN: no such journey: {', '.join(unknown)}")
        return 2

    from backend.agentic.principals import tenant_of
    from backend.playbook import repository as repo

    write = install_stand_in()
    scope = repo.Scope(tenant=tenant_of(None), user_id=None)

    print("Playbook soak")
    print("=" * 72)
    print(f"{len(letters)} journey(s) x {args.cycles} cycles, "
          "deterministic provider, no paid call")

    results = [run_journey(letter, args.cycles, write, scope)
               for letter in letters]

    passed = sum(r["passed"] for r in results)
    failed = sum(r["failed"] for r in results)

    EVIDENCE.parent.mkdir(parents=True, exist_ok=True)
    EVIDENCE.write_text(json.dumps({
        "cycles_each": args.cycles,
        "journeys": results,
        "checks_passed": passed,
        "checks_failed": failed,
        "provider": "deterministic stand-in; no paid call",
    }, indent=2) + "\n", encoding="utf-8")

    print("\n" + "=" * 72)
    print(f"{len(letters)} journeys x {args.cycles} cycles: "
          f"{passed} checks passed, {failed} failed. "
          f"Evidence: {EVIDENCE.relative_to(REPO)}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
