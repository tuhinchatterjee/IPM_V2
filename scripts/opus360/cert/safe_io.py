"""
Durable, safe writing: redaction, formula-injection-safe cells, fsync'd appends.

Every byte the harness writes to evidence passes through `redact`. The API key
is never read by this module and never placed anywhere; `redact` exists so a
provider error message that echoes a header cannot carry one into evidence.
"""

from __future__ import annotations

import csv
import json
import os
import re
from collections.abc import Iterable
from pathlib import Path
from typing import Any

_SECRET_PATTERNS = (
    re.compile(r"sk-ant-[A-Za-z0-9_\-]{8,}"),
    re.compile(r"sk-[A-Za-z0-9_\-]{16,}"),
    re.compile(r"(?i)(bearer\s+)[A-Za-z0-9._\-]{8,}"),
    re.compile(r"(?i)(x-api-key['\"]?\s*[:=]\s*['\"]?)[A-Za-z0-9._\-]{8,}"),
    re.compile(r"(?i)(api[_-]?key['\"]?\s*[:=]\s*['\"]?)[A-Za-z0-9._\-]{8,}"),
)


def redact(text: str) -> str:
    out = str(text)
    for pattern in _SECRET_PATTERNS:
        if pattern.groups:
            out = pattern.sub(lambda m: m.group(1) + "[REDACTED]", out)
        else:
            out = pattern.sub("[REDACTED]", out)
    return out


def redact_obj(obj: Any) -> Any:
    if isinstance(obj, str):
        return redact(obj)
    if isinstance(obj, dict):
        return {k: redact_obj(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [redact_obj(v) for v in obj]
    return obj


_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_NUMERIC = re.compile(r"^[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?$")


def cell(value: Any) -> Any:
    """A value safe to put in a CSV/XLSX cell.

    Numbers stay numbers. A string that a spreadsheet would evaluate as a
    formula is prefixed with a single quote. A negative number written as a
    string (e.g. "-1.5") is numeric text, not a formula, and is left alone.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).upper()
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    text = redact(str(value))
    if text.startswith(_FORMULA_START) and not _NUMERIC.match(text):
        return "'" + text
    return text


def write_csv(path: Path, rows: list[dict[str, Any]],
              columns: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = list(columns or [])
    if not cols:
        seen: dict[str, None] = {}
        for row in rows:
            for key in row:
                seen.setdefault(key, None)
        cols = list(seen)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with open(tmp, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=cols, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({c: cell(row.get(c)) for c in cols})
        fh.flush()
        os.fsync(fh.fileno())
    os.replace(tmp, path)


def append_jsonl(path: Path, record: dict[str, Any]) -> None:
    """Append one line and fsync. A crash loses at most the line in flight."""
    path.parent.mkdir(parents=True, exist_ok=True)
    line = json.dumps(redact_obj(record), ensure_ascii=False, default=str)
    with open(path, "a", encoding="utf-8") as fh:
        fh.write(line + "\n")
        fh.flush()
        os.fsync(fh.fileno())


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    """Every complete line. A torn final line (crash mid-write) is skipped."""
    if not path.exists():
        return []
    out: list[dict[str, Any]] = []
    with open(path, encoding="utf-8") as fh:
        for raw in fh:
            raw = raw.strip()
            if not raw:
                continue
            try:
                out.append(json.loads(raw))
            except json.JSONDecodeError:
                continue
    return out


def write_json_atomic(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(redact_obj(obj), fh, indent=1, ensure_ascii=False, default=str)
        fh.write("\n")
        fh.flush()
        os.fsync(fh.fileno())
    os.replace(tmp, path)


def write_text_atomic(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with open(tmp, "w", encoding="utf-8") as fh:
        fh.write(redact(text))
        fh.flush()
        os.fsync(fh.fileno())
    os.replace(tmp, path)


def rows_to_xlsx(path: Path, sheets: Iterable[tuple[str, list[dict[str, Any]]]]) -> None:
    """One formatted sheet per dataset. Cells pass through `cell`."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    header_font = Font(bold=True, color="FFFFFF")
    for title, rows in sheets:
        ws = wb.create_sheet(title=title[:31])
        cols: list[str] = []
        seen: dict[str, None] = {}
        for row in rows:
            for key in row:
                seen.setdefault(key, None)
        cols = list(seen)
        if not cols:
            ws.append(["(no rows)"])
            continue
        ws.append(cols)
        for i, _ in enumerate(cols, 1):
            c = ws.cell(row=1, column=i)
            c.fill, c.font = header_fill, header_font
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for row in rows:
            values = []
            for col in cols:
                v = cell(row.get(col))
                if isinstance(v, str) and len(v) > 32000:
                    v = v[:32000] + "…[truncated for xlsx; full value in CSV/JSONL]"
                values.append(v)
            ws.append(values)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(cols, 1):
            width = min(60, max(10, len(col) + 2))
            ws.column_dimensions[get_column_letter(i)].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, path)
