"""Recompute every formula in a workbook and compare it with its cached value.

The defect this exists for
--------------------------
xlsxwriter writes a formula AND a cached result. Passing the right cached
result makes a wrong formula invisible: the file opens showing the correct
number, and recalculates to a different one the first time somebody presses
F9. A residual written `=B9-B8` when the two values sat on rows 9 and 10
displayed SAR 0.00 and would have recomputed to the whole baseline.

So the formulas are evaluated here against the OTHER cells' cached values,
independently of the code that wrote them, and any disagreement is reported.

    .venv/bin/python scripts/retail_uat/check_workbook_formulas.py FILE
"""
from __future__ import annotations

import re
import sys

import openpyxl

CELL = re.compile(r"^([A-Z]{1,3})(\d+)$")
SIMPLE = re.compile(r"^=([A-Z]{1,3}\d+)([-+*/])([A-Z]{1,3}\d+)$")
REFERENCE = re.compile(r"^=([A-Z]{1,3}\d+)$")
SUBTOTAL = re.compile(
    r"^=SUBTOTAL\(1?09,([A-Z]{1,3})(\d+):([A-Z]{1,3})(\d+)\)$")
DIVIDE_GUARD = re.compile(
    r'^=IF\(([A-Z]{1,3}\d+)=0,"",([A-Z]{1,3}\d+)/([A-Z]{1,3}\d+)\)$')

TOLERANCE = 0.01


def _number(value) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def check(path: str) -> int:
    formulas = openpyxl.load_workbook(path)
    values = openpyxl.load_workbook(path, data_only=True)
    checked = agreed = 0
    problems: list[str] = []

    for sheet in formulas.worksheets:
        cached = values[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                text = cell.value
                if not isinstance(text, str) or not text.startswith("="):
                    continue
                shown = _number(cached[cell.coordinate].value)
                where = f"{sheet.title}!{cell.coordinate}"
                got = _evaluate(text, cached)
                if got is None:
                    problems.append(f"{where}: {text} — not evaluated here")
                    continue
                checked += 1
                if shown is None:
                    problems.append(f"{where}: {text} — no cached value")
                    continue
                if abs(got - shown) <= TOLERANCE:
                    agreed += 1
                else:
                    problems.append(
                        f"{where}: {text} recomputes to {got:,.2f} but the "
                        f"file shows {shown:,.2f}")

    print(f"formulas evaluated: {checked}   agreeing with their cached "
          f"value: {agreed}")
    for one in problems:
        print("  PROBLEM", one)
    return 1 if problems else 0


def _evaluate(text: str, cached) -> float | None:
    """Only the shapes this workbook writes. Anything else is reported."""
    hit = SUBTOTAL.match(text)
    if hit:
        column, first, _, last = hit.groups()
        total = 0.0
        for number in range(int(first), int(last) + 1):
            one = _number(cached[f"{column}{number}"].value)
            if one is not None:
                total += one
        return total
    hit = DIVIDE_GUARD.match(text)
    if hit:
        guard, top, bottom = hit.groups()
        if not _number(cached[guard].value):
            return None
        a, b = _number(cached[top].value), _number(cached[bottom].value)
        return None if (a is None or not b) else a / b
    hit = SIMPLE.match(text)
    if hit:
        left, operator, right = hit.groups()
        a, b = _number(cached[left].value), _number(cached[right].value)
        if a is None or b is None:
            return None
        return {"-": a - b, "+": a + b, "*": a * b,
                "/": (a / b if b else None)}[operator]
    hit = REFERENCE.match(text)
    if hit:
        return _number(cached[hit.group(1)].value)
    return None


if __name__ == "__main__":
    raise SystemExit(check(sys.argv[1]))
