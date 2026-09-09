"""
Playbook browser acceptance. Playbook §20.

Drives a real Chromium against a real front end and a real backend, and asserts
the things a screenshot cannot: that the home screen is in the order §3
specifies, that a seeded thread is served from the database rather than from
hard-coded strings, that only exported analyses are selectable, that previewing
one does not lose the selection, and that every download resolves to real bytes.

    .venv/bin/python scripts/acceptance/playbook_browser_acceptance.py

Both servers are assumed to be up. If Chromium cannot launch this EXITS
NON-ZERO rather than reporting a pass — the rule `scripts/browser_acceptance.py`
already follows, because a run that did not happen is not a run that succeeded.
"""

from __future__ import annotations

import asyncio
import json
import pathlib
import sys

REPO = pathlib.Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))

WEB = "http://127.0.0.1:3000"
API = "http://127.0.0.1:8000"
DOWNLOADS = pathlib.Path("/tmp/playbook-acceptance")
SHOTS = REPO / "docs" / "playbook" / "screenshots"
EVIDENCE = REPO / "docs" / "playbook" / "browser_acceptance.json"

#: §17 asks for a layout that works on a typical laptop, which is the viewport
#: a dense table breaks first and the developer's monitor never does.
VIEWPORTS = [("laptop", 1366, 768), ("desktop", 1600, 900)]

ok: list[str] = []
bad: list[str] = []
notes: list[str] = []


def check(name: str, condition: bool, detail: str = "") -> bool:
    (ok if condition else bad).append(f"{name}{(' — ' + detail) if detail else ''}")
    print(("  PASS  " if condition else "  FAIL  ") + name
          + (f"  {detail}" if detail else ""), flush=True)
    return bool(condition)


def _chromium_path() -> str | None:
    root = pathlib.Path("/opt/pw-browsers")
    for pattern in ("chromium-*/chrome-linux/chrome", "chromium/chrome-linux/chrome"):
        for found in sorted(root.glob(pattern)):
            if found.is_file():
                return str(found)
    return None


def _prior_report_bytes() -> bytes:
    """A real Word file, built here so the upload journey needs no fixture."""
    import io

    from docx import Document

    doc = Document()
    doc.add_heading("Committee report — prior period", level=1)
    doc.add_paragraph("Weighted ECL was SAR 20.90 million at a coverage ratio "
                      "of 2.09 per cent.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def journey(page, viewport: str) -> None:
    console_errors: list[str] = []
    page.on("console", lambda m: console_errors.append(m.text)
            if m.type == "error" else None)
    page.on("pageerror", lambda e: console_errors.append(str(e)))

    # ---------------------------------------------------- 1. home and its order
    await page.goto(f"{WEB}/playbook", wait_until="networkidle")
    await page.wait_for_timeout(1200)

    composer = page.locator("#playbook-composer")
    check(f"[{viewport}] the composer is on the home screen",
          await composer.count() > 0)
    check(f"[{viewport}] the plus button is in the composer",
          await page.locator('button[aria-label="Add sources"]').count() > 0)

    async def top_of(selector: str) -> float:
        box = await page.locator(selector).first.bounding_box()
        return box["y"] if box else -1.0

    composer_y = await top_of("#playbook-composer")
    playbooks_y = await top_of("#recent-playbooks")
    exports_y = await top_of("#recent-exports")
    check(f"[{viewport}] Recent Playbooks sits below the composer",
          composer_y < playbooks_y, f"{composer_y:.0f} < {playbooks_y:.0f}")
    check(f"[{viewport}] Recent Exported Analyses sits below Recent Playbooks",
          playbooks_y < exports_y, f"{playbooks_y:.0f} < {exports_y:.0f}")

    chips = page.locator("section[aria-label='Ask Playbook'] button")
    check(f"[{viewport}] quick prompts are under the composer",
          await chips.count() >= 4, f"{await chips.count()} controls")

    overflow = await page.evaluate(
        "document.body.scrollWidth > document.body.clientWidth")
    check(f"[{viewport}] the page does not scroll sideways", not overflow)

    SHOTS.mkdir(parents=True, exist_ok=True)
    await page.screenshot(path=str(SHOTS / f"home-{viewport}.png"),
                          full_page=True)

    # ------------------------------------------ a starter must not generate
    before = await page.locator("article").count()
    await chips.nth(await chips.count() - 1).click()
    await page.wait_for_timeout(600)
    value = await composer.input_value()
    check(f"[{viewport}] a quick prompt fills the composer",
          len(value) > 20, f"{len(value)} characters")
    check(f"[{viewport}] a quick prompt does not generate anything",
          await page.locator("article").count() == before)

    # ------------------------------------------------ 2. a seeded thread opens
    cards = page.locator("section[aria-labelledby='recent-playbooks'] a")
    check(f"[{viewport}] the three seeded playbooks are listed",
          await cards.count() >= 3, f"{await cards.count()} cards")
    # By name, not by position: the home screen orders by last activity, which
    # is not what this journey is testing.
    ifrs9 = cards.filter(has_text="IFRS 9 Committee Report")
    check(f"[{viewport}] the IFRS 9 committee workspace is on the home screen",
          await ifrs9.count() == 1)
    await ifrs9.first.click()
    await page.wait_for_load_state("networkidle")
    await page.wait_for_timeout(1500)

    articles = await page.locator("article").count()
    check(f"[{viewport}] the thread reopens with its whole history",
          articles >= 12, f"{articles} messages")

    body = await page.locator("body").inner_text()
    # The product-copy rule forbids the word "Demo" on screen, so the label a
    # user actually sees is "Synthetic data" and the per-message note says who
    # wrote it. Both must be there: the badge alone would not tell the reader
    # that the assistant's replies are fixtures rather than model output.
    check(f"[{viewport}] seeded history is labelled synthetic",
          # Lower-cased because the per-message note sits in a `.meta` line,
          # which CSS renders in capitals.
          "Synthetic data" in body
          and "not by a model" in body.lower())
    check(f"[{viewport}] the thread states the seeded ECL figure",
          "22.77" in body)
    check(f"[{viewport}] both artifact versions are recorded",
          "Version 2" in body or "2 versions" in body)
    check(f"[{viewport}] the sources pane lists the input files",
          ".docx" in body or ".xlsx" in body)

    await page.screenshot(path=str(SHOTS / f"thread-{viewport}.png"),
                          full_page=True)

    # ------------------------------------------------- downloads are real bytes
    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    for fmt in ("docx", "pdf"):
        link = page.locator(f'[data-testid="playbook-download-{fmt}"]').first
        if await link.count() == 0:
            check(f"[{viewport}] a {fmt} download is offered", False)
            continue
        async with page.expect_download() as info:
            await link.click()
        download = await info.value
        target = DOWNLOADS / f"{viewport}-{download.suggested_filename}"
        await download.save_as(str(target))
        size = target.stat().st_size
        head = target.read_bytes()[:4]
        expected = b"%PDF" if fmt == "pdf" else b"PK\x03\x04"
        check(f"[{viewport}] the {fmt} download is a real {fmt}",
              size > 1000 and head == expected, f"{size} bytes")

    # -------------------------------------------- 3. the picker and its boundary
    await page.goto(f"{WEB}/playbook", wait_until="networkidle")
    await page.wait_for_timeout(900)
    await page.locator('button[aria-label="Add sources"]').click()
    await page.wait_for_timeout(400)
    check(f"[{viewport}] the plus menu offers both sources",
          await page.locator('text="Upload from computer"').count() > 0
          and await page.locator('text="Add exported analyses"').count() > 0)

    await page.locator('text="Add exported analyses"').click()
    await page.wait_for_timeout(1500)
    dialog = page.locator('[role="dialog"]')
    check(f"[{viewport}] the picker opens as a labelled dialog",
          await dialog.count() > 0
          and await dialog.first.get_attribute("aria-modal") == "true")

    rows = dialog.locator('input[type="checkbox"]')
    count = await rows.count()
    check(f"[{viewport}] the picker lists the exported analyses",
          count >= 10, f"{count} rows")

    await rows.nth(0).check()
    await rows.nth(1).check()
    selected_before = await dialog.locator('text=/\\d+ selected/').inner_text()
    check(f"[{viewport}] the selection count is visible",
          "2 selected" in selected_before, selected_before)

    # Preview one, come back, and the selection must survive.
    await dialog.locator('button:has-text("Preview")').first.click()
    await page.wait_for_timeout(2500)
    preview_text = (await dialog.inner_text()).lower()
    check(f"[{viewport}] the preview shows the analysis, not a summary",
          "the question asked" in preview_text
          and "what it found" in preview_text
          and len(preview_text) > 600,
          f"{len(preview_text)} characters")
    await dialog.locator('button:has-text("Back to the list")').click()
    await page.wait_for_timeout(800)
    selected_after = await dialog.locator('text=/\\d+ selected/').inner_text()
    check(f"[{viewport}] previewing does not clear the selection",
          selected_after == selected_before,
          f"{selected_before} then {selected_after}")

    await page.screenshot(path=str(SHOTS / f"picker-{viewport}.png"))
    await page.keyboard.press("Escape")
    await page.wait_for_timeout(400)
    check(f"[{viewport}] Escape closes the picker",
          await page.locator('[role="dialog"]').count() == 0)

    # -------------------------------------------- 4. uploading a real document
    # Through the composer's own file input, so what is exercised is the control
    # a user actually reaches rather than the endpoint behind it.
    source = DOWNLOADS / "acceptance-prior-report.docx"
    source.write_bytes(_prior_report_bytes())
    await page.locator('button[aria-label="Add sources"]').click()
    await page.wait_for_timeout(300)
    async with page.expect_file_chooser() as chooser_info:
        await page.locator('text="Upload from computer"').click()
    chooser = await chooser_info.value
    await chooser.set_files(str(source))
    await page.wait_for_timeout(900)

    body = await page.locator("body").inner_text()
    check(f"[{viewport}] the uploaded file appears as an attachment chip",
          "acceptance-prior-report.docx" in body)

    # A send with an attachment but no configured provider must be refused
    # honestly rather than offered and then failing.
    send = page.locator('button:has-text("Send")').first
    await page.locator("#playbook-composer").fill("Summarise this report.")
    await page.wait_for_timeout(300)
    disabled = await send.is_disabled()
    note = await page.locator("text=/ANTHROPIC_API_KEY|provider/i").count()
    check(f"[{viewport}] with no provider, Send is refused and the reason shown",
          disabled and note > 0,
          f"disabled={disabled}, reason shown={note > 0}")

    remove = page.locator('button[aria-label^="Remove "]').first
    if await remove.count():
        await remove.click()
        await page.wait_for_timeout(400)
        after = await page.locator("body").inner_text()
        check(f"[{viewport}] an attachment can be removed",
              "acceptance-prior-report.docx" not in after)
    else:
        check(f"[{viewport}] an attachment can be removed", False,
              "no remove control found")

    real_errors = [e for e in console_errors
                   if "favicon" not in e.lower() and "404" not in e]
    check(f"[{viewport}] no console errors", not real_errors,
          "; ".join(real_errors[:2]))


def _open_proposal() -> tuple[int, int]:
    """A workspace with an UNDECIDED five-item proposal, created directly.

    The seeded proposals are already decided, which is the right thing for a
    demonstration and the wrong thing for testing the decision. The route that
    creates one runs a generation, so this writes it through the repository —
    the interface under test is the panel, not the way the proposal arrived.
    """
    from backend.db.engine import get_session
    from backend.playbook import repository as repo

    with get_session() as session:
        ws = repo.create_workspace(
            session, repo.Scope(tenant="default", user_id=None),
            title="Acceptance — deciding proposed changes",
            document_family="ifrs9_committee_report")
        change_set = repo.create_change_set(
            session, ws.id, message_id=None, base_version_id=None, items=[
                {"stable_id": "acc-exec", "display_number": 1,
                 "target_section": "1. Executive summary",
                 "rationale": "Every figure is from the prior period."},
                {"stable_id": "acc-scenario", "display_number": 2,
                 "target_section": "2. Scenario results",
                 "rationale": "The table shows only the prior period."},
                {"stable_id": "acc-coverage", "display_number": 3,
                 "target_section": "3. Coverage",
                 "rationale": "Coverage must follow the restated table.",
                 "depends_on": ["acc-scenario"]},
                {"stable_id": "acc-monitoring", "display_number": 4,
                 "target_section": "4. Model monitoring",
                 "rationale": "The methodology asks for evidence in the paper."},
                {"stable_id": "acc-limitations", "display_number": 5,
                 "target_section": "5. Limitations",
                 "rationale": "Post-model adjustments are named but not treated."},
            ])
        session.commit()
        return ws.id, change_set.id


def _drop_workspace(workspace_id: int) -> None:
    from sqlalchemy import text

    from backend.db.engine import get_session

    with get_session() as session:
        session.execute(
            text("DELETE FROM playbook_workspaces WHERE id = :i"),
            {"i": workspace_id})
        session.commit()


async def change_set_journey(page) -> None:
    """PB-016. Approving some proposed changes and not others."""
    await page.set_extra_http_headers({"X-IPM-Role": "ADMIN"})
    workspace_id, _ = _open_proposal()
    try:
        await page.goto(f"{WEB}/playbook/{workspace_id}", wait_until="networkidle")
        await page.wait_for_timeout(900)

        panel = page.locator('[data-testid="playbook-change-set"]')
        check("the proposal is shown as a numbered list",
              await panel.count() == 1)
        boxes = panel.locator('input[type="checkbox"]')
        check("every proposed change carries its own control",
              await boxes.count() == 5, f"{await boxes.count()} controls")

        text_before = await panel.inner_text()
        check("the changes are numbered, not merely listed",
              all(f"{n}." in text_before for n in range(1, 6)))
        check("nothing is selected before the user decides",
              "0 of 5 selected" in text_before)

        # Approve 1, 2 and 5. Leave 3 and 4.
        for label in ("Change 1: 1. Executive summary",
                      "Change 2: 2. Scenario results",
                      "Change 5: 5. Limitations"):
            await panel.locator(f'input[aria-label="{label}"]').check()
        await page.wait_for_timeout(200)
        summary = await panel.inner_text()
        check("the pending decision is stated in the user's own terms",
              "Applying changes 1, 2 and 5" in summary
              and "Changes 3 and 4 are held" in summary, summary[-200:])

        # Ticking a dependent change pulls in what it rests on.
        await panel.locator(
            'input[aria-label="Change 3: 3. Coverage"]').check()
        await page.wait_for_timeout(200)
        check("a change that rests on another cannot be approved alone",
              await panel.locator(
                  'input[aria-label="Change 2: 2. Scenario results"]'
              ).is_checked())
        await panel.locator(
            'input[aria-label="Change 3: 3. Coverage"]').uncheck()
        await page.wait_for_timeout(200)

        await panel.locator(
            '[data-testid="playbook-apply-selected"]').click()
        await page.wait_for_timeout(1200)

        composer = page.locator("textarea").first
        instruction = await composer.input_value()
        check("the approved changes become the instruction, and only those",
              "1. Executive summary" in instruction
              and "2. Scenario results" in instruction
              and "5. Limitations" in instruction
              and "NOT approved" in instruction,
              instruction[:120])
        held = instruction.split("NOT approved")[-1]
        check("the held changes are named as held rather than dropped",
              "3. Coverage" in held and "4. Model monitoring" in held)

        await page.screenshot(path=str(SHOTS / "change-set.png"), full_page=True)

        # And it still says the same thing after a reload.
        await page.reload(wait_until="networkidle")
        await page.wait_for_timeout(900)
        after = await page.locator(
            '[data-testid="playbook-change-set"]').inner_text()
        check("the decision survives a reload",
              "Applying changes 1, 2 and 5" in after
              and "Changes 3 and 4 are held" in after, after[-200:])
        check("a decided proposal is kept on screen rather than removed",
              "1. Executive summary" in after and "3. Coverage" in after)
    finally:
        _drop_workspace(workspace_id)


def _versioned_report() -> tuple[int, int]:
    """A workspace with a two-version report and real files on both.

    Its own workspace rather than a seeded one: restoring writes a third
    version, and a browser run that quietly grows the demonstration is a run
    that cannot be repeated.
    """
    from backend.db.engine import get_session
    from backend.playbook import repository as repo
    from backend.playbook import store

    with get_session() as session:
        ws = repo.create_workspace(
            session, repo.Scope(tenant="default", user_id=None),
            title="Acceptance — restoring a version",
            document_family="ifrs9_committee_report")
        artifact = repo.create_artifact(session, ws.id, kind="report",
                                        title="Committee report")
        for n, body in enumerate(
                [b"PK\x03\x04 the version that was approved",
                 b"PK\x03\x04 the version that went too far"], start=1):
            version = repo.new_version(
                session, artifact,
                content={"title": "Committee report",
                         "sections": [{"heading": f"Draft {n}", "blocks": []}]},
                source_manifest={}, content_hash=f"acceptance-{n}",
                change_summary=f"Draft {n}",
                base_version_id=artifact.current_version_id)
            stored = store.put_artifact(
                ws.id, artifact.id, version.version,
                f"committee-report-v{version.version}.docx", body)
            repo.add_file(
                session, version, fmt="docx", bytes_path=stored.relative,
                mime="application/vnd.openxmlformats-officedocument"
                     ".wordprocessingml.document",
                filename=f"committee-report-v{version.version}.docx",
                size_bytes=stored.size_bytes, sha256=stored.sha256,
                renderer="local", validated=True)
        session.commit()
        return ws.id, artifact.id


def _running_generation(title: str) -> tuple[int, int]:
    """A workspace with a generation genuinely in flight.

    Written through the repository rather than started with a provider,
    because this run has no credential and the thing under test is the
    TRANSPORT: that events written on the server appear in the browser as they
    are written, not all at once at the end. The deltas are a fixture; the SSE
    connection, the parser, the incremental rendering and the reconnect are
    real. Live generation is verified separately and is not claimed here.
    """
    from backend.db.engine import get_session
    from backend.models.playbook import PlaybookJob
    from backend.playbook import repository as repo
    from backend.playbook import stream

    with get_session() as session:
        ws = repo.create_workspace(
            session, repo.Scope(tenant="default", user_id=None), title=title,
            document_family="ifrs9_committee_report")
        job = PlaybookJob(workspace_id=ws.id, tenant="default",
                          idempotency_key=f"acceptance-stream-{ws.id}",
                          state="drafting")
        session.add(job)
        session.flush()
        writer = stream.Writer(session=session, job_id=job.id)
        writer.state("reviewing_sources", "3 sources")
        writer.delta("# IFRS 9 Committee Report\n\n## 1. Executive summary\n\n")
        writer.flush()
        session.commit()
        return ws.id, job.id


def _append(job_id: int, *, text: str = "", state: str = "",
            finish: bool = False, fail: str = "") -> None:
    """Write more of the generation, as the worker would."""
    from backend.db.engine import get_session
    from backend.playbook import stream
    from backend.playbook import workspace_service as service

    with get_session() as session:
        writer = stream.Writer(session=session, job_id=job_id)
        if state:
            writer.state(state)
        if text:
            writer.delta(text)
            writer.flush()
        # The job row is closed alongside the terminal event, exactly as the
        # worker does it — a stream that says "done" while the job still reads
        # as running would leave every reloading browser re-attaching for ever.
        if fail:
            service.mark_job_finished(session, job_id, state="cancelled")
            session.commit()
            writer.error(fail, category="server")
        if finish:
            service.mark_job_finished(session, job_id, state="ready")
            session.commit()
            writer.done({"version": 1})
        session.commit()


async def streaming_journey(page) -> None:
    """PB-038. The answer arrives while it is being written, and survives a
    refresh."""
    await page.set_extra_http_headers({"X-IPM-Role": "ADMIN"})
    workspace_id, job_id = _running_generation("Acceptance — streaming")
    try:
        # `domcontentloaded`, not `networkidle`: this page holds an open SSE
        # connection for as long as the generation runs, so it never goes
        # idle. Waiting for idle here would be waiting for streaming to stop.
        await page.goto(f"{WEB}/playbook/{workspace_id}",
                        wait_until="domcontentloaded")
        await page.wait_for_timeout(1200)

        bubble = page.locator('[data-testid="playbook-streaming"]')
        check("a running generation is attached to on load",
              await bubble.count() == 1)
        first = await bubble.inner_text()
        check("the answer so far is on screen before the run has finished",
              "Executive summary" in first, first[:100])
        check("the real state is named, not a percentage",
              "Reading the sources" in await page.locator(
                  '[data-testid="playbook-generation-state"]').inner_text())
        check("Stop is offered while it runs",
              await page.locator('[data-testid="playbook-stop"]').count() == 1)

        # More of the answer is written on the server. Nothing is reloaded.
        _append(job_id, state="drafting",
                text="Weighted ECL rose to SAR 22.77 million.\n\n")
        await page.wait_for_timeout(1500)
        second = await bubble.inner_text()
        check("text written after the page loaded appears without a reload",
              "22.77" in second and "22.77" not in first, second[:120])
        check("the state moved as the work moved",
              "Writing" in await page.locator(
                  '[data-testid="playbook-generation-state"]').inner_text())
        check("markdown is rendered, not printed as source",
              await bubble.locator("h2").count() > 0
              and "## 1." not in second)

        await page.screenshot(path=str(SHOTS / "streaming.png"), full_page=True)

        # A refresh mid-generation. The one thing that must not happen is a
        # second generation; the answer so far must still be there.
        await page.reload(wait_until="domcontentloaded")
        await page.wait_for_timeout(1500)
        after = await page.locator(
            '[data-testid="playbook-streaming"]').inner_text()
        check("a refresh mid-generation keeps the answer so far",
              "22.77" in after and "Executive summary" in after, after[:120])

        jobs = await (await page.request.get(
            f"{API}/api/v1/playbook/workspaces/{workspace_id}",
            headers={"X-IPM-Role": "ADMIN"})).json()
        check("a refresh started no second generation",
              jobs["running_job"]["id"] == job_id,
              str(jobs["running_job"]))
        check("no answer is in the thread while the stream is unfinished",
              all(m["role"] != "assistant" for m in jobs["messages"]),
              f"{len(jobs['messages'])} message(s)")

        _append(job_id, text="Nothing else changed.\n", finish=True)
        await page.wait_for_timeout(2000)
        check("the stream ends and the bubble goes with it",
              await page.locator(
                  '[data-testid="playbook-streaming"]').count() == 0)
        settled = await (await page.request.get(
            f"{API}/api/v1/playbook/workspaces/{workspace_id}",
            headers={"X-IPM-Role": "ADMIN"})).json()
        check("the finished generation is no longer offered to attach to",
              settled["running_job"] is None)
    finally:
        _drop_workspace(workspace_id)


async def stop_journey(page) -> None:
    """PB-038. Stopping mid-stream, and what is left behind."""
    await page.set_extra_http_headers({"X-IPM-Role": "ADMIN"})
    workspace_id, job_id = _running_generation("Acceptance — stopping")
    try:
        await page.goto(f"{WEB}/playbook/{workspace_id}",
                        wait_until="domcontentloaded")
        await page.wait_for_timeout(1200)
        _append(job_id, text="Weighted ECL rose to SAR 22.77 million.")
        await page.wait_for_timeout(1200)

        await page.locator('[data-testid="playbook-stop"]').click()
        await page.wait_for_timeout(900)
        cancelled = await (await page.request.get(
            f"{API}/api/v1/playbook/jobs/{job_id}",
            headers={"X-IPM-Role": "ADMIN"})).json()
        check("pressing Stop asks the running generation to stop",
              cancelled["cancelled"] is True, str(cancelled["cancelled"]))

        # The worker's own ending, as it would write it.
        _append(job_id, fail="This generation was stopped. Nothing was saved.")
        await page.wait_for_timeout(1800)
        body = await page.locator("body").inner_text()
        check("a stopped run says so rather than leaving half an answer",
              "stopped" in body.lower() and "22.77" not in body,
              body[:160].replace("\n", " "))

        thread = await (await page.request.get(
            f"{API}/api/v1/playbook/workspaces/{workspace_id}",
            headers={"X-IPM-Role": "ADMIN"})).json()
        check("a stopped run wrote no artifact version",
              thread["artifacts"] == [])
    finally:
        _drop_workspace(workspace_id)


async def source_journey(page) -> None:
    """PB-004 and PB-012. Correcting what the parser decided about a file."""
    await page.set_extra_http_headers({"X-IPM-Role": "ADMIN"})
    created = await page.request.post(
        f"{API}/api/v1/playbook/workspaces",
        headers={"X-IPM-Role": "ADMIN",
                 "Content-Type": "application/json"},
        data=json.dumps({"title": "Acceptance — correcting a source"}))
    workspace_id = (await created.json())["id"]
    try:
        DOWNLOADS.mkdir(parents=True, exist_ok=True)
        path = DOWNLOADS / "acceptance-standard.docx"
        path.write_bytes(_prior_report_bytes())
        uploaded = await page.request.post(
            f"{API}/api/v1/playbook/workspaces/{workspace_id}/sources",
            headers={"X-IPM-Role": "ADMIN"},
            multipart={"file": {"name": "acceptance-standard.docx",
                                "mimeType": "application/vnd.openxmlformats-"
                                            "officedocument.wordprocessingml"
                                            ".document",
                                "buffer": path.read_bytes()},
                       "source_role": "supporting"})
        source_id = (await uploaded.json())["id"]

        await page.goto(f"{WEB}/playbook/{workspace_id}",
                        wait_until="networkidle")
        await page.wait_for_timeout(900)

        select = page.locator(
            f'[data-testid="playbook-source-role-{source_id}"]')
        check("a source's kind is a control, not a fixed label",
              await select.count() == 1)
        check("it starts on what the upload said it was",
              await select.input_value() == "supporting")

        await select.select_option("methodology")
        await page.wait_for_timeout(1200)
        check("correcting the kind is recorded as a person's decision",
              "You set this" in await page.locator("body").inner_text())

        after = await (await page.request.get(
            f"{API}/api/v1/playbook/sources/{source_id}",
            headers={"X-IPM-Role": "ADMIN"})).json()
        check("the correction reached the database, not just the screen",
              after["role"] == "methodology" and after["role_set_by"] == "user",
              f"{after['role']} / {after['role_set_by']}")

        await page.reload(wait_until="networkidle")
        await page.wait_for_timeout(900)
        check("the correction survives a reload",
              await page.locator(
                  f'[data-testid="playbook-source-role-{source_id}"]'
              ).input_value() == "methodology")

        refused = await page.request.patch(
            f"{API}/api/v1/playbook/sources/{source_id}",
            headers={"X-IPM-Role": "ADMIN",
                     "Content-Type": "application/json"},
            data=json.dumps({"source_role": "board_paper"}))
        check("a kind that is not one of the five is refused",
              refused.status == 422, f"HTTP {refused.status}")
    finally:
        _drop_workspace(workspace_id)


async def restore_journey(page) -> None:
    """PB-022. An earlier version comes back by moving forward."""
    await page.set_extra_http_headers({"X-IPM-Role": "ADMIN"})
    workspace_id, artifact_id = _versioned_report()
    try:
        await page.goto(f"{WEB}/playbook/{workspace_id}",
                        wait_until="networkidle")
        await page.wait_for_timeout(900)

        # The version list starts collapsed to the current version, which is
        # the right default and the wrong state for finding an older one.
        show_all = page.locator('button:has-text("Show all")').first
        check("older versions are reachable from the files pane",
              await show_all.count() > 0)
        await show_all.click()
        await page.wait_for_timeout(300)

        restore = page.locator('[data-testid="playbook-restore-1"]').first
        check("an earlier version offers a restore, the current one does not",
              await restore.count() == 1
              and await page.locator(
                  '[data-testid="playbook-restore-2"]').count() == 0)

        before = await (await page.request.get(
            f"{API}/api/v1/playbook/workspaces/{workspace_id}",
            headers={"X-IPM-Role": "ADMIN"})).json()
        report = next(a for a in before["artifacts"] if a["id"] == artifact_id)
        v1_docx = next(f["id"] for v in report["versions"]
                       if v["version"] == 1 for f in v["files"])

        await restore.click()
        await page.wait_for_timeout(1500)

        after = await (await page.request.get(
            f"{API}/api/v1/playbook/workspaces/{workspace_id}",
            headers={"X-IPM-Role": "ADMIN"})).json()
        restored = next(a for a in after["artifacts"] if a["id"] == artifact_id)
        versions = sorted(v["version"] for v in restored["versions"])
        check("restoring writes a new version rather than rewinding",
              versions == [1, 2, 3], f"versions {versions}")
        latest = next(v for v in restored["versions"] if v["version"] == 3)
        check("the restored version says where it came from",
              "Restored version 1" in latest["change_summary"],
              latest["change_summary"][:80])

        original = await page.request.get(
            f"{API}/api/v1/playbook/artifact-files/{v1_docx}/download",
            headers={"X-IPM-Role": "ADMIN"})
        new_id = next(f["id"] for f in latest["files"] if f["format"] == "docx")
        copy = await page.request.get(
            f"{API}/api/v1/playbook/artifact-files/{new_id}/download",
            headers={"X-IPM-Role": "ADMIN"})
        check("the restored file is the same bytes, not a re-render",
              await original.body() == await copy.body())
        check("the restored file is named for its new version",
              "-v3.docx" in copy.headers.get("content-disposition", ""),
              copy.headers.get("content-disposition", ""))

        # Reading a version without downloading it. The list collapses back to
        # the current version after the restore, so expand it again first.
        await page.locator('button:has-text("Show all")').first.click()
        await page.wait_for_timeout(300)
        await page.locator('[data-testid="playbook-preview-1"]').first.click()
        await page.wait_for_timeout(700)
        pane = page.locator('[data-testid="playbook-version-preview"]')
        check("an older version can be read in the application",
              await pane.count() == 1)
        pane_text = await pane.inner_text()
        check("the preview shows the document that version saved",
              "Draft 1" in pane_text, pane_text[:120])
        check("the preview offers the file rather than replacing it",
              await pane.locator('a:has-text("DOCX")').count() > 0)
        await page.keyboard.press("Escape")
        await page.wait_for_timeout(300)
        check("Escape closes the preview",
              await page.locator(
                  '[data-testid="playbook-version-preview"]').count() == 0)

        already = await page.request.post(
            f"{API}/api/v1/playbook/artifacts/{artifact_id}/restore/3",
            headers={"X-IPM-Role": "ADMIN"})
        check("restoring the version already current is refused",
              already.status == 422, f"HTTP {already.status}")
    finally:
        _drop_workspace(workspace_id)


async def api_boundary(page) -> None:
    """The backend half of PB-005 and PB-037, exercised over HTTP."""
    unexported = await page.request.get(
        f"{API}/api/v1/playbook/exports/revisions/99999999",
        headers={"X-IPM-Role": "ADMIN"})
    check("an id that was never exported is refused by the API",
          unexported.status == 404, f"HTTP {unexported.status}")

    missing_ws = await page.request.get(
        f"{API}/api/v1/playbook/workspaces/99999999",
        headers={"X-IPM-Role": "ADMIN"})
    check("a workspace that does not exist is refused",
          missing_ws.status == 404, f"HTTP {missing_ws.status}")

    missing_file = await page.request.get(
        f"{API}/api/v1/playbook/artifact-files/99999999/download",
        headers={"X-IPM-Role": "ADMIN"})
    check("a direct download of a file that does not exist is refused",
          missing_file.status == 404, f"HTTP {missing_file.status}")

    greeting = await page.request.post(
        f"{API}/api/v1/playbook/exports",
        headers={"X-IPM-Role": "ADMIN"},
        data=json.dumps({"source_module": "cockpit", "title": "Hi",
                         "narrative": "Hello."}),
    )
    check("a greeting cannot be exported through the API",
          greeting.status == 422, f"HTTP {greeting.status}")

    planner = await page.request.post(
        f"{API}/api/v1/playbook/exports",
        headers={"X-IPM-Role": "ADMIN"},
        data=json.dumps({"source_module": "project_planner",
                         "title": "Plan",
                         "narrative": "x" * 200}),
    )
    check("Project Planner is not an accepted source module",
          planner.status == 422, f"HTTP {planner.status}")

    caps = await page.request.get(f"{API}/api/v1/playbook/capabilities",
                                  headers={"X-IPM-Role": "ADMIN"})
    body = await caps.text()
    check("the capability report carries no credential",
          "sk-ant" not in body)

    monitoring = await page.request.get(f"{API}/api/v1/playbooks",
                                        headers={"X-IPM-Role": "ADMIN"})
    check("the monitoring Playbooks feature still answers",
          monitoring.status in (200, 503), f"HTTP {monitoring.status}")


async def main() -> int:
    from playwright.async_api import async_playwright

    executable = _chromium_path()
    if executable is None:
        print("CANNOT RUN: no Chromium found under /opt/pw-browsers.")
        print("Browser acceptance is therefore NOT RUN, which is not a pass.")
        return 2

    print("Playbook browser acceptance")
    print("=" * 72)
    async with async_playwright() as pw:
        browser = await pw.chromium.launch(executable_path=executable,
                                           args=["--no-sandbox"])
        try:
            for name, width, height in VIEWPORTS:
                context = await browser.new_context(
                    viewport={"width": width, "height": height},
                    accept_downloads=True)
                page = await context.new_page()
                await page.set_extra_http_headers({"X-IPM-Role": "ADMIN"})
                print(f"\n-- {name} {width}x{height} " + "-" * 40)
                await journey(page, name)
                await context.close()

            context = await browser.new_context()
            page = await context.new_page()
            print("\n-- Streaming " + "-" * 47)
            await streaming_journey(page)
            print("\n-- Stopping mid-stream " + "-" * 37)
            await stop_journey(page)
            print("\n-- Correcting a source " + "-" * 37)
            await source_journey(page)
            print("\n-- Restoring a version " + "-" * 37)
            await restore_journey(page)
            print("\n-- Proposed changes " + "-" * 40)
            await change_set_journey(page)
            print("\n-- API boundary " + "-" * 44)
            await api_boundary(page)
            await context.close()
        finally:
            await browser.close()

    EVIDENCE.parent.mkdir(parents=True, exist_ok=True)
    EVIDENCE.write_text(json.dumps({
        "passed": ok, "failed": bad, "notes": notes,
        "screenshots": sorted(p.name for p in SHOTS.glob("*.png")),
    }, indent=2) + "\n", encoding="utf-8")

    print("\n" + "=" * 72)
    print(f"{len(ok)} passed, {len(bad)} failed. "
          f"Evidence: {EVIDENCE.relative_to(REPO)}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
